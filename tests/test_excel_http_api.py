"""HTTP adapter: Excel upload / process / export endpoints.

Runs without GEMINI_API_KEY set (the normal pytest invocation — nothing
here sources .env), which exercises the real "Gemini unavailable" path:
generation must fail each eligible row closed (REVIEW_ONLY), never invent
a canned reply.
"""

from __future__ import annotations

import pytest
import datetime as dt
import io
import json
import urllib.error
import urllib.request
from urllib.parse import quote

import openpyxl

from config import settings
from rcs.review.auth import create_user
from rcs.review.http_api import serve_in_thread
from rcs.review.repository import get_engine, init_db
from rcs.review.service import ReviewWorkflowService


pytestmark = pytest.mark.usefixtures("legacy_eligibility_policy")


def _json_req(url: str, *, method="GET", data=None, cookie=None):
    body = None if data is None else json.dumps(data).encode("utf-8")
    headers = {"Content-Type": "application/json"}
    if cookie:
        headers["Cookie"] = cookie
    req = urllib.request.Request(url, data=body, headers=headers, method=method)
    try:
        with urllib.request.urlopen(req, timeout=30) as resp:
            payload = json.loads(resp.read().decode("utf-8"))
            return resp.status, payload, resp.headers
    except urllib.error.HTTPError as err:
        raw = err.read()
        try:
            payload = json.loads(raw.decode("utf-8"))
        except Exception:
            payload = {"raw": raw}
        return err.code, payload, err.headers


def _upload_req(url: str, *, data: bytes, filename: str, cookie: str):
    req = urllib.request.Request(
        url,
        data=data,
        headers={"Content-Type": "application/octet-stream", "X-Filename": quote(filename), "Cookie": cookie},
        method="POST",
    )
    try:
        with urllib.request.urlopen(req, timeout=30) as resp:
            return resp.status, json.loads(resp.read().decode("utf-8")), resp.headers
    except urllib.error.HTTPError as err:
        return err.code, json.loads(err.read().decode("utf-8")), err.headers


def _xlsx_bytes(rows: list[dict]) -> bytes:
    wb = openpyxl.Workbook()
    sheet = wb.active
    headers = list(settings.SOURCE_COLUMNS.keys())
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(h) for h in headers])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _login(engine, base: str) -> str:
    # ADMIN (not REVIEWER): these are general workflow tests, one of
    # which exercises the unscoped /api/analytics view — an ADMIN-only
    # capability as of Phase 6. See tests/test_auth.py for the dedicated
    # REVIEWER-is-blocked-from-it security test.
    create_user(engine, email="reviewer@test.local", password="testpassword123", display_name="مراجع الاختبار", role="ADMIN")
    _, _, headers = _json_req(
        f"{base}/api/login", method="POST", data={"email": "reviewer@test.local", "password": "testpassword123"}
    )
    return headers.get("Set-Cookie").split(";")[0]


def test_upload_accepts_xlsx_and_returns_token(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "EXCEL_UPLOAD_DIR", tmp_path / "uploads")
    engine = get_engine(tmp_path / "r.sqlite")
    init_db(engine)
    server, base = serve_in_thread(engine)
    try:
        cookie = _login(engine, base)
        data = _xlsx_bytes(
            [
                {
                    "المحتوى": "شكرًا جزيلاً على هذا المعجم الرائع",
                    "التاريخ": dt.datetime(2026, 8, 1),
                }
            ]
        )
        status, payload, _ = _upload_req(f"{base}/api/excel/upload", data=data, filename="زوار.xlsx", cookie=cookie)
        assert status == 200
        assert payload["valid"] is True
        assert payload["total_rows"] == 1
        assert len(payload["upload_token"]) == 32
    finally:
        server.shutdown()


def test_upload_rejects_non_xlsx_file(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "EXCEL_UPLOAD_DIR", tmp_path / "uploads")
    engine = get_engine(tmp_path / "r.sqlite")
    init_db(engine)
    server, base = serve_in_thread(engine)
    try:
        cookie = _login(engine, base)
        status, payload, _ = _upload_req(
            f"{base}/api/excel/upload", data=b"not excel", filename="malware.exe", cookie=cookie
        )
        assert status == 400
        assert payload["error"] == "unsupported_file_type"
    finally:
        server.shutdown()


def test_upload_requires_authentication(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "EXCEL_UPLOAD_DIR", tmp_path / "uploads")
    engine = get_engine(tmp_path / "r.sqlite")
    init_db(engine)
    server, base = serve_in_thread(engine)
    try:
        status, payload, _ = _upload_req(f"{base}/api/excel/upload", data=b"x", filename="a.xlsx", cookie="")
        assert status == 401
    finally:
        server.shutdown()


def test_process_runs_full_pipeline_appears_in_queue_and_exports(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "EXCEL_UPLOAD_DIR", tmp_path / "uploads")
    engine = get_engine(tmp_path / "r.sqlite")
    init_db(engine)
    server, base = serve_in_thread(engine)
    try:
        cookie = _login(engine, base)
        data = _xlsx_bytes(
            [
                {
                    "المحتوى": "jhnjkn njnun Phone Number: 111111111111j",
                    "التاريخ": dt.datetime(2026, 8, 1),
                },
                {
                    "المحتوى": "أقترح إضافة كلمة جميلة إلى المعجم لفائدة الطلاب",
                    "التاريخ": dt.datetime(2026, 8, 2),
                },
            ]
        )
        status, upload, _ = _upload_req(f"{base}/api/excel/upload", data=data, filename="batch.xlsx", cookie=cookie)
        assert status == 200 and upload["valid"] is True

        status, summary, _ = _json_req(
            f"{base}/api/excel/process", method="POST", data={"upload_token": upload["upload_token"]}, cookie=cookie
        )
        assert status == 200
        assert summary["total_file_rows"] == 2
        assert summary["batch_processed"] == 2
        assert summary["processed_total"] == 2
        assert summary["remaining"] == 0
        assert summary["excluded_total"] == 1  # the junk row, deterministic prefilter
        # No GEMINI_API_KEY in this test process -> the eligible row must
        # fail closed, never get a canned/invented draft.
        assert summary["generated_drafts_total"] == 0
        assert summary["needs_review_total"] + summary["no_response_recommended_total"] == 1
        assert "model_name" not in json.dumps(summary)  # no internal ML details leaked to the UI contract

        # No Excel-specific queue: the SAME ReviewWorkflowService/engine
        # sees the new rows immediately.
        svc = ReviewWorkflowService(engine)
        counts = svc.workflow_counts()
        assert counts["all"] == 2

        status, export_headers = None, None
        req = urllib.request.Request(f"{base}/api/excel/export", headers={"Cookie": cookie})
        with urllib.request.urlopen(req, timeout=30) as resp:
            status = resp.status
            body = resp.read()
            export_headers = resp.headers
        assert status == 200
        assert export_headers.get("Content-Disposition", "").startswith("attachment;")
        assert "reviewed_responses_" in export_headers.get("Content-Disposition", "")
        wb = openpyxl.load_workbook(io.BytesIO(body))
        assert wb.active.max_row == 3  # header + 2 rows
    finally:
        server.shutdown()


def test_export_returns_conflict_when_no_rows_exist(tmp_path):
    engine = get_engine(tmp_path / "empty.sqlite")
    init_db(engine)
    server, base = serve_in_thread(engine)
    try:
        cookie = _login(engine, base)
        req = urllib.request.Request(f"{base}/api/excel/export", headers={"Cookie": cookie})
        try:
            urllib.request.urlopen(req, timeout=10)
            assert False, "expected HTTPError"
        except urllib.error.HTTPError as err:
            assert err.code == 409
    finally:
        server.shutdown()


def test_real_suggested_response_reaches_item_detail_and_analytics_same_engine(tmp_path, monkeypatch):
    """Locks in the full Excel -> ReviewItem -> /api/queue -> /api/items/{id}
    -> reviewer action -> /api/analytics chain, all against the ONE engine
    a single serve_in_thread() call is bound to -- there is no separate
    Excel-only queue or DB. A scripted (non-network) Gemini-shaped client
    stands in for a real one so this test is deterministic; the assertion
    that matters is that the EXACT string it returns is what /api/items/{id}
    persists and returns -- not a canned/offline fallback."""
    monkeypatch.setattr(settings, "EXCEL_UPLOAD_DIR", tmp_path / "uploads")
    engine = get_engine(tmp_path / "r.sqlite")
    init_db(engine)

    class FakeClient:
        def generate(self, prompt, schema):
            return json.dumps(
                {"suggested_response": "شكرًا لملاحظتكم بشأن هذا التعليق تحديدًا."}, ensure_ascii=False
            )

    import rcs.review.excel_review_flow as flow

    monkeypatch.setattr(flow, "build_live_gemini_client", lambda: FakeClient())

    server, base = serve_in_thread(engine)
    try:
        cookie = _login(engine, base)
        # Several distinct real-shaped comments: the frozen classifiers are
        # weak (documented in FINAL_HUMAN_TEST_REPORT.md) and may route any
        # single one of these to NO_RESPONSE_RECOMMENDED instead of
        # GENERATE_DRAFT -- using several makes it overwhelmingly likely at
        # least one lands on GENERATE_DRAFT, which is what this test needs
        # to verify (the wiring), independent of which exact row it is.
        data = _xlsx_bytes(
            [
                {"المحتوى": "أقترح إضافة كلمة جميلة إلى المعجم لفائدة الطلاب", "التاريخ": dt.datetime(2026, 8, 2)},
                {"المحتوى": "شكرًا جزيلاً لكم على هذا الجهد الرائع في خدمة اللغة العربية", "التاريخ": dt.datetime(2026, 8, 3)},
                {"المحتوى": "هل يمكن توضيح معنى كلمة استبانة وأصلها اللغوي من فضلكم؟", "التاريخ": dt.datetime(2026, 8, 4)},
            ]
        )
        status, upload, _ = _upload_req(f"{base}/api/excel/upload", data=data, filename="one.xlsx", cookie=cookie)
        assert status == 200 and upload["valid"] is True

        status, summary, _ = _json_req(
            f"{base}/api/excel/process", method="POST", data={"upload_token": upload["upload_token"]}, cookie=cookie
        )
        assert status == 200
        assert summary["generated_drafts_total"] >= 1

        status, page, _ = _json_req(f"{base}/api/queue?filter=pending&sort=priority&limit=10", cookie=cookie)
        assert status == 200
        real_item = next(i for i in page["items"] if i["suggested_response"])
        cid = real_item["comment_id"]

        status, detail, _ = _json_req(f"{base}/api/items/{cid}", cookie=cookie)
        assert status == 200
        assert detail["original_comment"]  # the real imported visitor text, not fabricated
        assert detail["suggested_response"] == "شكرًا لملاحظتكم بشأن هذا التعليق تحديدًا."
        assert "SEND_AS_IS" in detail["allowed_actions"]

        status, sent, _ = _json_req(
            f"{base}/api/items/{cid}/send-as-is",
            method="POST",
            data={"expected_version": detail["row_version"]},
            cookie=cookie,
        )
        assert status == 200
        assert sent["outbound_status"] == "PENDING"
        assert sent["review_status"] == "APPROVED"

        status, analytics, _ = _json_req(f"{base}/api/analytics?range=week", cookie=cookie)
        assert status == 200
        assert analytics["outcomes"].get("SEND_AS_IS", 0) >= 1
    finally:
        server.shutdown()


def test_process_rejects_invalid_upload_token(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "EXCEL_UPLOAD_DIR", tmp_path / "uploads")
    engine = get_engine(tmp_path / "r.sqlite")
    init_db(engine)
    server, base = serve_in_thread(engine)
    try:
        cookie = _login(engine, base)
        status, payload, _ = _json_req(
            f"{base}/api/excel/process", method="POST", data={"upload_token": "not-a-real-token"}, cookie=cookie
        )
        assert status == 400
        assert payload["error"] == "invalid_upload_token"
    finally:
        server.shutdown()
