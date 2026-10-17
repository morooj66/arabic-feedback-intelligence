"""Blind-evaluation sampler, reviewer workbook and evaluation harness (no network, no real data needed)."""

from __future__ import annotations

import sys
from pathlib import Path

import pytest

_BE = Path(__file__).resolve().parents[1] / "scripts" / "offline_eval" / "blind_eval"
sys.path.insert(0, str(_BE))

import build_sample as BS  # noqa: E402
import evaluate_blind as EV  # noqa: E402
from config import settings  # noqa: E402


def _pool(n=600):
    rows = []
    for k in range(n):
        base = f"تعليق رقم {k} عن المعجم والكلمات"
        if k % 11 == 0:
            base += " الموقع لا يفتح عندي"           # technical cue
        elif k % 13 == 0:
            base += " للأسف الخدمة سيئة"               # complaint cue
        elif k % 7 == 0:
            base += " التعريف غير صحيح والصواب كذا"     # entry-error cue
        elif k % 17 == 0:
            base += " أتمنى إضافة خاصية جديدة"          # feature cue
        rows.append({"sid": f"{k:064x}", "source_file": "f.xlsx", "source_row_index": k, "text": base, "masks": 0, "lemma_present": k % 5 == 0})
    return rows


def test_sampler_is_deterministic_correctly_sized_and_disjoint():
    a = BS.select_sample(_pool(), seed=7)
    b = BS.select_sample(_pool(), seed=7)
    c = BS.select_sample(_pool(), seed=8)
    assert [r["sid"] for r in a["records"]] == [r["sid"] for r in b["records"]]           # reproducible
    assert [r["sid"] for r in a["records"]] != [r["sid"] for r in c["records"]]           # seed matters
    recs = a["records"]
    assert len(recs) == 75 and len({r["sid"] for r in recs}) == 75
    comps = [r["component"] for r in recs]
    assert comps.count("natural") == 45 and comps.count("targeted") == 30
    plan = dict(BS.TARGETED_PLAN)
    for name, target in plan.items():
        assert sum(1 for r in recs if r["stratum"] == name) == target
    assert [r["record_id"] for r in recs] == [f"BE-{i:03d}" for i in range(1, 76)]        # ids carry no stratum information
    nat_sids = {r["sid"] for r in recs if r["component"] == "natural"}
    assert not any(r["sid"] in nat_sids for r in recs if r["component"] == "targeted")


def test_sampler_reports_and_tops_up_a_small_stratum_without_touching_the_natural_pool():
    pool = [r for r in _pool() if "للأسف" not in r["text"]]                                # almost no complaint cues
    out = BS.select_sample(pool, seed=3)
    assert len(out["records"]) == 75 and sum(1 for r in out["records"] if r["component"] == "natural") == 45
    assert out["shortfall"].get("complaint_cues", 0) == 8 and len(out["redistributed"]) == 8


def test_selection_never_looks_at_anything_but_text_and_declared_cues():
    import inspect
    src = inspect.getsource(BS.select_sample)
    for forbidden in ("label", "gold", "predict", "model"):
        assert forbidden not in src.lower().replace("labelled", "")


def test_reviewer_workbook_has_no_prefilled_labels_valid_dropdowns_and_no_leaks(tmp_path):
    from openpyxl import load_workbook
    sample = BS.select_sample(_pool(), seed=5)
    BS.write_outputs(sample, {"note": "synthetic"}, tmp_path / "out")
    wb = load_workbook(tmp_path / "out" / "reviewer" / "blind_eval_75_review.xlsx")
    assert wb.sheetnames == ["Read me first", "Guidelines", "Review", "Codes"]
    ws = wb["Review"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 75 and all(r[2:7] == (None,) * 5 for r in rows)
    lists = {str(dv.sqref): dv.formula1 for dv in ws.data_validations.dataValidation}
    assert set(lists["C2:C76"].strip('"').split(",")) == set(settings.PRIMARY_INTENTS)
    assert set(lists["D2:D76"].strip('"').split(",")) == set(settings.TOPIC_SUBTYPES)
    assert set(lists["E2:E76"].strip('"').split(",")) == set(settings.SENTIMENTS)
    allcells = " ".join(str(c.value) for s in wb.worksheets for row in s.iter_rows() for c in row if c.value).lower()
    for leak in ("natural_random", "targeted", "technical_cues", "stratum", "mapping"):
        assert leak not in allcells
    mapping = (tmp_path / "out" / "coordinator_only" / "blind_eval_mapping.local.csv").read_text(encoding="utf-8")
    assert "technical_cues" in mapping or "natural" in mapping                              # the mapping, not the workbook, holds it


def _full_labels(ids, **override):
    labels = {}
    for i in ids:
        labels[i] = {"text": "t", "intent": "SUGGESTION", "topic": "WORD_SUBMISSION", "sentiment": "NEUTRAL", "insufficient": None, "notes": None}
    for k, v in override.items():
        labels[k] = {**labels[k], **v}
    return labels


def test_label_validation_enforces_the_frozen_taxonomy_and_completeness():
    ids = [f"BE-{i:03d}" for i in range(1, 6)]
    ok = EV.validate_labels(_full_labels(ids, **{"BE-005": {"intent": None, "topic": None, "sentiment": None, "insufficient": "INSUFFICIENT_CONTEXT"}}), ids)
    assert len(ok["scored"]) == 4 and ok["insufficient"] == ["BE-005"]
    for bad in ({"BE-001": {"intent": "DUA_BLESSING"}},                                      # not in the frozen taxonomy
                {"BE-001": {"topic": None}},                                                  # blank field
                {"BE-001": {"insufficient": "INSUFFICIENT_CONTEXT"}},                         # labels must be empty when insufficient
                {"BE-001": {"insufficient": "maybe"}}):
        with pytest.raises(EV.LabelError):
            EV.validate_labels(_full_labels(ids, **bad), ids)
    with pytest.raises(EV.LabelError):
        EV.validate_labels({k: v for k, v in _full_labels(ids).items() if k != "BE-002"}, ids)


def test_exact_mcnemar_and_agreement_known_values():
    assert EV.mcnemar_exact_p(0, 5) == pytest.approx(0.0625) and EV.mcnemar_exact_p(2, 3) == 1.0 and EV.mcnemar_exact_p(0, 0) == 1.0
    a = {f"r{i}": {"intent": x, "topic": "GENERAL", "sentiment": "NEUTRAL"} for i, x in enumerate(["SPAM", "SPAM", "INQUIRY", "INQUIRY"])}
    b = {f"r{i}": {"intent": x, "topic": "GENERAL", "sentiment": "NEUTRAL"} for i, x in enumerate(["SPAM", "INQUIRY", "INQUIRY", "INQUIRY"])}
    ag = EV.agreement(a, b)
    assert ag["tasks"]["intent"]["raw_agreement"] == 0.75 and ag["disagreements"]["intent"] == ["r1"]
    assert ag["tasks"]["topic"]["cohen_kappa"] == 1.0


def _synthetic_eval():
    ids = [f"BE-{i:03d}" for i in range(1, 76)]
    component = {i: ("natural" if n < 45 else "targeted") for n, i in enumerate(ids)}
    gold = {}
    for n, i in enumerate(ids):
        topic = ["WORD_SUBMISSION", "ENTRY_CORRECTION", "GENERAL"][n % 3]
        gold[i] = {"intent": "COMPLAINT" if n % 10 == 0 else "SUGGESTION", "topic": "TECHNICAL_ISSUE" if n % 10 == 0 else topic,
                   "sentiment": "NEUTRAL"}
    base = {i: dict(g) for i, g in gold.items()}
    for n, i in enumerate(ids):
        if n % 4 == 0:
            base[i]["topic"] = "WORD_SUBMISSION"
            base[i]["intent"] = "SUGGESTION"
    return ids, component, gold, base


def test_natural_and_rare_class_results_are_separate_and_never_pooled():
    ids, component, gold, base = _synthetic_eval()
    cand = {i: dict(p) for i, p in base.items()}
    res1 = EV.evaluate(gold, {"base": base, "cand": cand}, component, [("base", "cand")])
    assert res1["never_pooled"] is True and set(res1) == {"A_natural", "B_rare_class_diagnostics", "never_pooled"}
    assert res1["A_natural"]["n_scored"] == 45 and res1["B_rare_class_diagnostics"]["n_scored"] == 30
    # changing ONLY targeted-record predictions must not move any natural (A) number, and vice versa
    base2 = {i: dict(p) for i, p in base.items()}
    for i in ids:
        if component[i] == "targeted":
            base2[i]["intent"] = "SPAM"
    res2 = EV.evaluate(gold, {"base": base2, "cand": cand}, component, [("base", "cand")])
    assert res2["A_natural"]["systems"]["base"] == res1["A_natural"]["systems"]["base"]
    assert res2["B_rare_class_diagnostics"]["systems"]["base"] != res1["B_rare_class_diagnostics"]["systems"]["base"]
    base3 = {i: dict(p) for i, p in base.items()}
    for i in ids:
        if component[i] == "natural":
            base3[i]["intent"] = "SPAM"
    res3 = EV.evaluate(gold, {"base": base3, "cand": cand}, component, [("base", "cand")])
    assert res3["B_rare_class_diagnostics"]["systems"]["base"] == res1["B_rare_class_diagnostics"]["systems"]["base"]
    text = EV.render(res1, {"sample_ids_sha256": "0" * 64, "labels_sha256": "0" * 64, "candidate_spec_sha256": "0" * 64, "reviewers": 1, "insufficient": 0})
    assert "never combined" in text and "NOT an accuracy estimate" in text and "overall accuracy" not in text.lower().replace("no pooled", "")


def test_regressions_and_boundary_confusion_are_counted_per_component():
    ids, component, gold, base = _synthetic_eval()
    cand = {i: dict(p) for i, p in base.items()}
    broke = [i for i in ids if component[i] == "natural" and base[i]["topic"] == gold[i]["topic"]][:3]
    for i in broke:
        cand[i]["topic"] = "GENERAL" if gold[i]["topic"] != "GENERAL" else "WORD_SUBMISSION"
    res = EV.evaluate(gold, {"base": base, "cand": cand}, component, [("base", "cand")])
    topic = res["A_natural"]["comparisons"]["base -> cand"]["topic"]
    assert topic["regressions"] == 3 and set(topic["regression_ids"]) == set(broke)
    conf = res["A_natural"]["ws_ec_confusion"]["base"]["matrix_rows_gold_cols_pred"]
    assert set(conf) == {"WORD_SUBMISSION", "ENTRY_CORRECTION"} and set(conf["WORD_SUBMISSION"]) == {"WORD_SUBMISSION", "ENTRY_CORRECTION", "OTHER"}


def test_candidate_layer_wrapper_only_changes_ws_ec_topics():
    preds = {"a": {"intent": "SPAM", "topic": "WORD_SUBMISSION", "sentiment": "NEUTRAL"}, "b": {"intent": "SPAM", "topic": "GENERAL", "sentiment": "NEUTRAL"}}
    out = EV.with_layer(preds, {"a": "كلمة", "b": "هذا خطأ"}, {"a": True, "b": True})
    assert out["a"]["topic"] == "ENTRY_CORRECTION" and out["b"]["topic"] == "GENERAL" and out["a"]["intent"] == "SPAM"


def test_missing_prediction_is_an_error_not_a_silent_skip():
    ids, component, gold, base = _synthetic_eval()
    partial = {i: p for i, p in base.items() if i != ids[0]}
    with pytest.raises(EV.LabelError):
        EV.evaluate(gold, {"base": partial}, component, [])


def _fill_workbook(src: Path, dst: Path, labeler):
    from openpyxl import load_workbook
    wb = load_workbook(src)
    ws = wb["Review"]
    for r in range(2, ws.max_row + 1):
        rid = ws.cell(row=r, column=1).value
        intent, topic, sent, insuff = labeler(rid)
        ws.cell(row=r, column=3, value=intent); ws.cell(row=r, column=4, value=topic)
        ws.cell(row=r, column=5, value=sent); ws.cell(row=r, column=6, value=insuff)
    wb.save(dst)


def test_cli_validates_labels_before_computing_any_prediction_and_then_reports_both_families(tmp_path, monkeypatch, capsys):
    out = tmp_path / "bev1"
    BS.write_outputs(BS.select_sample(_pool(), seed=9), {"note": "synthetic"}, out)
    monkeypatch.setattr(EV, "BASE", out)
    calls = {"n": 0}

    def fake_predict(texts):
        calls["n"] += 1
        return {i: {"intent": "SUGGESTION", "topic": "WORD_SUBMISSION", "sentiment": "NEUTRAL"} for i in texts}
    monkeypatch.setattr(EV, "compute_rcs_predictions", fake_predict)
    src = out / "reviewer" / "blind_eval_75_review.xlsx"

    blank = tmp_path / "blank.xlsx"                                           # un-labelled workbook: must be refused first
    _fill_workbook(src, blank, lambda rid: (None, None, None, None))
    monkeypatch.setattr(sys, "argv", ["evaluate_blind.py", "--labels", str(blank), "--out-dir", str(tmp_path / "res")])
    with pytest.raises(EV.LabelError):
        EV.main()
    assert calls["n"] == 0                                                    # no prediction was produced for an invalid label set

    done = tmp_path / "done.xlsx"
    _fill_workbook(src, done, lambda rid: ("SUGGESTION", "ENTRY_CORRECTION" if int(rid[3:]) % 2 else "WORD_SUBMISSION", "NEUTRAL", None))
    monkeypatch.setattr(sys, "argv", ["evaluate_blind.py", "--labels", str(done), "--out-dir", str(tmp_path / "res")])
    EV.main()
    report = (tmp_path / "res" / "blind_eval_report.md").read_text(encoding="utf-8")
    assert calls["n"] == 1 and "## A. Natural-distribution results (n = 45" in report and "## B. Rare-class diagnostics (n = 30" in report
    assert "RCS + WS/EC layer (candidate)" in report and "WORD_SUBMISSION vs ENTRY_CORRECTION confusion" in report


def test_disagreement_listing_has_ids_and_codes_only_and_matches_the_counts():
    ids, component, gold, base = _synthetic_eval()
    rows = EV.disagreement_rows(gold, {"base": base}, component)
    assert set(rows[0]) == {"system", "component", "record_id", "task", "human_label", "system_prediction"}   # no comment text
    assert all(r["human_label"] != r["system_prediction"] for r in rows)
    wrong = sum(1 for i in ids for t in EV.TASKS if base[i][t] != gold[i][t])
    assert len(rows) == wrong
    pairs = EV.top_confusions(rows)["base"]["natural"]["topic"]
    assert pairs and pairs[0][1] >= pairs[-1][1]


def test_intent_candidate_is_loaded_only_after_its_hash_manifest_verifies(tmp_path):
    import hashlib
    import joblib
    from rcs.classification import intent_candidate as IC
    X = ["الموقع ما يفتح", "كيف أبحث؟", "أتمنى خاصية", "شكرا لكم", "نص عشوائي"] * 4
    y = ["COMPLAINT", "INQUIRY", "SUGGESTION", "THANKS_PRAISE", "SPAM"] * 4
    d = tmp_path / "cand"; d.mkdir()
    joblib.dump(IC.IntentCandidate().fit(X, y), d / "model.joblib")
    (d / "config.json").write_text("{}")
    sums = "".join(f"{hashlib.sha256((d / n).read_bytes()).hexdigest()}  {n}\n" for n in ("config.json", "model.joblib"))
    (d / "SHA256SUMS").write_text(sums)
    model, sha = EV.load_verified_intent_candidate(d)
    assert len(sha) == 64
    texts = {"BE-001": "الموقع ما يفتح", "BE-002": "شكرا لكم"}
    rcs = {i: {"intent": "SPAM", "topic": "GENERAL", "sentiment": "NEUTRAL"} for i in texts}
    out = EV.with_intent_candidate(rcs, texts, model)
    assert all(out[i]["topic"] == "GENERAL" and out[i]["sentiment"] == "NEUTRAL" for i in texts)   # only intent may change
    (d / "config.json").write_text('{"tampered": true}')                                          # any modification is refused
    with pytest.raises(EV.LabelError):
        EV.load_verified_intent_candidate(d)
