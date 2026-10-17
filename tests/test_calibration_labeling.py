"""Tests for the B.2 calibration human-labeling wizard (no Gradio UI, no Gemini)."""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook

from rcs.evaluation.calibration_labeling import (
    CALIBRATION_XLSX,
    REVIEW_JSON,
    apply_human_answers,
    count_completed_rows,
    empty_review_state,
    evaluate_calibration,
    finalize_calibration_review,
    is_review_complete,
    load_calibration_workbook,
    load_or_init_review_state,
    save_review_state,
    topic_requires_subcategory,
)


def _write_mini_workbook(path: Path, rows: list[dict[str, str]]) -> None:
    columns = [
        "comment_text",
        "b2_v2_primary_intent",
        "b2_v2_topic_subtype",
        "b2_v2_sentiment",
        "b2_v3_primary_intent",
        "b2_v3_topic_subtype",
        "b2_v3_sentiment",
        "gold_primary_intent",
        "gold_topic_subtype",
        "gold_sentiment",
        "gold_annotation_status",
        "comment_id",
        "calibration_stratum",
    ]
    wb = Workbook()
    ws = wb.active
    ws.title = "Review"
    for col_idx, name in enumerate(columns, start=1):
        ws.cell(row=1, column=col_idx, value=name)
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, name in enumerate(columns, start=1):
            ws.cell(row=row_idx, column=col_idx, value=row.get(name, ""))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _mini_rows() -> list[dict[str, str]]:
    return [
        {
            "comment_id": "id1",
            "comment_text": "شكرا",
            "calibration_stratum": "short_bare_word",
            "b2_v2_primary_intent": "THANKS_PRAISE",
            "b2_v2_topic_subtype": "GENERAL",
            "b2_v2_sentiment": "POSITIVE",
            "b2_v3_primary_intent": "THANKS_PRAISE",
            "b2_v3_topic_subtype": "GENERAL",
            "b2_v3_sentiment": "POSITIVE",
        },
        {
            "comment_id": "id2",
            "comment_text": "الموقع لا يعمل",
            "calibration_stratum": "complaint_entry_correction",
            "b2_v2_primary_intent": "COMPLAINT",
            "b2_v2_topic_subtype": "TECHNICAL_ISSUE",
            "b2_v2_sentiment": "NEGATIVE",
            "b2_v3_primary_intent": "COMPLAINT",
            "b2_v3_topic_subtype": "TECHNICAL_ISSUE",
            "b2_v3_sentiment": "NEGATIVE",
        },
        {
            "comment_id": "id3",
            "comment_text": "؟؟",
            "calibration_stratum": "short_bare_word",
            "b2_v2_primary_intent": "SPAM",
            "b2_v2_topic_subtype": "GENERAL",
            "b2_v2_sentiment": "NEUTRAL",
            "b2_v3_primary_intent": "SPAM",
            "b2_v3_topic_subtype": "GENERAL",
            "b2_v3_sentiment": "NEUTRAL",
        },
    ]


@pytest.fixture()
def mini_calibration(tmp_path: Path):
    workbook = tmp_path / "sample.xlsx"
    manifest = tmp_path / "manifest.json"
    review = tmp_path / "review.json"
    labeled = tmp_path / "labeled.xlsx"
    eval_json = tmp_path / "eval.json"
    eval_txt = tmp_path / "eval.txt"
    gold_csv = tmp_path / "gold.csv"
    rows = _mini_rows()
    _write_mini_workbook(workbook, rows)
    manifest.write_text(
        json.dumps({"comment_ids_in_workbook_order": [r["comment_id"] for r in rows]}),
        encoding="utf-8",
    )
    pd.DataFrame(
        [
            {
                "comment_id": "id1",
                "annotation_stage": "ai_assisted_train",
                "ai_suggested_confidence": "0.95",
            },
            {
                "comment_id": "id2",
                "annotation_stage": "ai_assisted_train",
                "ai_suggested_confidence": "0.90",
            },
            {
                "comment_id": "id3",
                "annotation_stage": "ai_assisted_train",
                "ai_suggested_confidence": "0.99",
            },
        ]
    ).to_csv(gold_csv, index=False, encoding="utf-8-sig")
    return {
        "workbook": workbook,
        "manifest": manifest,
        "review": review,
        "labeled": labeled,
        "eval_json": eval_json,
        "eval_txt": eval_txt,
        "gold_csv": gold_csv,
        "rows": rows,
    }


def test_apply_human_answers_maps_arabic_labels_deterministically(mini_calibration):
    row = load_calibration_workbook(
        mini_calibration["workbook"],
        mini_calibration["manifest"],
        mini_calibration["gold_csv"],
    )[0]
    updated = apply_human_answers(
        row,
        intent_label="شكر / مدح",
        topic_label="موضوع عام",
        sentiment_label="إيجابية",
    )
    assert updated["gold_primary_intent"] == "THANKS_PRAISE"
    assert updated["gold_topic_subtype"] == "GENERAL"
    assert updated["gold_sentiment"] == "POSITIVE"
    assert updated["gold_annotation_status"] == "LABELED"


def test_insufficient_context_clears_gold_fields(mini_calibration):
    row = load_calibration_workbook(
        mini_calibration["workbook"],
        mini_calibration["manifest"],
        mini_calibration["gold_csv"],
    )[2]
    updated = apply_human_answers(row, intent_label="النص غير كافٍ للحكم")
    assert updated["gold_annotation_status"] == "INSUFFICIENT_CONTEXT"
    assert updated["gold_primary_intent"] == ""
    assert updated["gold_topic_subtype"] == ""
    assert updated["gold_sentiment"] == ""


def test_subcategory_required_for_entry_correction(mini_calibration):
    row = load_calibration_workbook(
        mini_calibration["workbook"],
        mini_calibration["manifest"],
        mini_calibration["gold_csv"],
    )[1]
    assert topic_requires_subcategory("ENTRY_CORRECTION") is True
    with pytest.raises(ValueError, match="subcategory required"):
        apply_human_answers(
            row,
            intent_label="شكوى / انتقاد",
            topic_label="تصحيح مدخل / معنى / كتابة",
            sentiment_label="سلبية",
        )
    updated = apply_human_answers(
        row,
        intent_label="شكوى / انتقاد",
        topic_label="مشكلة تقنية",
        sentiment_label="سلبية",
        subcategory_label="مشكلة تقنية",
    )
    assert updated["gold_subcategory"] == "TechFault"


def test_save_and_resume_roundtrip(mini_calibration):
    state = empty_review_state(
        mini_calibration["workbook"],
        mini_calibration["manifest"],
        mini_calibration["gold_csv"],
    )
    state["rows"][0] = apply_human_answers(
        state["rows"][0],
        intent_label="شكر / مدح",
        topic_label="موضوع عام",
        sentiment_label="إيجابية",
    )
    state["current_index"] = 1
    save_review_state(state, mini_calibration["review"])

    resumed = load_or_init_review_state(
        mini_calibration["review"],
        mini_calibration["workbook"],
        mini_calibration["manifest"],
        mini_calibration["gold_csv"],
    )
    assert resumed["current_index"] == 1
    assert resumed["rows"][0]["gold_annotation_status"] == "LABELED"
    assert resumed["rows"][0]["b2_v2_primary_intent"] == "THANKS_PRAISE"


def test_finalize_writes_labeled_workbook_and_evaluation(mini_calibration):
    state = empty_review_state(
        mini_calibration["workbook"],
        mini_calibration["manifest"],
        mini_calibration["gold_csv"],
    )
    state["rows"][0] = apply_human_answers(
        state["rows"][0],
        intent_label="شكر / مدح",
        topic_label="موضوع عام",
        sentiment_label="إيجابية",
    )
    state["rows"][1] = apply_human_answers(
        state["rows"][1],
        intent_label="شكوى / انتقاد",
        topic_label="مشكلة تقنية",
        sentiment_label="سلبية",
        subcategory_label="مشكلة تقنية",
    )
    state["rows"][2] = apply_human_answers(state["rows"][2], intent_label="النص غير كافٍ للحكم")
    assert is_review_complete(state)

    labeled_path, eval_json_path, eval_txt_path, report_text = finalize_calibration_review(
        state,
        review_path=mini_calibration["review"],
        labeled_xlsx_path=mini_calibration["labeled"],
        eval_json_path=mini_calibration["eval_json"],
        eval_txt_path=mini_calibration["eval_txt"],
    )
    assert labeled_path.exists()
    assert eval_json_path.exists()
    assert eval_txt_path.exists()
    assert "Macro-F1" in report_text or "macro_f1" in report_text.lower() or "Macro-F1" in report_text

    evaluation = evaluate_calibration(state)
    assert evaluation.labeled_count == 2
    assert evaluation.insufficient_context_count == 1
    assert evaluation.intent_b2_v2.per_class_f1["THANKS_PRAISE"] == pytest.approx(1.0)
    assert evaluation.intent_b2_v2.per_class_f1["COMPLAINT"] == pytest.approx(1.0)


def test_existing_calibration_workbook_loads():
    assert CALIBRATION_XLSX.exists()
    rows = load_calibration_workbook()
    assert len(rows) == 40
    assert rows[0]["comment_id"]
    assert rows[0]["b2_v2_primary_intent"]
    assert rows[0]["b2_v3_primary_intent"]


def test_review_does_not_modify_source_workbook(tmp_path, mini_calibration):
    workbook = mini_calibration["workbook"]
    before = workbook.read_bytes()
    state = empty_review_state(
        workbook,
        mini_calibration["manifest"],
        mini_calibration["gold_csv"],
    )
    state["rows"][0] = apply_human_answers(
        state["rows"][0],
        intent_label="شكر / مدح",
        topic_label="موضوع عام",
        sentiment_label="إيجابية",
    )
    save_review_state(state, mini_calibration["review"])
    assert workbook.read_bytes() == before
