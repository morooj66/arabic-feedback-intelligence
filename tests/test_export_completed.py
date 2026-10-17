"""Completed-only export (GET /api/excel/export/completed): server-side across all matching persisted decisions, read-only, admin-only. Synthetic data, real HTTP server."""
from __future__ import annotations

import datetime as dt
import hashlib
import io
import json
import urllib.error
import urllib.parse
import urllib.request

import openpyxl
import pytest
from sqlalchemy.orm import Session

from rcs.review.auth import create_user
from rcs.review.http_api import serve_in_thread
from rcs.review.models import ReviewItem, ReviewItemEvent
from rcs.review.repository import get_engine, init_db

T0 = dt.datetime(2026, 8, 1, 9, 0, 0)
DRAFT = "نشكركم على ملاحظتكم القيّمة."
EDITED = "نشكركم على تواصلكم الكريم،\nوسنأخذ ملاحظتكم بالاعتبار — شكرًا."       # diacritics, a line break and a dash must survive exactly
N_COMPLETED = 620                                                                  # more than one 500-row page of the queue query


def call(base, path, cookie=None):
    headers = {"Cookie": cookie} if cookie else {}
    r = urllib.request.Request(base + path, headers=headers)
    try:
        with urllib.request.urlopen(r, timeout=120) as resp:
            return resp.status, resp.read(), resp.headers
    except urllib.error.HTTPError as err:
        return err.code, err.read(), err.headers


def sheets(blob):
    wb = openpyxl.load_workbook(io.BytesIO(blob))
    out = {}
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        out[ws.title] = (list(rows[0]), [list(r) for r in rows[1:]])
    return out


def state(engine):
    with Session(engine) as s:
        items = [(r.comment_id, r.review_status, r.reviewer_action, r.reviewer_user_id, r.row_version, r.final_reply, r.reviewer_edit, r.no_response_reason, r.updated_at)
                 for r in s.query(ReviewItem).order_by(ReviewItem.comment_id).all()]
        events = [(e.event_id, e.review_item_id, e.event_type) for e in s.query(ReviewItemEvent).order_by(ReviewItemEvent.event_id).all()]
    return hashlib.sha256(repr((items, events)).encode()).hexdigest(), len(items), len(events)


class Env:
    pass


@pytest.fixture(scope="module")
def env(tmp_path_factory):
    tmp = tmp_path_factory.mktemp("exp")
    e = get_engine(tmp / "x.sqlite")
    init_db(e)
    users = {k: create_user(e, email=f"{k}@t.local", password="testpassword123", display_name=k, role=r) for k, r in (("admin", "ADMIN"), ("r1", "REVIEWER"), ("r2", "REVIEWER"))}
    intents, topics = ["SUGGESTION", "INQUIRY", "COMPLAINT", "THANKS_PRAISE"], ["WORD_SUBMISSION", "GENERAL", "TECHNICAL_ISSUE", "GENERAL"]
    expected = {}
    with Session(e) as s:
        def add(cid, i, sf, owner, status, **kw):
            k = i % 4
            s.add(ReviewItem(comment_id=cid, source_file=sf, source_row_index=i, original_comment=kw.pop("text", f"تعليق اختباري {cid}"), review_status=status,
                             ai_primary_intent=intents[k], primary_intent=intents[k], ai_topic_subtype=topics[k], topic_subtype=topics[k], sentiment="NEUTRAL",
                             ai_suggested_final_reply=DRAFT, suggested_final_reply=DRAFT, reviewer_user_id=owner, review_priority=(i % 3) + 1,
                             ingested_at=T0 + dt.timedelta(minutes=i), comment_created_at=T0 + dt.timedelta(minutes=i), reviewed_at=T0 + dt.timedelta(days=1, minutes=i) if kw.get("reviewer_action") else None,
                             reviewer_display_name=("r1" if owner == users["r1"].user_id else "r2") if kw.get("reviewer_action") else None, row_version=2 if kw.get("reviewer_action") else 1, **kw))
        for i in range(N_COMPLETED):
            sf, owner = ("one.xlsx", users["r1"].user_id) if i % 2 == 0 else ("two.xlsx", users["r2"].user_id)
            cid = f"c{i:04d}"
            kind = i % 3
            if kind == 0:
                add(cid, i, sf, owner, "APPROVED", reviewer_action="SEND_AS_IS", final_reply=DRAFT)
                expected[cid] = ("SEND_AS_IS", DRAFT, None)
            elif kind == 1:
                add(cid, i, sf, owner, "EDITED_AND_APPROVED", reviewer_action="EDIT_AND_SEND", reviewer_edit=EDITED, final_reply=EDITED)
                expected[cid] = ("EDIT_AND_SEND", EDITED, None)
            else:
                add(cid, i, sf, owner, "NO_REPLY", reviewer_action="NO_RESPONSE_NEEDED", no_response_reason="لا يتطلب ردًا", final_reply=None)
                expected[cid] = ("NO_RESPONSE_NEEDED", None, "لا يتطلب ردًا")
        base_i = 1000
        for j in range(15):                                                         # pending / undecided / in review / system recommendations / excluded
            add(f"p{j:03d}", base_i + j, "one.xlsx", users["r1"].user_id, "DRAFT_READY")
            add(f"q{j:03d}", base_i + j, "two.xlsx", users["r2"].user_id, "IN_REVIEW")
        for j in range(5):
            add(f"n{j:03d}", base_i + j, "one.xlsx", users["r1"].user_id, "NO_RESPONSE_RECOMMENDED")
            add(f"x{j:03d}", base_i + j, "one.xlsx", None, "EXCLUDED")
        for j in range(4):                                                          # exact duplicates of completed comments: never independent reviews
            add(f"dup{j}", base_i + 50 + j, "one.xlsx", None, "DUPLICATE_CONTENT", dedupe_canonical_id=f"c{j * 6:04d}", text=f"تعليق اختباري c{j * 6:04d}")
        add("evil", base_i + 99, "one.xlsx", users["r1"].user_id, "EDITED_AND_APPROVED", reviewer_action="EDIT_AND_SEND", reviewer_edit="=HYPERLINK(\"http://example.invalid\",\"x\")",
            final_reply="=HYPERLINK(\"http://example.invalid\",\"x\")", text="=1+1 تعليق يبدأ بعلامة المعادلة")
        from rcs.review.audit_events import build_event_fields

        for k, cid in enumerate(("c0000", "c0001", "c0002")):                        # real audit history that the export must leave untouched
            s.add(ReviewItemEvent(**build_event_fields(review_item_id=cid, event_type=("SEND_AS_IS", "EDIT_AND_SEND", "NO_RESPONSE_NEEDED")[k], actor_user_id=users["r1"].user_id,
                                                       actor_role="REVIEWER", from_status="IN_REVIEW", to_status="APPROVED", from_assignee_user_id=None, to_assignee_user_id=None, created_at=T0)))
        s.commit()
    expected["evil"] = ("EDIT_AND_SEND", "=HYPERLINK(\"http://example.invalid\",\"x\")", None)
    server, base = serve_in_thread(e)
    env_ = Env()
    env_.engine, env_.base, env_.users, env_.expected = e, base, users, expected
    env_.cookie = {}
    for k in users:
        req = urllib.request.Request(base + "/api/login", data=json.dumps({"email": f"{k}@t.local", "password": "testpassword123"}).encode(), headers={"Content-Type": "application/json"}, method="POST")
        with urllib.request.urlopen(req) as resp:
            env_.cookie[k] = resp.headers.get("Set-Cookie").split(";")[0]
    yield env_
    server.shutdown()


def export(env, qs, who="admin"):
    s, blob, h = call(env.base, "/api/excel/export/completed?" + qs, env.cookie.get(who))
    return s, blob, h


def ids_of(blob):
    sh = sheets(blob)
    hdr2, rows2 = sh["بيانات التدقيق"]
    return [r[hdr2.index("comment_id")] for r in rows2]


# ------------------------------------------------------------------------------------------------------------------------ quantities
@pytest.mark.parametrize("n", [1, 5, 20, 150])
def test_exports_exactly_the_requested_number_of_completed_reviews(env, n):
    s, blob, h = export(env, f"count={n}")
    assert s == 200 and len(ids_of(blob)) == n
    assert (h["X-Export-Rows"], h["X-Export-Requested"], h["X-Export-Shortfall"]) == (str(n), str(n), "0") and int(h["X-Export-Matched"]) == N_COMPLETED + 1
    assert set(ids_of(blob)) <= set(env.expected)                                    # only completed reviews, never pending / excluded / duplicate rows


def test_export_all_spans_more_than_one_page_and_has_no_duplicates_or_gaps(env):
    s, blob, h = export(env, "all=1")
    ids = ids_of(blob)
    assert s == 200 and len(ids) == len(set(ids)) == N_COMPLETED + 1 and set(ids) == set(env.expected)
    assert h["X-Export-Requested"] == "all" and h["X-Export-Rows"] == str(N_COMPLETED + 1)
    assert export(env, "all=1")[1] is not None and ids_of(export(env, "all=1")[1]) == ids        # deterministic order across calls


def test_count_above_available_exports_everything_and_reports_the_shortfall(env):
    s, blob, h = export(env, "count=5000")
    assert s == 200 and len(ids_of(blob)) == N_COMPLETED + 1
    assert h["X-Export-Shortfall"] == str(5000 - (N_COMPLETED + 1))


def test_count_selection_is_deterministic_and_prefix_stable(env):
    a, b = ids_of(export(env, "count=10")[1]), ids_of(export(env, "count=25")[1])
    assert b[:10] == a and len(set(b)) == 25


# ------------------------------------------------------------------------------------------------------------------------ filters across all pages
def test_filters_apply_server_side_across_the_whole_dataset(env):
    s, blob, h = export(env, "all=1&source_file=one.xlsx")
    assert s == 200
    got = ids_of(blob)
    exp = {c for c in env.expected if c[0] == "c" and int(c[1:]) % 2 == 0} | {"evil"}
    assert set(got) == exp and int(h["X-Export-Matched"]) == len(exp)
    s, blob, _ = export(env, f"all=1&reviewer_user_id={env.users['r2'].user_id}")
    assert set(ids_of(blob)) == {c for c in env.expected if c[0] == "c" and int(c[1:]) % 2 == 1}
    s, blob, _ = export(env, "all=1&reviewer_action=NO_RESPONSE_NEEDED")
    assert set(ids_of(blob)) == {c for c, v in env.expected.items() if v[0] == "NO_RESPONSE_NEEDED"} and len(ids_of(blob)) == sum(1 for v in env.expected.values() if v[0] == "NO_RESPONSE_NEEDED")
    s, blob, _ = export(env, "all=1&primary_intent=COMPLAINT")
    hdr, rows = sheets(blob)["الردود المراجعة"]
    assert rows and {r[hdr.index("التصنيف الرئيسي")] for r in rows} == {"COMPLAINT"}
    s, blob, h = export(env, "count=7&source_file=two.xlsx&primary_intent=SUGGESTION,INQUIRY")
    assert s == 200 and len(ids_of(blob)) == 7


def test_review_status_can_only_narrow_never_widen_the_completed_set(env):
    s, blob, _ = export(env, "all=1&review_status=DRAFT_READY,IN_REVIEW")                  # pending statuses asked for -> nothing completed matches
    assert s == 409 and json.loads(blob)["error"] == "no_results_to_export"
    s, blob, _ = export(env, "all=1&review_status=NO_REPLY")
    assert s == 200 and {env.expected[c][0] for c in ids_of(blob)} == {"NO_RESPONSE_NEEDED"}


# ------------------------------------------------------------------------------------------------------------------------ what is (not) in the file
def test_mixed_pending_and_completed_only_completed_rows_appear(env):
    _, blob, _ = export(env, "all=1")
    ids = set(ids_of(blob))
    for prefix in ("p", "q", "n", "x"):
        assert not any(i.startswith(prefix) and i[1:].isdigit() for i in ids)
    assert not any(i.startswith("dup") for i in ids)                                          # duplicate rows are not independent completed reviews
    hdr, rows = sheets(blob)["الردود المراجعة"]
    statuses = {r[hdr.index("حالة المراجعة")] for r in rows}
    assert statuses <= {"APPROVED", "EDITED_AND_APPROVED", "NO_REPLY"} and all(r[hdr.index("إجراء المراجع")] for r in rows)


def test_decisions_and_final_replies_are_exact(env):
    _, blob, _ = export(env, "all=1")
    sh = sheets(blob)
    h1, r1 = sh["الردود المراجعة"]
    h2, r2 = sh["بيانات التدقيق"]
    assert len(r1) == len(r2)
    seen = {"SEND_AS_IS": 0, "EDIT_AND_SEND": 0, "NO_RESPONSE_NEEDED": 0}
    for a, b in zip(r1, r2):
        cid = b[h2.index("comment_id")]
        action, final, reason = env.expected[cid]
        assert a[h1.index("إجراء المراجع")] == action
        seen[action] += 1
        if action == "SEND_AS_IS":
            assert a[h1.index("الرد النهائي")] == DRAFT == a[h1.index("الرد المقترح")] and a[h1.index("الرد بعد التعديل")] is None
        elif action == "EDIT_AND_SEND":
            assert a[h1.index("الرد النهائي")] == final == a[h1.index("الرد بعد التعديل")]       # exact approved text, byte for byte
            assert a[h1.index("الرد المقترح")] == DRAFT                                       # the original draft stays visible
        else:
            assert a[h1.index("الرد النهائي")] is None and a[h1.index("الرد بعد التعديل")] is None   # no fabricated response
            assert b[h2.index("no_response_reason")] == reason
        assert a[h1.index("تعليق الزائر")] and a[h1.index("اسم المراجع")] and a[h1.index("تاريخ المراجعة")]
    assert all(v > 100 for v in seen.values())


def test_text_starting_with_an_equals_sign_stays_text_not_a_formula(env):
    _, blob, _ = export(env, "all=1")
    wb = openpyxl.load_workbook(io.BytesIO(blob))
    ws = wb["الردود المراجعة"]
    hdr = [c.value for c in ws[1]]
    row = next(r for r in ws.iter_rows(min_row=2) if r[hdr.index("الرد النهائي")].value and str(r[hdr.index("الرد النهائي")].value).startswith("="))
    assert row[hdr.index("الرد النهائي")].data_type == "s" and row[hdr.index("تعليق الزائر")].data_type == "s"
    assert row[hdr.index("تعليق الزائر")].value.startswith("=1+1")
    assert not [c for r in wb.worksheets for rr in r.iter_rows() for c in rr if c.data_type == "f"]          # no formula cell anywhere in the workbook


# ------------------------------------------------------------------------------------------------------------------------ read-only, authorization, validation
def test_export_is_read_only_status_assignments_versions_and_audit_history_unchanged(env):
    before = state(env.engine)
    for qs in ("all=1", "count=5", "all=1&source_file=one.xlsx", "count=3&reviewer_action=SEND_AS_IS"):
        assert export(env, qs)[0] == 200
    assert export(env, "all=1", who=None)[0] == 401
    assert before[2] == 3
    assert state(env.engine) == before                                                         # statuses, assignees, row versions, texts, events byte-identical


@pytest.mark.parametrize("who,code", [(None, 401), ("r1", 403), ("r2", 403)])
def test_only_admin_may_export(env, who, code):
    s, blob, _ = export(env, "all=1", who=who)
    assert s == code and b"PK" != blob[:2]


@pytest.mark.parametrize("qs", ["", "count=0", "count=-3", "count=abc", "count=2.5", "count=3&all=1", "all=0", "count=3&filter=pending", "count=3&limit=5", "count=3&offset=2",
                                "count=3&primary_intent=NOPE", "count=3&reviewer_user_id=nobody", "count=3&sort=bogus"])
def test_invalid_requests_are_rejected_without_a_file(env, qs):
    s, blob, _ = export(env, qs)
    assert s == 400 and blob[:2] != b"PK"


def test_empty_result_is_a_clear_409_not_an_empty_file(env):
    s, blob, _ = export(env, "all=1&search=" + urllib.parse.quote("لا-يوجد-تعليق-بهذا-النص"))
    assert s == 409 and json.loads(blob)["error"] == "no_results_to_export"
    s, blob, _ = export(env, "count=3&source_file=missing.xlsx")
    assert s == 409


def test_the_existing_full_batch_export_is_unchanged_and_still_includes_pending_rows(env):
    s, blob, _ = call(env.base, "/api/excel/export?source_file=one.xlsx", env.cookie["admin"])
    assert s == 200
    hdr, rows = sheets(blob)["الردود المراجعة"]
    assert any(not r[hdr.index("إجراء المراجع")] for r in rows) and any(r[hdr.index("إجراء المراجع")] for r in rows)
    assert call(env.base, "/api/excel/export?source_file=one.xlsx", env.cookie["r1"])[0] == 403
