"""Regression tests for calibration wizard navigation (no Gradio runtime)."""

from __future__ import annotations

import copy
import json
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook

from rcs.evaluation.calibration_labeling import (
    REVIEW_JSON,
    apply_human_answers,
    current_comment_id,
    current_row,
    empty_review_state,
    ensure_resume_index,
    first_incomplete_index,
    go_previous_row,
    load_or_init_review_state,
    row_to_form_labels,
    save_and_advance_row,
    save_current_row,
    save_review_state,
)


def _write_workbook(path: Path, rows: list[dict[str, str]]) -> None:
    columns = [
        "comment_text",
        "b2_v2_primary_intent",
        "b2_v2_topic_subtype",
        "b2_v2_sentiment",
        "b2_v3_primary_intent",
        "b2_v3_topic_subtype",
        "b2_v3_sentiment",
        "gold_primary_intent",
        "gold_topic_subtype",
        "gold_sentiment",
        "gold_annotation_status",
        "comment_id",
        "calibration_stratum",
    ]
    wb = Workbook()
    ws = wb.active
    ws.title = "Review"
    for col_idx, name in enumerate(columns, start=1):
        ws.cell(row=1, column=col_idx, value=name)
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, name in enumerate(columns, start=1):
            ws.cell(row=row_idx, column=col_idx, value=row.get(name, ""))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _base_row(comment_id: str, text: str) -> dict[str, str]:
    return {
        "comment_id": comment_id,
        "comment_text": text,
        "calibration_stratum": "test",
        "b2_v2_primary_intent": "SUGGESTION",
        "b2_v2_topic_subtype": "GENERAL",
        "b2_v2_sentiment": "NEUTRAL",
        "b2_v3_primary_intent": "SUGGESTION",
        "b2_v3_topic_subtype": "GENERAL",
        "b2_v3_sentiment": "NEUTRAL",
    }


def _forty_rows() -> list[dict[str, str]]:
    return [_base_row(f"id{i:02d}", f"comment {i}") for i in range(40)]


def _labeled_kwargs() -> dict[str, str]:
    return {
        "intent_label": "شكر / مدح",
        "topic_label": "موضوع عام",
        "sentiment_label": "إيجابية",
    }


@pytest.fixture()
def nav_fixture(tmp_path: Path):
    rows = _forty_rows()
    workbook = tmp_path / "sample40.xlsx"
    manifest = tmp_path / "manifest.json"
    review = tmp_path / "review.json"
    gold_csv = tmp_path / "gold.csv"
    _write_workbook(workbook, rows)
    manifest.write_text(
        json.dumps({"comment_ids_in_workbook_order": [r["comment_id"] for r in rows]}),
        encoding="utf-8",
    )
    pd.DataFrame(
        [{"comment_id": r["comment_id"], "annotation_stage": "ai_assisted_train", "ai_suggested_confidence": "0.9"} for r in rows]
    ).to_csv(gold_csv, index=False, encoding="utf-8-sig")
    return {"workbook": workbook, "manifest": manifest, "review": review, "gold_csv": gold_csv}


def test_save_and_next_advances_exactly_once(nav_fixture):
    state = empty_review_state(nav_fixture["workbook"], nav_fixture["manifest"], nav_fixture["gold_csv"])
    assert state["current_index"] == 0
    state, error = save_and_advance_row(state, review_path=nav_fixture["review"], **_labeled_kwargs())
    assert error is None
    assert state["current_index"] == 1
    assert current_comment_id(state) == "id01"


def test_saved_comment_id_matches_displayed_comment(nav_fixture):
    state = empty_review_state(nav_fixture["workbook"], nav_fixture["manifest"], nav_fixture["gold_csv"])
    displayed_id = current_comment_id(state)
    state, error = save_and_advance_row(state, review_path=nav_fixture["review"], **_labeled_kwargs())
    assert error is None
    assert state["rows"][0]["comment_id"] == displayed_id
    assert state["rows"][0]["gold_annotation_status"] == "LABELED"
    assert current_comment_id(state) != displayed_id


def test_resume_opens_first_incomplete(nav_fixture):
    state = empty_review_state(nav_fixture["workbook"], nav_fixture["manifest"], nav_fixture["gold_csv"])
    state["rows"][0] = apply_human_answers(state["rows"][0], intent_label="النص غير كافٍ للحكم")
    state["current_index"] = 7
    save_review_state(state, nav_fixture["review"])

    resumed = load_or_init_review_state(
        nav_fixture["review"],
        nav_fixture["workbook"],
        nav_fixture["manifest"],
        nav_fixture["gold_csv"],
    )
    assert first_incomplete_index(resumed["rows"]) == 1
    assert resumed["current_index"] == 1
    assert current_comment_id(resumed) == "id01"


def test_previous_restores_saved_answers(nav_fixture):
    state = empty_review_state(nav_fixture["workbook"], nav_fixture["manifest"], nav_fixture["gold_csv"])
    state, _ = save_and_advance_row(state, review_path=nav_fixture["review"], **_labeled_kwargs())
    state, _ = save_and_advance_row(state, review_path=nav_fixture["review"], **_labeled_kwargs())
    state = go_previous_row(state, nav_fixture["review"])
    labels = row_to_form_labels(current_row(state))
    assert labels["intent_label"] == "شكر / مدح"
    assert labels["topic_label"] == "موضوع عام"
    assert labels["sentiment_label"] == "إيجابية"


def test_incomplete_form_does_not_advance(nav_fixture):
    state = empty_review_state(nav_fixture["workbook"], nav_fixture["manifest"], nav_fixture["gold_csv"])
    before_index = state["current_index"]
    state, error = save_and_advance_row(
        state,
        review_path=nav_fixture["review"],
        intent_label="اقتراح",
        topic_label="",
        sentiment_label="",
    )
    assert error == "اختر موضوع التعليق."
    assert state["current_index"] == before_index
    assert state["rows"][0]["gold_annotation_status"] == ""


def test_insufficient_context_advances_without_topic_or_sentiment(nav_fixture):
    state = empty_review_state(nav_fixture["workbook"], nav_fixture["manifest"], nav_fixture["gold_csv"])
    state, error = save_and_advance_row(
        state,
        review_path=nav_fixture["review"],
        intent_label="النص غير كافٍ للحكم",
    )
    assert error is None
    assert state["current_index"] == 1
    assert state["rows"][0]["gold_annotation_status"] == "INSUFFICIENT_CONTEXT"
    assert state["rows"][0]["gold_primary_intent"] == ""


def test_forty_sequential_saves_visit_unique_ids_without_repeat_or_skip(nav_fixture):
    state = empty_review_state(nav_fixture["workbook"], nav_fixture["manifest"], nav_fixture["gold_csv"])
    visited: list[str] = []
    for _ in range(40):
        visited.append(current_comment_id(state))
        state, error = save_and_advance_row(state, review_path=nav_fixture["review"], **_labeled_kwargs())
        assert error is None
    assert len(visited) == 40
    assert len(set(visited)) == 40
    assert visited == [f"id{i:02d}" for i in range(40)]
    assert all(row["gold_annotation_status"] == "LABELED" for row in state["rows"])


def test_repeated_save_and_next_does_not_overwrite_prior_row(nav_fixture):
    """Reproduces the stale-index failure mode: without advancing current_index,
    a second save would rewrite row 0 instead of row 1."""
    state = empty_review_state(nav_fixture["workbook"], nav_fixture["manifest"], nav_fixture["gold_csv"])
    state, _ = save_and_advance_row(state, review_path=nav_fixture["review"], **_labeled_kwargs())
    first_saved_at = state["rows"][0]["labeled_at"]

    stale_state = copy.deepcopy(state)
    stale_state["current_index"] = 0
    stale_state, _ = save_and_advance_row(
        stale_state,
        review_path=nav_fixture["review"],
        intent_label="سبام / خارج الموضوع",
        topic_label="موضوع عام",
        sentiment_label="محايدة",
    )
    assert stale_state["rows"][0]["gold_primary_intent"] == "SPAM"
    assert stale_state["current_index"] == 1

    state, _ = save_and_advance_row(state, review_path=nav_fixture["review"], **_labeled_kwargs())
    assert state["rows"][0]["labeled_at"] == first_saved_at
    assert state["rows"][0]["gold_primary_intent"] == "THANKS_PRAISE"
    assert state["current_index"] == 2


def test_existing_saved_human_review_is_preserved():
    if not REVIEW_JSON.exists():
        pytest.skip("no persisted human review artifact yet")
    before = REVIEW_JSON.read_text(encoding="utf-8")
    payload = json.loads(before)
    first_row = payload["rows"][0]
    assert first_row["gold_annotation_status"] == "INSUFFICIENT_CONTEXT"
    assert first_row["comment_id"] == "01dc4643e00059568afcd22974c74ad3cf1749af4f021da57422fc9b8fea120d"

    state = load_or_init_review_state()
    assert state["rows"][0]["gold_annotation_status"] == "INSUFFICIENT_CONTEXT"
    assert state["rows"][0]["labeled_at"] == first_row["labeled_at"]
    assert ensure_resume_index(state)["current_index"] == first_incomplete_index(state["rows"])

    after = REVIEW_JSON.read_text(encoding="utf-8")
    assert after == before
