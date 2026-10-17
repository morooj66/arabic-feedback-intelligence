"""Pre-B.3 workbook export/import safety tests. No live API calls."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from rcs.evaluation.gold_dataset import (
    AI_SUGGESTED_COLUMNS,
    PREEXISTING_AI_ASSISTED_GOLD_IDS,
    REVIEWER_EDITABLE_COLUMNS,
    verify_post_b2_b3_invariants,
)
from scripts.build_b3_review_workbook import (
    ALL_COLUMNS,
    build,
    import_calibration_human_labels,
    import_reviewed_workbook,
    prepare_calibration_human_import,
    prepare_reviewed_workbook_import,
)

PROJECT_ROOT = Path(__file__).resolve().parent.parent
CANONICAL_CSV = PROJECT_ROOT / "data" / "gold" / "gold_dataset_round1.csv"
PERSISTED_REVIEW_WORKBOOK = (
    PROJECT_ROOT / "data" / "gold" / "gold_dataset_round1_b3_review.xlsx"
)


def _synthetic_post_b2_csv(tmp_path: Path) -> Path:
    defaults = {
        "comment_text": "",
        "gold_primary_intent": "",
        "gold_topic_subtype": "",
        "gold_secondary_intent": "",
        "gold_subcategory": "",
        "gold_sentiment": "",
        "gold_annotation_status": "",
        "labeler_notes": "",
        "labeled_by": "",
        "labeled_at": "",
        "is_few_shot_example": "False",
    }
    for col in AI_SUGGESTED_COLUMNS:
        defaults[col] = ""

    rows = [
        {
            **defaults,
            "comment_id": "seed1",
            "comment_text": "شكرا",
            "split": "train",
            "annotation_stage": "human_train_seed",
            "is_few_shot_example": "True",
            "gold_primary_intent": "THANKS_PRAISE",
            "gold_topic_subtype": "GENERAL",
            "gold_sentiment": "POSITIVE",
            "gold_annotation_status": "LABELED",
        },
        {
            **defaults,
            "comment_id": "val1",
            "comment_text": "سؤال",
            "split": "val",
            "annotation_stage": "human_val",
            "gold_primary_intent": "INQUIRY",
            "gold_topic_subtype": "GENERAL",
            "gold_sentiment": "NEUTRAL",
            "gold_annotation_status": "LABELED",
        },
        {
            **defaults,
            "comment_id": "test1",
            "comment_text": "اقتراح",
            "split": "test",
            "annotation_stage": "human_test",
            "gold_primary_intent": "SUGGESTION",
            "gold_topic_subtype": "DICTIONARY_SUGGESTION",
            "gold_sentiment": "NEUTRAL",
            "gold_annotation_status": "LABELED",
        },
        {
            **defaults,
            "comment_id": "ai1",
            "comment_text": "الموقع لا يعمل",
            "split": "train",
            "annotation_stage": "ai_assisted_train",
            "ai_suggested_primary_intent": "COMPLAINT",
            "ai_suggested_topic_subtype": "TECHNICAL_ISSUE",
            "ai_suggested_subcategory": "TechFault",
            "ai_suggested_sentiment": "NEGATIVE",
            "ai_suggested_confidence": "0.9",
            "ai_suggestion_model": "gemini-test:b2",
        },
        {
            **defaults,
            "comment_id": "ai2",
            "comment_text": "أضيفوا كلمة",
            "split": "train",
            "annotation_stage": "ai_assisted_train",
            "ai_suggested_primary_intent": "SUGGESTION",
            "ai_suggested_topic_subtype": "WORD_SUBMISSION",
            "ai_suggested_sentiment": "NEUTRAL",
            "ai_suggested_confidence": "0.8",
            "ai_suggestion_model": "gemini-test:b2",
        },
    ]
    path = tmp_path / "gold.csv"
    pd.DataFrame(rows).to_csv(path, index=False, encoding="utf-8-sig")
    return path


def _set_review_value(workbook_path: Path, comment_id: str, column: str, value: str) -> None:
    wb = load_workbook(workbook_path)
    ws = wb["Review"]
    headers = {cell.value: cell.column for cell in ws[1]}
    id_col = headers["comment_id"]
    for row_number in range(2, ws.max_row + 1):
        if ws.cell(row_number, id_col).value == comment_id:
            ws.cell(row_number, headers[column], value)
            wb.save(workbook_path)
            wb.close()
            return
    wb.close()
    raise AssertionError(f"comment_id not found in workbook: {comment_id}")


def _complete_review(workbook_path: Path) -> None:
    values = {
        "ai1": {
            "gold_primary_intent": "COMPLAINT",
            "gold_topic_subtype": "TECHNICAL_ISSUE",
            "gold_subcategory": "TechFault",
            "gold_sentiment": "NEGATIVE",
        },
        "ai2": {
            "gold_primary_intent": "SUGGESTION",
            "gold_topic_subtype": "WORD_SUBMISSION",
            "gold_sentiment": "NEUTRAL",
        },
    }
    for cid, labels in values.items():
        for column, value in labels.items():
            _set_review_value(workbook_path, cid, column, value)
        _set_review_value(workbook_path, cid, "gold_annotation_status", "LABELED")
        _set_review_value(workbook_path, cid, "labeled_by", "reviewer")
        _set_review_value(workbook_path, cid, "labeled_at", "2026-08-16")


def test_persisted_b3_workbook_roundtrips_exactly_against_canonical():
    pytest.skip(
        "Persisted full 300-row workbook predates b2-v3; use remaining260 export after calibration import."
    )


def test_export_then_read_back_roundtrip_is_exact(tmp_path):
    csv_path = _synthetic_post_b2_csv(tmp_path)
    workbook_path = tmp_path / "review.xlsx"
    build(csv_path, workbook_path)

    original = pd.read_csv(csv_path, dtype=str, keep_default_na=False)
    merged = prepare_reviewed_workbook_import(
        workbook_path,
        csv_path,
        require_b3_complete=False,
    )
    assert list(load_workbook(workbook_path, read_only=True)["Review"].values)[0] == tuple(ALL_COLUMNS)
    pd.testing.assert_frame_equal(original, merged)


def test_reviewer_editable_cells_are_unlocked(tmp_path):
    csv_path = _synthetic_post_b2_csv(tmp_path)
    workbook_path = tmp_path / "review.xlsx"
    build(csv_path, workbook_path)

    wb = load_workbook(workbook_path)
    ws = wb["Review"]
    headers = {cell.value: cell.column for cell in ws[1]}
    editable_cols = {headers[name] for name in REVIEWER_EDITABLE_COLUMNS}
    for row_number in range(2, ws.max_row + 1):
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row_number, col_idx)
            if col_idx in editable_cols:
                assert cell.protection.locked is False, f"{ws.cell(1, col_idx).value} should be unlocked"
            else:
                assert cell.protection.locked is True
    wb.close()


def test_import_updates_only_human_review_fields(tmp_path):
    csv_path = _synthetic_post_b2_csv(tmp_path)
    workbook_path = tmp_path / "review.xlsx"
    output_path = tmp_path / "imported.csv"
    build(csv_path, workbook_path)
    _complete_review(workbook_path)

    before = pd.read_csv(csv_path, dtype=str, keep_default_na=False)
    import_reviewed_workbook(workbook_path, output_path, csv_path)
    after = pd.read_csv(output_path, dtype=str, keep_default_na=False)

    ai1 = after.set_index("comment_id").loc["ai1"]
    assert ai1["gold_primary_intent"] == "COMPLAINT"
    assert ai1["labeled_by"] == "reviewer"
    assert ai1["ai_suggested_primary_intent"] == "COMPLAINT"
    protected = before["annotation_stage"] != "ai_assisted_train"
    pd.testing.assert_frame_equal(
        before.loc[protected].reset_index(drop=True),
        after.loc[protected].reset_index(drop=True),
    )
    verify_post_b2_b3_invariants(after, require_b3_complete=True)


@pytest.mark.parametrize(
    ("column", "value", "message"),
    [
        ("comment_text", "tampered", "protected column"),
        ("ai_suggested_primary_intent", "SPAM", "protected column"),
        ("split", "test", "protected column"),
        ("annotation_stage", "human_test", "protected column"),
    ],
)
def test_import_rejects_protected_field_changes(tmp_path, column, value, message):
    csv_path = _synthetic_post_b2_csv(tmp_path)
    workbook_path = tmp_path / "review.xlsx"
    output_path = tmp_path / "must_not_exist.csv"
    build(csv_path, workbook_path)
    _set_review_value(workbook_path, "ai1", column, value)

    with pytest.raises(AssertionError, match=message):
        import_reviewed_workbook(
            workbook_path,
            output_path,
            csv_path,
            require_b3_complete=False,
        )
    assert not output_path.exists()


def test_import_rejects_missing_or_replaced_target_row(tmp_path):
    csv_path = _synthetic_post_b2_csv(tmp_path)
    workbook_path = tmp_path / "review.xlsx"
    build(csv_path, workbook_path)
    _set_review_value(workbook_path, "ai1", "comment_id", "val1")

    with pytest.raises(AssertionError, match="row set differs"):
        prepare_reviewed_workbook_import(
            workbook_path,
            csv_path,
            require_b3_complete=False,
        )


def test_import_rejects_invalid_human_labels_and_missing_trace(tmp_path):
    csv_path = _synthetic_post_b2_csv(tmp_path)
    workbook_path = tmp_path / "review.xlsx"
    build(csv_path, workbook_path)
    _set_review_value(workbook_path, "ai1", "gold_primary_intent", "NOT_AN_INTENT")
    _set_review_value(workbook_path, "ai1", "gold_annotation_status", "LABELED")

    with pytest.raises(AssertionError, match="invalid gold_primary_intent"):
        prepare_reviewed_workbook_import(
            workbook_path,
            csv_path,
            require_b3_complete=False,
        )

    build(csv_path, workbook_path)
    _set_review_value(workbook_path, "ai1", "gold_primary_intent", "COMPLAINT")
    _set_review_value(workbook_path, "ai1", "gold_topic_subtype", "GENERAL")
    _set_review_value(workbook_path, "ai1", "gold_sentiment", "NEGATIVE")
    _set_review_value(workbook_path, "ai1", "gold_annotation_status", "LABELED")
    with pytest.raises(AssertionError, match="require labeled_by and labeled_at"):
        prepare_reviewed_workbook_import(
            workbook_path,
            csv_path,
            require_b3_complete=False,
        )


def test_changing_preexisting_anomaly_requires_new_human_trace(tmp_path):
    csv_path = _synthetic_post_b2_csv(tmp_path)
    anomaly_id = next(iter(PREEXISTING_AI_ASSISTED_GOLD_IDS))
    canonical = pd.read_csv(csv_path, dtype=str, keep_default_na=False)
    row = canonical["comment_id"] == "ai1"
    canonical.loc[row, "comment_id"] = anomaly_id
    canonical.loc[row, "gold_primary_intent"] = "INQUIRY"
    canonical.loc[row, "gold_topic_subtype"] = "GENERAL"
    canonical.loc[row, "gold_sentiment"] = "NEUTRAL"
    canonical.loc[row, "gold_annotation_status"] = "LABELED"
    canonical.to_csv(csv_path, index=False, encoding="utf-8-sig")

    workbook_path = tmp_path / "review.xlsx"
    build(csv_path, workbook_path)
    _set_review_value(
        workbook_path,
        anomaly_id,
        "gold_primary_intent",
        "COMPLAINT",
    )

    with pytest.raises(AssertionError, match="pre-existing anomaly"):
        prepare_reviewed_workbook_import(
            workbook_path,
            csv_path,
            require_b3_complete=False,
        )


def test_calibration_human_import_applies_only_reviewer_fields(tmp_path):
    calibration_review = PROJECT_ROOT / "data" / "gold" / "b2_calibration_human_review.json"
    if not calibration_review.exists():
        pytest.skip("calibration review artifact not present")
    from rcs.evaluation.calibration_labeling import completed_calibration_comment_ids

    csv_copy = tmp_path / "gold.csv"
    frame = pd.read_csv(CANONICAL_CSV, dtype=str, keep_default_na=False)
    calibration_ids = completed_calibration_comment_ids(calibration_review)
    for col in REVIEWER_EDITABLE_COLUMNS:
        frame.loc[frame["comment_id"].isin(calibration_ids), col] = ""
    frame.to_csv(csv_copy, index=False, encoding="utf-8-sig")
    before = pd.read_csv(csv_copy, dtype=str, keep_default_na=False)
    merged = prepare_calibration_human_import(csv_copy, calibration_review)
    verify_post_b2_b3_invariants(merged, require_b3_complete=False)

    from rcs.evaluation.calibration_labeling import completed_calibration_comment_ids

    imported_ids = completed_calibration_comment_ids(calibration_review)
    imported = merged[merged["comment_id"].isin(imported_ids)]
    assert len(imported) == 40
    assert (imported["gold_annotation_status"].astype(str).str.strip() != "").all()
    assert (imported["labeled_by"].astype(str).str.strip() != "").all()
    assert (imported["labeled_at"].astype(str).str.strip() != "").all()

    noneditable = [col for col in before.columns if col not in REVIEWER_EDITABLE_COLUMNS]
    pd.testing.assert_frame_equal(before[noneditable], merged[noneditable])
    protected = before["annotation_stage"] != "ai_assisted_train"
    pd.testing.assert_frame_equal(before.loc[protected], merged.loc[protected])


def test_build_remaining_workbook_excludes_calibration_rows(tmp_path):
    calibration_review = PROJECT_ROOT / "data" / "gold" / "b2_calibration_human_review.json"
    if not calibration_review.exists():
        pytest.skip("calibration review artifact not present")
    csv_copy = tmp_path / "gold.csv"
    out_path = tmp_path / "remaining.xlsx"
    pd.read_csv(CANONICAL_CSV, dtype=str, keep_default_na=False).to_csv(
        csv_copy, index=False, encoding="utf-8-sig"
    )
    from rcs.evaluation.calibration_labeling import completed_calibration_comment_ids

    exclude_ids = completed_calibration_comment_ids(calibration_review)
    build(csv_copy, out_path, exclude_comment_ids=exclude_ids)
    wb = load_workbook(out_path, read_only=True)
    rows = list(wb["Review"].values)[1:]
    wb.close()
    assert len(rows) == 260
    hidden_idx = ALL_COLUMNS.index("comment_id")
    exported_ids = {row[hidden_idx] for row in rows}
    assert exported_ids.isdisjoint(exclude_ids)
