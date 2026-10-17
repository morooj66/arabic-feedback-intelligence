"""
Safe local/demo Excel upload handling for the reviewer application.

Scope: accept an .xlsx file the reviewer app can then hand to the existing
Excel E2E flow (rcs.review.excel_review_flow). This module only deals with
receiving/validating/storing bytes safely — it does not ingest, classify,
or generate anything.

Safety properties:
  - Only .xlsx is accepted.
  - The client-supplied filename is NEVER used to build a filesystem path
    directly. Each upload is written under its own fresh uuid4 directory
    inside EXCEL_UPLOAD_DIR; only the basename (stripped of any directory
    components, and re-checked for traversal) is kept, purely so the
    original name survives for display and for comment_id stability (see
    rcs.ingest.loader._comment_id, which is keyed on the file's name) —
    re-uploading the exact same file/name is what makes duplicate-import
    idempotency work.
  - resolve_uploaded_workbook() re-validates that a token-derived path
    stays inside EXCEL_UPLOAD_DIR before ever touching the filesystem.
  - Nothing here writes to data/raw, data/gold, data/models, or any other
    frozen/gold artifact location.
"""

from __future__ import annotations

import re
import uuid
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from config import settings

_SAFE_NAME_RE = re.compile(r"[^A-Za-z0-9؀-ۿ._-]+")


class UploadRejected(Exception):
    """Client-facing reason code; message is safe to show as-is."""

    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code


def sanitize_upload_filename(name: str | None) -> str:
    """Strip any path component and unsafe characters; force a .xlsx name.

    Never trusts the client filename as a filesystem path — only the
    basename survives, and only a conservative character set.
    """
    base = Path((name or "").strip() or "upload.xlsx").name  # drop any dir components
    base = base.replace("\\", "_")
    stem = Path(base).stem or "upload"
    stem = _SAFE_NAME_RE.sub("_", stem).strip("._") or "upload"
    return f"{stem[:120]}.xlsx"


@dataclass(frozen=True)
class UploadedWorkbook:
    token: str
    path: Path
    original_filename: str
    size_bytes: int


def save_uploaded_workbook(filename: str | None, data: bytes) -> UploadedWorkbook:
    """Write raw upload bytes to a fresh, isolated directory. Never
    overwrites anything; never touches data/raw or any other source path."""
    if not data:
        raise UploadRejected("empty_upload", "الملف المرفوع فارغ")
    if len(data) > settings.EXCEL_UPLOAD_MAX_BYTES:
        raise UploadRejected(
            "file_too_large",
            f"حجم الملف يتجاوز الحد المسموح ({settings.EXCEL_UPLOAD_MAX_BYTES // (1024 * 1024)} ميجابايت)",
        )
    original_suffix = Path((filename or "").strip()).suffix.lower()
    if original_suffix != ".xlsx":
        raise UploadRejected("unsupported_file_type", "الملفات المدعومة حاليًا: .xlsx فقط")
    safe_name = sanitize_upload_filename(filename)

    token = uuid.uuid4().hex
    upload_dir = settings.EXCEL_UPLOAD_DIR / token
    upload_dir.mkdir(parents=True, exist_ok=False)
    dest = upload_dir / safe_name
    dest.write_bytes(data)  # preserves the exact uploaded bytes, unmodified, for audit
    return UploadedWorkbook(token=token, path=dest, original_filename=safe_name, size_bytes=len(data))


def resolve_uploaded_workbook(token: str) -> Path:
    """Resolve an opaque upload token back to its stored file.

    Defense in depth: even though the token is server-generated, re-verify
    the resolved path is actually contained inside EXCEL_UPLOAD_DIR before
    returning it (rejects any '..'/absolute-path token trickery).
    """
    if not re.fullmatch(r"[0-9a-f]{32}", token or ""):
        raise UploadRejected("invalid_upload_token", "رمز الملف المرفوع غير صالح")
    upload_dir = (settings.EXCEL_UPLOAD_DIR / token).resolve()
    root = settings.EXCEL_UPLOAD_DIR.resolve()
    if root not in upload_dir.parents and upload_dir != root:
        raise UploadRejected("invalid_upload_token", "رمز الملف المرفوع غير صالح")
    if not upload_dir.is_dir():
        raise UploadRejected("upload_not_found", "لم يتم العثور على الملف المرفوع")
    candidates = sorted(p for p in upload_dir.iterdir() if p.suffix.lower() == ".xlsx")
    if not candidates:
        raise UploadRejected("upload_not_found", "لم يتم العثور على الملف المرفوع")
    return candidates[0]


@dataclass(frozen=True)
class WorkbookValidation:
    ok: bool
    total_rows: int
    errors: list[str] = field(default_factory=list)


def validate_workbook(path: Path) -> WorkbookValidation:
    """Inspect an uploaded workbook before ingest: readable, has the
    expected header row, and at least one data row. Never raises for a
    malformed file — always returns a WorkbookValidation with Arabic error
    messages the UI can show directly.
    """
    errors: list[str] = []
    if not path.exists():
        return WorkbookValidation(False, 0, ["لم يتم العثور على الملف المرفوع"])

    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return WorkbookValidation(False, 0, ["تعذّرت قراءة الملف — تأكد أنه ملف Excel (.xlsx) سليم"])

    try:
        sheet = wb.active
        rows_iter = sheet.iter_rows(values_only=True)
        try:
            header = next(rows_iter)
        except StopIteration:
            return WorkbookValidation(False, 0, ["الملف فارغ ولا يحتوي على صف عناوين"])

        # Exact match against settings.SOURCE_COLUMNS keys, no whitespace
        # normalization — one header is documented to include a leading
        # space (" البريد الالكتروني للزائر"), matching rcs.ingest.loader.
        header_set = {str(h) for h in header if h is not None}
        missing = [col for col in settings.SOURCE_COLUMNS if col not in header_set]
        if missing:
            errors.append(
                "الملف لا يحتوي على الأعمدة المتوقعة: " + "، ".join(missing)
            )

        total_rows = 0
        for row in rows_iter:
            if any(cell is not None and str(cell).strip() for cell in row):
                total_rows += 1

        if total_rows == 0 and not errors:
            errors.append("لا توجد بيانات صالحة في الملف (صف العناوين فقط)")

        return WorkbookValidation(ok=not errors, total_rows=total_rows, errors=errors)
    finally:
        wb.close()
