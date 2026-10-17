"""
Phase B.2 calibration — human labeling persistence and evaluation.

Reads the frozen 40-row calibration workbook (b2-v2/b2-v3 predictions are
never modified). Saves human gold labels to a separate review artifact only.
On completion, writes a labeled export and an evaluation report comparing
both Gemini prompt versions against human gold via rcs.evaluation.metrics.
"""

from __future__ import annotations

import datetime as dt
import json
import os
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from config import settings
from rcs.evaluation.metrics import EvaluationReport, evaluate

CALIBRATION_XLSX = settings.DATA_DIR / "gold" / "b2_calibration_sample40.xlsx"
CALIBRATION_MANIFEST = settings.DATA_DIR / "gold" / "b2_calibration_sample40_manifest.json"
REVIEW_JSON = settings.DATA_DIR / "gold" / "b2_calibration_human_review.json"
LABELED_XLSX = settings.DATA_DIR / "gold" / "b2_calibration_sample40_labeled.xlsx"
EVAL_REPORT_JSON = settings.DATA_DIR / "gold" / "b2_calibration_evaluation_report.json"
EVAL_REPORT_TXT = settings.DATA_DIR / "gold" / "b2_calibration_evaluation_report.txt"
GOLD_CSV = settings.DATA_DIR / "gold" / "gold_dataset_round1.csv"

HIGH_CONFIDENCE_THRESHOLD = 0.85
CALIBRATION_LABELER_ID = "b2_calibration_wizard"
CALIBRATION_IMPORT_NOTES = "Imported from completed B.2 calibration sample (40-row wizard)."

# Arabic display label -> taxonomy code (deterministic, no free-form parsing).
# Note: "سبام / خارج الموضوع" is primary_intent=SPAM only — topic OFF_TOPIC_PERSONAL
# is a separate judgment under "موضوع شخصي / خارج المعجم".
INTENT_OPTIONS: tuple[tuple[str, str], ...] = (
    ("اقتراح", "SUGGESTION"),
    ("سؤال / استفسار", "INQUIRY"),
    ("شكوى / انتقاد", "COMPLAINT"),
    ("شكر / مدح", "THANKS_PRAISE"),
    ("دعاء", "DUA_BLESSING"),
    ("سبام / خارج الموضوع", "SPAM"),
    ("النص غير كافٍ للحكم", "INSUFFICIENT_CONTEXT"),
)

TOPIC_OPTIONS: tuple[tuple[str, str], ...] = (
    ("إضافة كلمة", "WORD_SUBMISSION"),
    ("تصحيح مدخل / معنى / كتابة", "ENTRY_CORRECTION"),
    ("اقتراح متعلق بالمعجم أو ميزة", "DICTIONARY_SUGGESTION"),
    ("مشكلة تقنية", "TECHNICAL_ISSUE"),
    ("ترشيح كلمة العام", "WORD_OF_YEAR_NOMINATION"),
    ("موضوع شخصي / خارج المعجم", "OFF_TOPIC_PERSONAL"),
    ("موضوع عام", "GENERAL"),
)

SENTIMENT_OPTIONS: tuple[tuple[str, str], ...] = (
    ("إيجابية", "POSITIVE"),
    ("محايدة", "NEUTRAL"),
    ("سلبية", "NEGATIVE"),
)

SUBCATEGORY_OPTIONS: tuple[tuple[str, str], ...] = (
    ("خطأ في المعنى", "MeaningFault"),
    ("مشكلة تقنية", "TechFault"),
    ("خطأ إملائي", "typo"),
    ("خطأ نحوي", "GrammerFault"),
    ("أخرى", "Other"),
)

# One-line Arabic hints for reviewers (Gradio wizard / labeling guide alignment).
INTENT_HINTS: dict[str, str] = {
    "SUGGESTION": "اقتراح أو تقديم هادئ للمعجم — بدون استياء في الصياغة.",
    "INQUIRY": "سؤال حقيقي يطلب معلومة («كيف»، «هل»، «ما معنى»).",
    "COMPLAINT": "استياء أو لوم أو انتقاد في النبرة — حتى لو الموضوع تصحيح.",
    "THANKS_PRAISE": "شكر أو مدح secular للمنصة — بدون دعاء ديني.",
    "DUA_BLESSING": "دعاء أو تبريك («جزاكم الله»، «بارك الله») — حتى لو يقرأ كشكر.",
    "SPAM": "ليس تقديمًا معجميًا حقيقيًا — spam/إعلان/محتوى شخصي بلا غرض معجمي (≠ topic).",
    "INSUFFICIENT_CONTEXT": "لا يمكن الحكم — اترك gold_* فارغًا.",
}

TOPIC_HINTS: dict[str, str] = {
    "WORD_SUBMISSION": "كلمة أو مدخل جديد غير موجود بعد في المعجم.",
    "ENTRY_CORRECTION": "خطأ في مدخل موجود (معنى، إملاء، entryId).",
    "DICTIONARY_SUGGESTION": "فكرة لميزة أو عملية في المنصة — لا لكلمة بعينها.",
    "TECHNICAL_ISSUE": "عطل أو bug في الموقع/التطبيق.",
    "WORD_OF_YEAR_NOMINATION": "ترشيح صريح لحملة «كلمة العام» فقط.",
    "OFF_TOPIC_PERSONAL": "محتوى شخصي أو خارج المعجم — ليس شكرًا/general فقط.",
    "GENERAL": "شكر/دعاء/ملاحظة عامة بلا موضوع معجمي محدد.",
}

MANUAL_EDIT_GUIDANCE_MD = """**تذكير سريع**
- **شكوى vs اقتراح:** استياء/لوم → شكوى؛ تقديم هادئ → اقتراح.
- **إضافة كلمة vs تصحيح:** موجود وفيه خطأ → تصحيح؛ جديد → إضافة.
- **عام vs خارج المعجم:** شكر/دعاء فقط → عام؛ قصة/slang/محتوى شخصي → خارج المعجم.
- **نص غامض:** استخدم «النص غير كافٍ للحكم» — لا تخمّن."""

INTENT_LABEL_TO_CODE = {label: code for label, code in INTENT_OPTIONS}
INTENT_CODE_TO_LABEL = {code: label for label, code in INTENT_OPTIONS}
TOPIC_LABEL_TO_CODE = {label: code for label, code in TOPIC_OPTIONS}
TOPIC_CODE_TO_LABEL = {code: label for label, code in TOPIC_OPTIONS}
SENTIMENT_LABEL_TO_CODE = {label: code for label, code in SENTIMENT_OPTIONS}
SENTIMENT_CODE_TO_LABEL = {code: label for label, code in SENTIMENT_OPTIONS}
SUBCATEGORY_LABEL_TO_CODE = {label: code for label, code in SUBCATEGORY_OPTIONS}
SUBCATEGORY_CODE_TO_LABEL = {code: label for label, code in SUBCATEGORY_OPTIONS}

SECONDARY_INTENT_OPTIONS: tuple[tuple[str, str], ...] = tuple(
    (label, code) for label, code in INTENT_OPTIONS if code != "INSUFFICIENT_CONTEXT"
)

ROW_FIELDS = (
    "comment_id",
    "comment_text",
    "calibration_stratum",
    "b2_v2_primary_intent",
    "b2_v2_topic_subtype",
    "b2_v2_sentiment",
    "b2_v3_primary_intent",
    "b2_v3_topic_subtype",
    "b2_v3_sentiment",
    "b2_v3_confidence",
    "gold_primary_intent",
    "gold_topic_subtype",
    "gold_secondary_intent",
    "gold_subcategory",
    "gold_sentiment",
    "gold_annotation_status",
    "labeled_at",
)


def _atomic_write_text(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(content, encoding="utf-8")
    os.replace(tmp, path)


def _atomic_write_json(path: Path, payload: dict[str, Any]) -> None:
    _atomic_write_text(path, json.dumps(payload, ensure_ascii=False, indent=2))


def _load_confidence_by_comment_id(gold_csv_path: Path = GOLD_CSV) -> dict[str, str]:
    if not gold_csv_path.exists():
        return {}
    df = pd.read_csv(gold_csv_path, dtype=str, keep_default_na=False)
    subset = df[df["annotation_stage"] == "ai_assisted_train"]
    return dict(zip(subset["comment_id"], subset["ai_suggested_confidence"], strict=False))


def load_calibration_workbook(
    workbook_path: Path = CALIBRATION_XLSX,
    manifest_path: Path = CALIBRATION_MANIFEST,
    gold_csv_path: Path = GOLD_CSV,
) -> list[dict[str, str]]:
    if not workbook_path.exists():
        raise FileNotFoundError(f"calibration workbook not found: {workbook_path}")
    df = pd.read_excel(workbook_path, sheet_name="Review", dtype=str).fillna("")
    confidence = _load_confidence_by_comment_id(gold_csv_path)

    order: list[str] = []
    if manifest_path.exists():
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        order = manifest.get("comment_ids_in_workbook_order", [])

    rows: list[dict[str, str]] = []
    for _, record in df.iterrows():
        row = {field: str(record.get(field, "")).strip() for field in ROW_FIELDS if field != "b2_v3_confidence"}
        row["b2_v3_confidence"] = confidence.get(row["comment_id"], "")
        rows.append(row)

    if order:
        index = {comment_id: i for i, comment_id in enumerate(order)}
        rows.sort(key=lambda r: index.get(r["comment_id"], 10_000))
    else:
        rows.sort(key=lambda r: r["comment_id"])
    return rows


def empty_review_state(
    workbook_path: Path = CALIBRATION_XLSX,
    manifest_path: Path = CALIBRATION_MANIFEST,
    gold_csv_path: Path = GOLD_CSV,
) -> dict[str, Any]:
    rows = load_calibration_workbook(workbook_path, manifest_path, gold_csv_path)
    for row in rows:
        row["gold_primary_intent"] = ""
        row["gold_topic_subtype"] = ""
        row["gold_secondary_intent"] = ""
        row["gold_subcategory"] = ""
        row["gold_sentiment"] = ""
        row["gold_annotation_status"] = ""
        row["labeled_at"] = ""
    return {
        "version": 1,
        "source_workbook": str(workbook_path),
        "manifest": str(manifest_path) if manifest_path.exists() else "",
        "current_index": 0,
        "rows": rows,
        "updated_at": "",
        "completed_at": "",
    }


def load_or_init_review_state(
    review_path: Path = REVIEW_JSON,
    workbook_path: Path = CALIBRATION_XLSX,
    manifest_path: Path = CALIBRATION_MANIFEST,
    gold_csv_path: Path = GOLD_CSV,
) -> dict[str, Any]:
    if review_path.exists():
        state = json.loads(review_path.read_text(encoding="utf-8"))
        source_rows = load_calibration_workbook(workbook_path, manifest_path, gold_csv_path)
        source_by_id = {row["comment_id"]: row for row in source_rows}
        merged_rows: list[dict[str, str]] = []
        for saved in state.get("rows", []):
            cid = saved["comment_id"]
            if cid not in source_by_id:
                raise AssertionError(f"review artifact references unknown comment_id: {cid}")
            merged = dict(source_by_id[cid])
            for field in (
                "gold_primary_intent",
                "gold_topic_subtype",
                "gold_secondary_intent",
                "gold_subcategory",
                "gold_sentiment",
                "gold_annotation_status",
                "labeled_at",
            ):
                merged[field] = str(saved.get(field, "")).strip()
            merged_rows.append(merged)
        if len(merged_rows) != len(source_rows):
            raise AssertionError("review artifact row count does not match calibration workbook")
        state["rows"] = merged_rows
        state["current_index"] = min(int(state.get("current_index", 0)), len(merged_rows) - 1)
        return ensure_resume_index(state)
    return empty_review_state(workbook_path, manifest_path, gold_csv_path)


def save_review_state(state: dict[str, Any], review_path: Path = REVIEW_JSON) -> None:
    state = dict(state)
    state["updated_at"] = dt.datetime.now(dt.timezone.utc).replace(microsecond=0).isoformat()
    _atomic_write_json(review_path, state)


COMPLETED_ANNOTATION_STATUSES = frozenset({"LABELED", "INSUFFICIENT_CONTEXT"})


def first_incomplete_index(rows: list[dict[str, str]]) -> int:
    for index, row in enumerate(rows):
        if row.get("gold_annotation_status", "") not in COMPLETED_ANNOTATION_STATUSES:
            return index
    return max(0, len(rows) - 1)


def ensure_resume_index(state: dict[str, Any]) -> dict[str, Any]:
    """Position the wizard on the first row that still needs human input."""
    state = dict(state)
    state["current_index"] = first_incomplete_index(state["rows"])
    return state


def current_row(state: dict[str, Any]) -> dict[str, str]:
    return state["rows"][state["current_index"]]


def current_comment_id(state: dict[str, Any]) -> str:
    return current_row(state)["comment_id"]


def validate_human_answers(
    *,
    intent_label: str,
    topic_label: str = "",
    sentiment_label: str = "",
    subcategory_label: str = "",
    secondary_intent_label: str = "",
) -> str | None:
    if not intent_label:
        return "اختر الغرض الأساسي من التعليق."
    if intent_label == INTENT_CODE_TO_LABEL["INSUFFICIENT_CONTEXT"]:
        return None
    if not topic_label:
        return "اختر موضوع التعليق."
    if not sentiment_label:
        return "اختر نبرة التعليق."
    topic_code = TOPIC_LABEL_TO_CODE[topic_label]
    if topic_requires_subcategory(topic_code) and not subcategory_label:
        return "اختر نوع التصحيح/المشكلة."
    if secondary_intent_label and secondary_intent_label == intent_label:
        return "النية الثانية يجب أن تختلف عن النية الأساسية."
    return None


def save_current_row(
    state: dict[str, Any],
    *,
    intent_label: str,
    topic_label: str = "",
    sentiment_label: str = "",
    subcategory_label: str = "",
    secondary_intent_label: str = "",
    review_path: Path = REVIEW_JSON,
) -> tuple[dict[str, Any], str | None]:
    error = validate_human_answers(
        intent_label=intent_label,
        topic_label=topic_label,
        sentiment_label=sentiment_label,
        subcategory_label=subcategory_label,
        secondary_intent_label=secondary_intent_label,
    )
    if error:
        return state, error

    state = dict(state)
    rows = [dict(row) for row in state["rows"]]
    index = state["current_index"]
    rows[index] = apply_human_answers(
        rows[index],
        intent_label=intent_label,
        topic_label=topic_label,
        sentiment_label=sentiment_label,
        subcategory_label=subcategory_label,
        secondary_intent_label=secondary_intent_label,
    )
    labeler_id = str(state.get("labeler_id", "")).strip()
    if labeler_id:
        rows[index]["labeled_by"] = labeler_id
        if not str(rows[index].get("labeler_notes", "")).strip():
            rows[index]["labeler_notes"] = str(state.get("import_notes", "")).strip()
    state["rows"] = rows
    save_review_state(state, review_path)
    return state, None


def save_and_advance_row(
    state: dict[str, Any],
    *,
    intent_label: str,
    topic_label: str = "",
    sentiment_label: str = "",
    subcategory_label: str = "",
    secondary_intent_label: str = "",
    review_path: Path = REVIEW_JSON,
) -> tuple[dict[str, Any], str | None]:
    state, error = save_current_row(
        state,
        intent_label=intent_label,
        topic_label=topic_label,
        sentiment_label=sentiment_label,
        subcategory_label=subcategory_label,
        secondary_intent_label=secondary_intent_label,
        review_path=review_path,
    )
    if error:
        return state, error

    state = dict(state)
    if state["current_index"] < len(state["rows"]) - 1:
        state["current_index"] += 1
        save_review_state(state, review_path)
    return state, None


def go_previous_row(state: dict[str, Any], review_path: Path = REVIEW_JSON) -> dict[str, Any]:
    state = dict(state)
    if state["current_index"] > 0:
        state["current_index"] -= 1
        save_review_state(state, review_path)
    return state


def normalize_calibration_labeled_at(raw: str) -> str:
    """Gold trace validator accepts YYYY-MM-DD only."""
    value = str(raw).strip()
    if not value:
        return ""
    if "T" in value:
        return value.split("T", 1)[0]
    return value[:10]


def completed_calibration_rows(review_path: Path = REVIEW_JSON) -> list[dict[str, str]]:
    if not review_path.exists():
        raise FileNotFoundError(f"calibration review artifact not found: {review_path}")
    state = json.loads(review_path.read_text(encoding="utf-8"))
    return [
        row
        for row in state.get("rows", [])
        if row.get("gold_annotation_status", "") in COMPLETED_ANNOTATION_STATUSES
    ]


def completed_calibration_comment_ids(review_path: Path = REVIEW_JSON) -> set[str]:
    return {row["comment_id"] for row in completed_calibration_rows(review_path)}


def calibration_human_labels_by_id(review_path: Path = REVIEW_JSON) -> dict[str, dict[str, str]]:
    labels: dict[str, dict[str, str]] = {}
    for row in completed_calibration_rows(review_path):
        labels[row["comment_id"]] = {
            "gold_primary_intent": row.get("gold_primary_intent", ""),
            "gold_topic_subtype": row.get("gold_topic_subtype", ""),
            "gold_secondary_intent": row.get("gold_secondary_intent", ""),
            "gold_subcategory": row.get("gold_subcategory", ""),
            "gold_sentiment": row.get("gold_sentiment", ""),
            "gold_annotation_status": row.get("gold_annotation_status", ""),
            "labeler_notes": row.get("labeler_notes", "") or CALIBRATION_IMPORT_NOTES,
            "labeled_by": row.get("labeled_by", "") or CALIBRATION_LABELER_ID,
            "labeled_at": normalize_calibration_labeled_at(row.get("labeled_at", "")),
        }
    return labels


def topic_requires_subcategory(topic_code: str) -> bool:
    return topic_code in settings.TOPICS_REQUIRING_SUBCATEGORY


def apply_human_answers(
    row: dict[str, str],
    *,
    intent_label: str,
    topic_label: str = "",
    sentiment_label: str = "",
    subcategory_label: str = "",
    secondary_intent_label: str = "",
    labeled_by: str = "",
    labeler_notes: str = "",
    use_date_only_labeled_at: bool = False,
) -> dict[str, str]:
    updated = dict(row)
    intent_code = INTENT_LABEL_TO_CODE[intent_label]
    if use_date_only_labeled_at:
        labeled_at = dt.date.today().isoformat()
    else:
        labeled_at = dt.datetime.now(dt.timezone.utc).replace(microsecond=0).isoformat()

    if intent_code == "INSUFFICIENT_CONTEXT":
        updated["gold_primary_intent"] = ""
        updated["gold_topic_subtype"] = ""
        updated["gold_secondary_intent"] = ""
        updated["gold_subcategory"] = ""
        updated["gold_sentiment"] = ""
        updated["gold_annotation_status"] = "INSUFFICIENT_CONTEXT"
        updated["labeled_at"] = labeled_at
        if labeled_by:
            updated["labeled_by"] = labeled_by
        if labeler_notes:
            updated["labeler_notes"] = labeler_notes
        return updated

    if intent_label not in INTENT_LABEL_TO_CODE:
        raise ValueError(f"unknown intent label: {intent_label!r}")
    if topic_label not in TOPIC_LABEL_TO_CODE:
        raise ValueError(f"unknown topic label: {topic_label!r}")
    if sentiment_label not in SENTIMENT_LABEL_TO_CODE:
        raise ValueError(f"unknown sentiment label: {sentiment_label!r}")

    topic_code = TOPIC_LABEL_TO_CODE[topic_label]
    if topic_requires_subcategory(topic_code):
        if subcategory_label not in SUBCATEGORY_LABEL_TO_CODE:
            raise ValueError(f"subcategory required for topic {topic_code}")
        updated["gold_subcategory"] = SUBCATEGORY_LABEL_TO_CODE[subcategory_label]
    else:
        updated["gold_subcategory"] = ""

    updated["gold_primary_intent"] = intent_code
    updated["gold_topic_subtype"] = topic_code
    updated["gold_sentiment"] = SENTIMENT_LABEL_TO_CODE[sentiment_label]
    updated["gold_annotation_status"] = "LABELED"

    if secondary_intent_label:
        secondary_code = INTENT_LABEL_TO_CODE[secondary_intent_label]
        if secondary_code == intent_code:
            raise ValueError("secondary intent must differ from primary intent")
        updated["gold_secondary_intent"] = secondary_code
    else:
        updated["gold_secondary_intent"] = ""

    updated["labeled_at"] = labeled_at
    if labeled_by:
        updated["labeled_by"] = labeled_by
    if labeler_notes:
        updated["labeler_notes"] = labeler_notes
    return updated


def row_to_form_labels(row: dict[str, str]) -> dict[str, str]:
    if row.get("gold_annotation_status") == "INSUFFICIENT_CONTEXT":
        return {
            "intent_label": INTENT_CODE_TO_LABEL["INSUFFICIENT_CONTEXT"],
            "topic_label": "",
            "sentiment_label": "",
            "subcategory_label": "",
            "secondary_intent_label": "",
        }
    return {
        "intent_label": INTENT_CODE_TO_LABEL.get(row.get("gold_primary_intent", ""), ""),
        "topic_label": TOPIC_CODE_TO_LABEL.get(row.get("gold_topic_subtype", ""), ""),
        "sentiment_label": SENTIMENT_CODE_TO_LABEL.get(row.get("gold_sentiment", ""), ""),
        "subcategory_label": SUBCATEGORY_CODE_TO_LABEL.get(row.get("gold_subcategory", ""), ""),
        "secondary_intent_label": INTENT_CODE_TO_LABEL.get(row.get("gold_secondary_intent", ""), ""),
    }


def count_completed_rows(state: dict[str, Any]) -> int:
    done = 0
    for row in state["rows"]:
        status = row.get("gold_annotation_status", "")
        if status in {"LABELED", "INSUFFICIENT_CONTEXT"}:
            done += 1
    return done


def is_review_complete(state: dict[str, Any]) -> bool:
    return count_completed_rows(state) == len(state["rows"])


def _labeled_rows(state: dict[str, Any]) -> list[dict[str, str]]:
    return [row for row in state["rows"] if row.get("gold_annotation_status") == "LABELED"]


def _format_confusion(report: EvaluationReport) -> list[list[str]]:
    header = [""] + list(report.labels)
    matrix = [header]
    for i, true_label in enumerate(report.labels):
        row = [true_label] + [str(report.confusion_matrix[i][j]) for j in range(len(report.labels))]
        matrix.append(row)
    return matrix


def _report_to_dict(name: str, report: EvaluationReport) -> dict[str, Any]:
    return {
        "name": name,
        "macro_f1": report.macro_f1,
        "labels": report.labels,
        "per_class_precision": report.per_class_precision,
        "per_class_recall": report.per_class_recall,
        "per_class_f1": report.per_class_f1,
        "per_class_support": report.per_class_support,
        "confusion_matrix": report.confusion_matrix,
    }


@dataclass(frozen=True)
class CalibrationEvaluation:
    insufficient_context_count: int
    labeled_count: int
    intent_b2_v2: EvaluationReport
    intent_b2_v3: EvaluationReport
    topic_b2_v2: EvaluationReport
    topic_b2_v3: EvaluationReport
    sentiment_b2_v2: EvaluationReport
    sentiment_b2_v3: EvaluationReport
    high_confidence_mismatch_b2_v2: list[dict[str, str]]
    high_confidence_mismatch_b2_v3: list[dict[str, str]]
    high_confidence_insufficient_b2_v2: list[dict[str, str]]
    high_confidence_insufficient_b2_v3: list[dict[str, str]]


def evaluate_calibration(state: dict[str, Any]) -> CalibrationEvaluation:
    labeled = _labeled_rows(state)
    insufficient = [
        row for row in state["rows"] if row.get("gold_annotation_status") == "INSUFFICIENT_CONTEXT"
    ]

    y_true_intent = [row["gold_primary_intent"] for row in labeled]
    y_true_topic = [row["gold_topic_subtype"] for row in labeled]
    y_true_sentiment = [row["gold_sentiment"] for row in labeled]

    b2_v2_intent = [row["b2_v2_primary_intent"] for row in labeled]
    b2_v3_intent = [row["b2_v3_primary_intent"] for row in labeled]
    b2_v2_topic = [row["b2_v2_topic_subtype"] for row in labeled]
    b2_v3_topic = [row["b2_v3_topic_subtype"] for row in labeled]
    b2_v2_sentiment = [row["b2_v2_sentiment"] for row in labeled]
    b2_v3_sentiment = [row["b2_v3_sentiment"] for row in labeled]

    intent_b2_v2 = evaluate(y_true_intent, b2_v2_intent, labels=list(settings.PRIMARY_INTENTS))
    intent_b2_v3 = evaluate(y_true_intent, b2_v3_intent, labels=list(settings.PRIMARY_INTENTS))
    topic_b2_v2 = evaluate(y_true_topic, b2_v2_topic, labels=list(settings.TOPIC_SUBTYPES))
    topic_b2_v3 = evaluate(y_true_topic, b2_v3_topic, labels=list(settings.TOPIC_SUBTYPES))
    sentiment_b2_v2 = evaluate(y_true_sentiment, b2_v2_sentiment, labels=list(settings.SENTIMENTS))
    sentiment_b2_v3 = evaluate(y_true_sentiment, b2_v3_sentiment, labels=list(settings.SENTIMENTS))

    def _high_conf_mismatch(version: str) -> list[dict[str, str]]:
        out: list[dict[str, str]] = []
        for row in labeled:
            conf_raw = row.get("b2_v3_confidence", "")
            try:
                conf = float(conf_raw)
            except ValueError:
                continue
            if conf < HIGH_CONFIDENCE_THRESHOLD:
                continue
            pred_intent = row[f"{version}_primary_intent"]
            pred_topic = row[f"{version}_topic_subtype"]
            if pred_intent != row["gold_primary_intent"] or pred_topic != row["gold_topic_subtype"]:
                out.append(
                    {
                        "comment_id": row["comment_id"],
                        "comment_text": row["comment_text"][:120],
                        "human": f"{row['gold_primary_intent']}/{row['gold_topic_subtype']}",
                        "predicted": f"{pred_intent}/{pred_topic}",
                        "confidence": conf_raw,
                    }
                )
        return out

    def _high_conf_insufficient(version: str) -> list[dict[str, str]]:
        out: list[dict[str, str]] = []
        for row in insufficient:
            conf_raw = row.get("b2_v3_confidence", "")
            try:
                conf = float(conf_raw)
            except ValueError:
                continue
            if conf < HIGH_CONFIDENCE_THRESHOLD:
                continue
            pred_intent = row[f"{version}_primary_intent"]
            pred_topic = row[f"{version}_topic_subtype"]
            out.append(
                {
                    "comment_id": row["comment_id"],
                    "comment_text": row["comment_text"][:120],
                    "predicted": f"{pred_intent}/{pred_topic}",
                    "confidence": conf_raw,
                }
            )
        return out

    return CalibrationEvaluation(
        insufficient_context_count=len(insufficient),
        labeled_count=len(labeled),
        intent_b2_v2=intent_b2_v2,
        intent_b2_v3=intent_b2_v3,
        topic_b2_v2=topic_b2_v2,
        topic_b2_v3=topic_b2_v3,
        sentiment_b2_v2=sentiment_b2_v2,
        sentiment_b2_v3=sentiment_b2_v3,
        high_confidence_mismatch_b2_v2=_high_conf_mismatch("b2_v2"),
        high_confidence_mismatch_b2_v3=_high_conf_mismatch("b2_v3"),
        high_confidence_insufficient_b2_v2=_high_conf_insufficient("b2_v2"),
        high_confidence_insufficient_b2_v3=_high_conf_insufficient("b2_v3"),
    )


def format_evaluation_report(result: CalibrationEvaluation) -> str:
    lines = [
        "B.2 Calibration Evaluation — Human Gold vs Gemini",
        "=" * 60,
        f"Human LABELED rows: {result.labeled_count}",
        f"Human INSUFFICIENT_CONTEXT rows: {result.insufficient_context_count}",
        "",
        "Primary intent Macro-F1:",
        f"  b2-v2 vs Human Gold: {result.intent_b2_v2.macro_f1:.4f}",
        f"  b2-v3 vs Human Gold: {result.intent_b2_v3.macro_f1:.4f}",
        "",
        "Topic subtype Macro-F1:",
        f"  b2-v2 vs Human Gold: {result.topic_b2_v2.macro_f1:.4f}",
        f"  b2-v3 vs Human Gold: {result.topic_b2_v3.macro_f1:.4f}",
        "",
        "Sentiment Macro-F1:",
        f"  b2-v2 vs Human Gold: {result.sentiment_b2_v2.macro_f1:.4f}",
        f"  b2-v3 vs Human Gold: {result.sentiment_b2_v3.macro_f1:.4f}",
    ]

    for title, report in (
        ("Intent b2-v2", result.intent_b2_v2),
        ("Intent b2-v3", result.intent_b2_v3),
        ("Topic b2-v2", result.topic_b2_v2),
        ("Topic b2-v3", result.topic_b2_v3),
    ):
        lines.extend(["", f"{title} per-class F1:"])
        for label in report.labels:
            if report.per_class_support.get(label, 0) == 0 and report.per_class_f1.get(label, 0) == 0:
                continue
            lines.append(
                f"  {label:24} P={report.per_class_precision[label]:.3f} "
                f"R={report.per_class_recall[label]:.3f} F1={report.per_class_f1[label]:.3f} "
                f"(support={report.per_class_support[label]})"
            )
        lines.append(f"  Confusion matrix ({title}):")
        for row in _format_confusion(report):
            lines.append("    " + " | ".join(row))

    lines.extend(
        [
            "",
            f"High-confidence (>= {HIGH_CONFIDENCE_THRESHOLD}) mismatches vs Human Gold:",
            f"  b2-v2: {len(result.high_confidence_mismatch_b2_v2)}",
            f"  b2-v3: {len(result.high_confidence_mismatch_b2_v3)}",
            "",
            f"High-confidence predictions when Human marked INSUFFICIENT_CONTEXT:",
            f"  b2-v2: {len(result.high_confidence_insufficient_b2_v2)}",
            f"  b2-v3: {len(result.high_confidence_insufficient_b2_v3)}",
        ]
    )
    return "\n".join(lines)


def evaluation_to_json(result: CalibrationEvaluation) -> dict[str, Any]:
    return {
        "labeled_count": result.labeled_count,
        "insufficient_context_count": result.insufficient_context_count,
        "high_confidence_threshold": HIGH_CONFIDENCE_THRESHOLD,
        "intent": {
            "b2_v2": _report_to_dict("intent_b2_v2", result.intent_b2_v2),
            "b2_v3": _report_to_dict("intent_b2_v3", result.intent_b2_v3),
        },
        "topic": {
            "b2_v2": _report_to_dict("topic_b2_v2", result.topic_b2_v2),
            "b2_v3": _report_to_dict("topic_b2_v3", result.topic_b2_v3),
        },
        "sentiment": {
            "b2_v2": _report_to_dict("sentiment_b2_v2", result.sentiment_b2_v2),
            "b2_v3": _report_to_dict("sentiment_b2_v3", result.sentiment_b2_v3),
        },
        "high_confidence_mismatch_b2_v2": result.high_confidence_mismatch_b2_v2,
        "high_confidence_mismatch_b2_v3": result.high_confidence_mismatch_b2_v3,
        "high_confidence_insufficient_b2_v2": result.high_confidence_insufficient_b2_v2,
        "high_confidence_insufficient_b2_v3": result.high_confidence_insufficient_b2_v3,
    }


def write_labeled_workbook(state: dict[str, Any], out_path: Path = LABELED_XLSX) -> Path:
    columns = [
        "comment_id",
        "comment_text",
        "calibration_stratum",
        "b2_v2_primary_intent",
        "b2_v2_topic_subtype",
        "b2_v2_sentiment",
        "b2_v3_primary_intent",
        "b2_v3_topic_subtype",
        "b2_v3_sentiment",
        "gold_primary_intent",
        "gold_topic_subtype",
        "gold_secondary_intent",
        "gold_subcategory",
        "gold_sentiment",
        "gold_annotation_status",
        "labeled_at",
    ]
    wb = Workbook()
    ws = wb.active
    ws.title = "Labeled"
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for col_idx, name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.fill = header_fill
        cell.font = header_font
    for row_idx, row in enumerate(state["rows"], start=2):
        for col_idx, name in enumerate(columns, start=1):
            value = row.get(name, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if name == "comment_text":
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def finalize_calibration_review(
    state: dict[str, Any],
    *,
    review_path: Path = REVIEW_JSON,
    labeled_xlsx_path: Path = LABELED_XLSX,
    eval_json_path: Path = EVAL_REPORT_JSON,
    eval_txt_path: Path = EVAL_REPORT_TXT,
) -> tuple[Path, Path, Path, str]:
    if not is_review_complete(state):
        raise AssertionError("all 40 rows must be labeled before finalization")
    evaluation = evaluate_calibration(state)
    report_text = format_evaluation_report(evaluation)
    state = dict(state)
    state["completed_at"] = dt.datetime.now(dt.timezone.utc).replace(microsecond=0).isoformat()
    save_review_state(state, review_path)
    labeled_path = write_labeled_workbook(state, labeled_xlsx_path)
    _atomic_write_json(eval_json_path, evaluation_to_json(evaluation))
    _atomic_write_text(eval_txt_path, report_text)
    return labeled_path, eval_json_path, eval_txt_path, report_text
