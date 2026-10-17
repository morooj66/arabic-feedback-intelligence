"""Demo-mode upload policy: only workbooks that carry the synthetic-data marker, and only small ones. This stops an ACCIDENTAL upload of real comments into a hosted demo;
it is not a defence against a determined person (the demo is invite-only for that reason)."""

from __future__ import annotations

from pathlib import Path

MARKER = "RCS-SYNTHETIC-DEMO-WORKBOOK-v1"
MAX_ROWS = 200
MAX_BYTES = 256 * 1024
REFUSAL = "في وضع العرض التجريبي يُقبل فقط ملف البيانات الاصطناعية الذي يوفّره المشروع (examples/synthetic_comments.xlsx)."


def stamp(path: Path) -> None:
    """Writes the marker into the workbook's document properties (used by the project's own synthetic workbooks)."""
    from openpyxl import load_workbook

    wb = load_workbook(path)
    wb.properties.title = MARKER
    wb.save(path)


def has_marker(path: Path) -> bool:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(path, read_only=True)
        try:
            return (wb.properties.title or "") == MARKER
        finally:
            wb.close()
    except Exception:  # noqa: BLE001
        return False


def wrap_validator(original):
    """Returns a validate_workbook replacement that applies the base validation first, then the demo policy."""

    def validate(path: Path):
        base = original(path)
        if not base.ok:
            return base
        errors: list[str] = []
        if not has_marker(path):
            errors.append(REFUSAL)
        elif base.total_rows > MAX_ROWS:
            errors.append(f"الحد الأقصى في العرض التجريبي {MAX_ROWS} صفًا.")
        if errors:
            from rcs.review.excel_upload import WorkbookValidation

            return WorkbookValidation(False, base.total_rows, errors)
        return base

    return validate
