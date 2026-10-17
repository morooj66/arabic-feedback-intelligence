#!/usr/bin/env python3
"""One-shot response-quality inspection batch (30 comments).

Does not change prompt, model, routing, classifiers, reviewer workflow,
or gold/human_test artifacts. Does not send replies to visitors.
"""

from __future__ import annotations

import json
import os
import re
import sys
import time
from collections import Counter
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent


def _load_dotenv_if_present(env_path: Path) -> None:
    if not env_path.exists():
        return
    for line in env_path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, _, value = line.partition("=")
        key, value = key.strip(), value.strip()
        if key and key not in os.environ:
            os.environ[key] = value


_load_dotenv_if_present(PROJECT_ROOT / ".env")
sys.path.insert(0, str(PROJECT_ROOT / "src"))
sys.path.insert(0, str(PROJECT_ROOT))

import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402
from sklearn.feature_extraction.text import TfidfVectorizer  # noqa: E402
from sklearn.linear_model import LogisticRegression  # noqa: E402
from sklearn.pipeline import Pipeline  # noqa: E402

from config import settings  # noqa: E402
from rcs.classification.redact import redact_phone_numbers  # noqa: E402
from rcs.response_generation.prompt import _EMAIL_RE  # noqa: E402
from rcs.response_generation.eligibility import EligibilityInput, determine_response_eligibility  # noqa: E402
from rcs.response_generation.service import (  # noqa: E402
    GenerationRequest,
    build_live_gemini_client,
    generate_suggested_response,
)
from rcs.response_generation.validate import _LEAK_PATTERNS, validate_draft  # noqa: E402
from rcs.routing.engine import determine_routing  # noqa: E402
from rcs.training.gold_splits import LabeledExample, load_supervised_gold  # noqa: E402
from rcs.training.hf_loading import load_sequence_classifier, load_tokenizer  # noqa: E402
from rcs.training.sentiment_pretrained import _normalize_sentiment  # noqa: E402

GOLD_CSV = PROJECT_ROOT / "data" / "gold" / "gold_dataset_final_training.csv"
EXCLUSION_JSON = PROJECT_ROOT / "data" / "gold" / "human_test_exclusion_ids.json"
FREEZE_JSON = PROJECT_ROOT / "data" / "models" / "selection" / "FROZEN_MODEL_SELECTION.json"
OUT_XLSX = PROJECT_ROOT / "response_quality_review_30.xlsx"
SAMPLE_SEED = 43
SAMPLE_N = 30
INTENT_EVAL_REMAP = {"DUA_BLESSING": "THANKS_PRAISE"}

HUMAN_COLS = (
    "human_relevance_score",
    "human_tone_score",
    "human_safety_factuality_score",
    "human_edit_level",
    "hallucination_or_unsupported_claim",
    "human_notes",
)

SYSTEM_COLS = (
    "sample_no",
    "comment_id",
    "comment_text",
    "intent",
    "topic",
    "sentiment",
    "routing_decision",
    "gemini_suggested_response",
    "generation_status",
    "generation_reason",
)

PROMISE_PATTERNS = (
    re.compile(r"سيتم"),
    re.compile(r"سوف (ن|يتم)"),
    re.compile(r"خلال\s*\d+"),
    re.compile(r"خلال\s*(ساعة|يوم|أيام|أسبوع)"),
    re.compile(r"تم (إصلاح|حل|تنفيذ|اعتماد)"),
    re.compile(r"نعدك"),
    re.compile(r"نؤكد أن"),
    re.compile(r"24\s*ساعة"),
    re.compile(r"we will (fix|implement|resolve)", re.I),
)


class CountingGeminiClient:
    def __init__(self, inner):
        self.inner = inner
        self.call_comment_ids: list[str] = []
        self._pending_id: str | None = None

    def set_comment_id(self, comment_id: str) -> None:
        self._pending_id = comment_id

    def generate(self, prompt: str, response_schema: dict) -> str:
        self.call_comment_ids.append(self._pending_id or "")
        return self.inner.generate(prompt, response_schema)


def remap_intent(examples: list[LabeledExample]) -> list[LabeledExample]:
    out: list[LabeledExample] = []
    for ex in examples:
        label = INTENT_EVAL_REMAP.get(ex.label, ex.label)
        out.append(
            LabeledExample(
                comment_id=ex.comment_id,
                text=ex.text,
                label=label,
                annotation_stage=ex.annotation_stage,
                split=ex.split,
            )
        )
    return out


def fit_frozen_intent_pipeline(train: list[LabeledExample]) -> Pipeline:
    freeze = json.loads(FREEZE_JSON.read_text())
    vec_cfg = freeze["intent"]["preprocessing"]["vectorizer"]
    clf_cfg = freeze["intent"]["preprocessing"]["classifier"]
    pipe = Pipeline(
        [
            (
                "tfidf",
                TfidfVectorizer(
                    analyzer=vec_cfg["analyzer"],
                    ngram_range=tuple(vec_cfg["ngram_range"]),
                    min_df=vec_cfg["min_df"],
                    max_features=vec_cfg["max_features"],
                    sublinear_tf=vec_cfg["sublinear_tf"],
                ),
            ),
            (
                "clf",
                LogisticRegression(
                    max_iter=clf_cfg["max_iter"],
                    random_state=clf_cfg["random_state"],
                    class_weight=clf_cfg["class_weight"],
                ),
            ),
        ]
    )
    pipe.fit([ex.text for ex in train], [ex.label for ex in train])
    return pipe


def find_topic_checkpoint() -> Path:
    env = os.environ.get("RCS_MARBERTV2_TOPIC_CHECKPOINT")
    candidates = []
    if env:
        candidates.append(Path(env))
    candidates.extend(
        [
            PROJECT_ROOT / "data" / "models" / "final_test" / "checkpoints" / "marbertv2_topic" / "best",
            PROJECT_ROOT
            / "data"
            / "models"
            / "selection"
            / "expanded_val_108_2026-08-25"
            / "marbertv2_topic"
            / "best",
        ]
    )
    for path in candidates:
        if (path / "config.json").exists() and (
            (path / "model.safetensors").exists() or (path / "pytorch_model.bin").exists()
        ):
            return path
    raise FileNotFoundError("Frozen MARBERTv2 topic checkpoint not found")


def _device() -> str:
    import torch

    if torch.cuda.is_available():
        return "cuda"
    if getattr(torch.backends, "mps", None) and torch.backends.mps.is_available():
        return "mps"
    return "cpu"


def predict_topic(texts: list[str], checkpoint: Path, freeze: dict) -> tuple[list[str], list[float]]:
    import torch

    labels = list(freeze["topic"]["label_order"])
    max_length = int(freeze["topic"]["train_config"]["max_length"])
    batch_size = int(freeze["topic"]["train_config"]["batch_size"])
    device = _device()
    tokenizer = load_tokenizer(str(checkpoint))
    model = load_sequence_classifier(str(checkpoint))
    model.to(device)
    model.eval()
    id2label = {int(k): v for k, v in model.config.id2label.items()}
    preds: list[str] = []
    confs: list[float] = []
    for start in range(0, len(texts), batch_size):
        batch = texts[start : start + batch_size]
        enc = tokenizer(
            batch,
            padding=True,
            truncation=True,
            max_length=max_length,
            return_tensors="pt",
        )
        enc = {k: v.to(device) for k, v in enc.items()}
        with torch.no_grad():
            logits = model(**enc).logits
            probs = torch.softmax(logits, dim=-1).cpu().numpy()
        for i in range(len(batch)):
            pred_idx = int(probs[i].argmax())
            preds.append(id2label[pred_idx])
            row = []
            for lab in labels:
                idx = next((j for j, name in id2label.items() if name == lab), None)
                row.append(float(probs[i][idx]) if idx is not None else 0.0)
            confs.append(max(row) if row else float(probs[i].max()))
    return preds, confs


def predict_sentiment(texts: list[str], freeze: dict) -> tuple[list[str], list[float]]:
    import torch

    model_id = freeze["sentiment"]["model_id"]
    max_length = int(freeze["sentiment"]["preprocessing"]["max_length"])
    batch_size = int(freeze["sentiment"]["preprocessing"]["batch_size"])
    device = _device()
    tokenizer = load_tokenizer(model_id)
    model = load_sequence_classifier(model_id)
    model.to(device)
    model.eval()
    id2label = {int(k): v for k, v in model.config.id2label.items()}
    label_order = list(settings.SENTIMENTS)
    preds: list[str] = []
    confs: list[float] = []
    for start in range(0, len(texts), batch_size):
        batch = texts[start : start + batch_size]
        enc = tokenizer(
            batch,
            padding=True,
            truncation=True,
            max_length=max_length,
            return_tensors="pt",
        )
        enc = {k: v.to(device) for k, v in enc.items()}
        with torch.no_grad():
            logits = model(**enc).logits
            probs = torch.softmax(logits, dim=-1).cpu().numpy()
        for i in range(len(batch)):
            pred_idx = int(probs[i].argmax())
            preds.append(_normalize_sentiment(id2label[pred_idx]))
            row = []
            for project_label in label_order:
                idx = next(
                    (j for j, lbl in id2label.items() if _normalize_sentiment(lbl) == project_label),
                    None,
                )
                row.append(float(probs[i][idx]) if idx is not None else 0.0)
            total = sum(row) or 1.0
            row = [p / total for p in row]
            confs.append(max(row))
    return preds, confs


def _take(rng: np.random.Generator, frame: pd.DataFrame, n: int, used: set[str]) -> list[str]:
    avail = frame[~frame["comment_id"].isin(used)]
    if avail.empty or n <= 0:
        return []
    ids = avail["comment_id"].tolist()
    rng.shuffle(ids)
    chosen = ids[:n]
    used.update(chosen)
    return chosen


def select_30(
    scored: pd.DataFrame,
    rng: np.random.Generator,
) -> tuple[pd.DataFrame, dict]:
    """Diversity sample: GENERATE_DRAFT cases and NO_RESPONSE_RECOMMENDED skips."""
    used: set[str] = set()
    buckets: dict[str, int] = {}
    skip = scored[scored["pred_routing"] == "NO_RESPONSE_RECOMMENDED"]
    gen = scored[scored["pred_routing"] == "GENERATE_DRAFT"]

    buckets["pred_NO_RESPONSE_RECOMMENDED"] = len(_take(rng, skip, 8, used))

    def take_intent(label: str, n: int) -> None:
        ids = _take(rng, gen[gen["gold_primary_intent"] == label], n, used)
        buckets[f"gold_intent:{label}"] = len(ids)

    def take_topic(label: str, n: int) -> None:
        ids = _take(rng, gen[gen["gold_topic_subtype"] == label], n, used)
        buckets[f"gold_topic:{label}"] = len(ids)

    take_intent("THANKS_PRAISE", 4)
    take_intent("SUGGESTION", 4)
    take_intent("COMPLAINT", 4)
    take_intent("INQUIRY", 4)
    take_topic("TECHNICAL_ISSUE", 3)
    take_topic("WORD_OF_YEAR_NOMINATION", 2)
    leftover_n = SAMPLE_N - len(used)
    leftover_ids = _take(rng, gen, leftover_n, used)
    buckets["leftover_generate_fill"] = len(leftover_ids)
    if len(used) < SAMPLE_N:
        buckets["leftover_any_fill"] = len(_take(rng, scored, SAMPLE_N - len(used), used))
    if len(used) != SAMPLE_N:
        raise SystemExit(f"sample size {len(used)} != {SAMPLE_N}")
    order = list(used)
    rng2 = np.random.default_rng(SAMPLE_SEED + 1)
    rng2.shuffle(order)
    ordered = scored.set_index("comment_id").loc[order].reset_index()
    return ordered, {
        "pool_size": int(len(scored)),
        "selection_seed": SAMPLE_SEED,
        "bucket_fills": buckets,
        "source_csv": str(GOLD_CSV.relative_to(PROJECT_ROOT)),
        "pool_pred_routing": dict(Counter(scored["pred_routing"])),
        "note": (
            "Comments drawn from gold_dataset_final_training.csv excluding "
            "sealed human_test. New sample under final eligibility: mix of "
            "GENERATE_DRAFT and NO_RESPONSE_RECOMMENDED. Gold labels diversify "
            "the generate set; frozen-model signals drive the pipeline."
        ),
    }


def precheck_row(row: dict, gemini_called: bool) -> list[str]:
    findings: list[str] = []
    routing = row["routing_decision"]
    status = row["generation_status"]
    draft = row.get("gemini_suggested_response") or ""
    draft = "" if draft is None else str(draft)

    if routing == "NO_RESPONSE_RECOMMENDED":
        if gemini_called:
            findings.append("gemini_called_for_no_response_recommended")
        if draft.strip():
            findings.append("draft_present_when_generation_skipped")
        if status != "SKIPPED":
            findings.append("expected_skip_for_no_response_recommended")
    elif routing == "GENERATE_DRAFT":
        if status == "SKIPPED":
            findings.append("unexpected_skip_when_generate_draft")
        if status == "SUCCEEDED" and not draft.strip():
            findings.append("missing_draft")
        if status == "SUCCEEDED" and not gemini_called:
            findings.append("succeeded_without_gemini_call")
        if status == "FAILED_REVIEW_ONLY":
            findings.append("generation_failed")
    if routing == "GENERATE_DRAFT" and gemini_called is False and status not in (
        "FAILED_REVIEW_ONLY",
        "SKIPPED",
    ):
        findings.append("gemini_not_called_when_generate_draft")

    if draft:
        if len(draft) > settings.RESPONSE_DRAFT_MAX_CHARS:
            findings.append("excessive_response_length")
        if redact_phone_numbers(draft) != draft:
            findings.append("phone_pii_in_draft")
        if _EMAIL_RE.search(draft):
            findings.append("email_pii_in_draft")
        for pattern in _LEAK_PATTERNS:
            if pattern.search(draft):
                findings.append("internal_metadata_leak")
                break
        for pattern in PROMISE_PATTERNS:
            if pattern.search(draft):
                findings.append("unsupported_promise_pattern")
                break
        ok, why = validate_draft(draft)
        if status == "SUCCEEDED" and not ok:
            findings.append(f"malformed_or_invalid_draft:{why}")
    return findings


def write_workbook(rows: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Review"
    ws.sheet_view.rightToLeft = True
    ws.freeze_panes = "A2"
    header_fill = PatternFill("solid", fgColor="1F4E78")
    human_fill = PatternFill("solid", fgColor="FFF2CC")
    header_font = Font(color="FFFFFF", bold=True, name="Calibri", size=11)
    body_font = Font(name="Calibri", size=11)
    wrap = Alignment(wrap_text=True, vertical="top", horizontal="right")
    top = Alignment(vertical="top", wrap_text=True)
    thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    cols = list(SYSTEM_COLS) + list(HUMAN_COLS)
    for col_idx, name in enumerate(cols, 1):
        cell = ws.cell(1, col_idx, name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = thin
    for r_i, row in enumerate(rows, 2):
        for c_i, name in enumerate(SYSTEM_COLS, 1):
            val = row.get(name)
            cell = ws.cell(r_i, c_i, val if val is not None else "")
            cell.font = body_font
            cell.border = thin
            cell.alignment = wrap if name in ("comment_text", "gemini_suggested_response") else top
            if name == "comment_id":
                cell.number_format = "@"
        for c_i, name in enumerate(HUMAN_COLS, 1 + len(SYSTEM_COLS)):
            cell = ws.cell(r_i, c_i, None)
            cell.fill = human_fill
            cell.font = body_font
            cell.border = thin
            cell.alignment = wrap if name == "human_notes" else top
        ws.row_dimensions[r_i].height = 72
    widths = {
        "A": 11,
        "B": 44,
        "C": 48,
        "D": 16,
        "E": 24,
        "F": 12,
        "G": 26,
        "H": 52,
        "I": 16,
        "J": 28,
        "K": 18,
        "L": 16,
        "M": 24,
        "N": 16,
        "O": 28,
        "P": 32,
    }
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{1 + len(rows)}"
    last = 1 + len(rows)
    score_dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True)
    score_dv.error = "Use 1–5"
    score_dv.prompt = "1–5"
    ws.add_data_validation(score_dv)
    score_dv.add(f"K2:M{last}")
    edit_dv = DataValidation(
        type="list",
        formula1='"SEND_AS_IS,MINOR_EDIT,MAJOR_EDIT,UNUSABLE"',
        allow_blank=True,
    )
    ws.add_data_validation(edit_dv)
    edit_dv.add(f"N2:N{last}")
    hall_dv = DataValidation(type="list", formula1='"YES,NO"', allow_blank=True)
    ws.add_data_validation(hall_dv)
    hall_dv.add(f"O2:O{last}")

    ins = wb.create_sheet("Instructions")
    ins.sheet_view.rightToLeft = True
    ins["A1"] = "Response quality review — scoring rubric"
    ins["A1"].font = Font(bold=True, size=14)
    instructions = """This workbook is for human inspection of Gemini drafts under the final eligibility policy.

GENERATE_DRAFT rows have a suggested response. NO_RESPONSE_RECOMMENDED rows correctly have no Gemini draft (spam / off-topic / empty-like / non-actionable). Use NOT sending (NO_RESPONSE_NEEDED) as the human outcome for those; leave scores blank or skip them if you prefer to score drafts only.

human_relevance_score (1–5)
human_tone_score (1–5)
human_safety_factuality_score (1–5)

human_edit_level: SEND_AS_IS / MINOR_EDIT / MAJOR_EDIT / UNUSABLE
hallucination_or_unsupported_claim: YES / NO
human_notes — free text.

Nothing in this file was sent to visitors.
"""
    ins["A3"] = instructions.strip()
    ins["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    ins.column_dimensions["A"].width = 100
    ins.row_dimensions[3].height = 420
    wb.save(OUT_XLSX)


def main() -> None:
    if settings.RESPONSE_GENERATION_MODEL_NAME != "gemini-3.6-flash":
        raise SystemExit(
            f"Configured model is {settings.RESPONSE_GENERATION_MODEL_NAME!r}, "
            "expected gemini-3.6-flash. Not overriding."
        )
    if settings.RESPONSE_GENERATION_PROMPT_VERSION != "rg-v2":
        raise SystemExit("Unexpected prompt version; aborting without changing it.")
    exclusion = set(json.loads(EXCLUSION_JSON.read_text())["test_ids"])
    gold = pd.read_csv(GOLD_CSV, dtype=str, keep_default_na=False)
    human_test_ids = set(gold.loc[gold["annotation_stage"] == "human_test", "comment_id"])
    if human_test_ids - exclusion:
        raise SystemExit("human_test IDs missing from exclusion manifest")
    pool = gold[
        (gold["annotation_stage"] != "human_test")
        & (~gold["comment_id"].isin(exclusion))
        & (gold["comment_text"].astype(str).str.strip() != "")
    ].copy()

    freeze = json.loads(FREEZE_JSON.read_text())
    data = load_supervised_gold(GOLD_CSV, include_test=False)
    intent_pipe = fit_frozen_intent_pipeline(remap_intent(data.intent.train))
    texts = pool["comment_text"].tolist()
    intent_labels = list(intent_pipe.classes_)
    intent_probs = intent_pipe.predict_proba(texts)
    pool["pred_intent"] = [intent_labels[int(np.argmax(p))] for p in intent_probs]
    pool["intent_conf"] = [float(np.max(p)) for p in intent_probs]
    topic_pred, topic_conf = predict_topic(texts, find_topic_checkpoint(), freeze)
    sent_pred, sent_conf = predict_sentiment(texts, freeze)
    pool["pred_topic"] = topic_pred
    pool["topic_conf"] = topic_conf
    pool["pred_sentiment"] = sent_pred
    pool["sent_conf"] = sent_conf
    pool["pred_confidence"] = [
        min(i, t, s) for i, t, s in zip(pool["intent_conf"], pool["topic_conf"], pool["sent_conf"])
    ]
    routes = []
    reasons = []
    class_routes = []
    needs_flags = []
    for rec in pool.itertuples(index=False):
        class_route = determine_routing(rec.pred_intent, rec.pred_topic, None, None, True)
        needs_review = rec.pred_confidence < settings.RESPONSE_LOW_CONFIDENCE_THRESHOLD
        el = determine_response_eligibility(
            EligibilityInput(
                comment_text=rec.comment_text,
                primary_intent=rec.pred_intent,
                topic_subtype=rec.pred_topic,
                sentiment=rec.pred_sentiment,
                confidence=rec.pred_confidence,
                needs_review=needs_review,
                is_content_meaningful=True,
                insufficient_context=rec.gold_annotation_status == "INSUFFICIENT_CONTEXT",
                trusted_grounding_available=False,
                classification_routing_decision=class_route.routing_decision,
            )
        )
        routes.append(el.routing_decision)
        reasons.append(el.reason_code)
        class_routes.append(class_route.routing_decision)
        needs_flags.append(needs_review)
    pool["pred_routing"] = routes
    pool["pred_reason"] = reasons
    pool["class_route"] = class_routes
    pool["pred_needs_review"] = needs_flags

    rng = np.random.default_rng(SAMPLE_SEED)
    sample, selection_meta = select_30(pool, rng)
    selection_meta["excluded_human_test_ids"] = len(exclusion)
    if set(sample["comment_id"]) & exclusion:
        raise SystemExit("sampled a sealed human_test id")
    if set(sample["comment_id"]) & human_test_ids:
        raise SystemExit("sampled annotation_stage=human_test")

    inner = build_live_gemini_client()
    client = CountingGeminiClient(inner)
    rows: list[dict] = []
    for i, rec in enumerate(sample.itertuples(index=False), 1):
        cid = str(rec.comment_id)
        text = str(rec.comment_text)
        intent = rec.pred_intent
        topic = rec.pred_topic
        sentiment = rec.pred_sentiment
        confidence = float(rec.pred_confidence)
        insufficient = str(rec.gold_annotation_status) == "INSUFFICIENT_CONTEXT"
        class_route = rec.class_route
        needs_review = bool(rec.pred_needs_review)
        req = GenerationRequest(
            comment_id=cid,
            comment_text=text,
            primary_intent=intent,
            topic_subtype=topic,
            sentiment=sentiment,
            confidence=confidence,
            needs_review=needs_review,
            is_content_meaningful=True,
            insufficient_context=insufficient,
            trusted_grounding_available=False,
            classification_routing_decision=class_route,
        )
        before = len(client.call_comment_ids)
        client.set_comment_id(cid)
        result = generate_suggested_response(req, client=client)
        called = len(client.call_comment_ids) > before
        if called:
            time.sleep(0.4)
        rows.append(
            {
                "sample_no": i,
                "comment_id": cid,
                "comment_text": text,
                "intent": intent,
                "topic": topic,
                "sentiment": sentiment,
                "confidence": confidence,
                "needs_review": needs_review,
                "routing_decision": result.routing_decision,
                "generation_status": result.generation_status,
                "generation_reason": result.reason_code,
                "gemini_suggested_response": result.suggested_response,
                "_gemini_called": called,
                "_class_route": class_route,
                "_gold_intent": rec.gold_primary_intent,
                "_gold_topic": rec.gold_topic_subtype,
                "_gold_status": rec.gold_annotation_status,
                "_model": result.model_name,
                "_prompt": result.prompt_version,
            }
        )

    write_workbook(rows)

    findings: dict[str, list[int]] = {}
    for row in rows:
        issues = precheck_row(row, bool(row["_gemini_called"]))
        for issue in issues:
            findings.setdefault(issue, []).append(int(row["sample_no"]))

    gemini_n = len(client.call_comment_ids)
    report = {
        "selection": selection_meta,
        "n": len(rows),
        "intent_distribution": dict(Counter(r["intent"] for r in rows)),
        "topic_distribution": dict(Counter(r["topic"] for r in rows)),
        "sentiment_distribution": dict(Counter(r["sentiment"] for r in rows)),
        "routing_distribution": dict(Counter(r["routing_decision"] for r in rows)),
        "reason_code_distribution": dict(Counter(r["generation_reason"] for r in rows)),
        "generation_status_distribution": dict(Counter(r["generation_status"] for r in rows)),
        "gold_intent_of_sampled_comments": dict(Counter(r["_gold_intent"] for r in rows)),
        "gold_topic_of_sampled_comments": dict(Counter(r["_gold_topic"] for r in rows)),
        "gemini_calls": gemini_n,
        "gemini_call_comment_ids": client.call_comment_ids,
        "model": settings.RESPONSE_GENERATION_MODEL_NAME,
        "prompt_version": settings.RESPONSE_GENERATION_PROMPT_VERSION,
        "xlsx": str(OUT_XLSX),
        "human_columns_blank": True,
        "visitor_send": False,
        "optimization": False,
        "automated_precheck": findings,
        "precheck_is_not_human_score": True,
    }
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
