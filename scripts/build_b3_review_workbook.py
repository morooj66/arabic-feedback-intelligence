#!/usr/bin/env python3
"""
Phase B.2 post-run audit, deliverable 8 — build a reviewer-friendly .xlsx
for the (not-yet-started) Phase B.3 human review step.

The default build command is READ-ONLY with respect to the canonical gold
CSV: it reads data/gold/gold_dataset_round1.csv and writes a NEW workbook
(data/gold/gold_dataset_round1_b3_review.xlsx).

The explicit --import-reviewed path is the separate B.3 safety boundary. It
accepts human edits only in gold_*/labeler_* fields, rejects any changed
technical, source-text, split, stage, or ai_suggested_* value, validates the
fully merged dataset, and only then atomically writes the requested output.

Scope: the 300 annotation_stage == "ai_assisted_train" rows only (these are
exactly the rows Phase B.3 review will act on). human_val/human_test/
human_train_seed rows are not included in this workbook.

Column order (as specified): comment_text, ai_suggested_primary_intent,
ai_suggested_topic_subtype, ai_suggested_sentiment, ai_suggested_confidence,
followed by the human gold_* fields a reviewer would fill in to confirm or
correct. Technical/provenance columns are present but hidden (never
deleted), consistent with every prior workbook delivered in this project.

Rows are sorted by ai_suggested_confidence ascending (lowest confidence
first) purely for reviewer convenience/prioritization — this workbook is a
display artifact; it does not change row order or any value in the CSV.
"""

from __future__ import annotations

import argparse
import datetime as dt
import os
import sys
import tempfile
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT / "src"))
sys.path.insert(0, str(PROJECT_ROOT))

import pandas as pd  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill, Protection  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

from config import settings  # noqa: E402
from rcs.evaluation.gold_dataset import (  # noqa: E402
    GOLD_LABEL_COLUMNS,
    HUMAN_TRACE_COLUMNS,
    PREEXISTING_AI_ASSISTED_GOLD_IDS,
    merge_partial_reviewer_labels,
    verify_post_b2_b3_invariants,
)
from rcs.evaluation.calibration_labeling import (  # noqa: E402
    CALIBRATION_LABELER_ID,
    REVIEW_JSON as CALIBRATION_REVIEW_JSON,
    calibration_human_labels_by_id,
    completed_calibration_comment_ids,
)
from rcs.evaluation.b3_labeling import (  # noqa: E402
    B3_REVIEW_JSON,
    import_b3_wizard_labels,
)

GOLD_CSV = settings.DATA_DIR / "gold" / "gold_dataset_round1.csv"
OUT_XLSX = settings.DATA_DIR / "gold" / "gold_dataset_round1_b3_review.xlsx"
OUT_REMAINING_XLSX = settings.DATA_DIR / "gold" / "gold_dataset_round1_b3_review_remaining260.xlsx"

# Visible, reviewer-facing columns, in the exact order requested.
VISIBLE_COLUMNS = [
    "comment_text",
    "ai_suggested_primary_intent",
    "ai_suggested_topic_subtype",
    "ai_suggested_sentiment",
    "ai_suggested_confidence",
    "gold_primary_intent",
    "gold_topic_subtype",
    "gold_secondary_intent",
    "gold_subcategory",
    "gold_sentiment",
    "gold_annotation_status",
    "labeler_notes",
    "labeled_by",
    "labeled_at",
]

# Present but hidden — technical/provenance columns, never deleted.
HIDDEN_COLUMNS = [
    "comment_id",
    "split",
    "annotation_stage",
    "is_few_shot_example",
    "ai_suggested_secondary_intent",
    "ai_suggested_subcategory",
    "ai_suggestion_model",
]

ALL_COLUMNS = VISIBLE_COLUMNS + HIDDEN_COLUMNS
REVIEWER_EDITABLE_COLUMNS = list(GOLD_LABEL_COLUMNS) + list(HUMAN_TRACE_COLUMNS)
PROTECTED_WORKBOOK_COLUMNS = [
    col for col in ALL_COLUMNS if col not in REVIEWER_EDITABLE_COLUMNS
]

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
LOW_CONF_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")  # <0.85
ANOMALY_FILL = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")  # pre-existing gold_* anomaly


def build(
    gold_csv_path: Path = GOLD_CSV,
    out_path: Path = OUT_XLSX,
    *,
    exclude_comment_ids: set[str] | None = None,
    instructions_title: str | None = None,
) -> Path:
    df = pd.read_csv(gold_csv_path, dtype=str, keep_default_na=False)
    review = df[df["annotation_stage"] == "ai_assisted_train"].copy()
    if exclude_comment_ids:
        review = review[~review["comment_id"].isin(exclude_comment_ids)].copy()
    if len(review) == 0:
        raise AssertionError("no ai_assisted_train rows found — nothing to build")

    review["_conf_sort"] = review["ai_suggested_confidence"].astype(float)
    review = review.sort_values("_conf_sort", ascending=True).drop(columns="_conf_sort")
    review = review[ALL_COLUMNS]

    wb = Workbook()

    # --- Instructions sheet -------------------------------------------------
    ws_i = wb.active
    ws_i.title = "Instructions"
    lines = [
        (instructions_title or "Phase B.3 Reviewer Workbook — AI-Suggested Responses for Human Review", True),
        ("", False),
        (f"Scope: the {len(review)} comments in annotation_stage = ai_assisted_train.", False),
        ("These are AI-suggested classifications only. Nothing here has been", False),
        ("published or sent to any user, and nothing here is a final gold label", False),
        ("until a human reviewer confirms or corrects it.", False),
        ("", False),
        ("Column order (left to right, visible columns):", True),
        ("  1. comment_text — the original visitor comment", False),
        ("  2. ai_suggested_primary_intent — AI's suggested intent (draft)", False),
        ("  3. ai_suggested_topic_subtype — AI's suggested topic (draft)", False),
        ("  4. ai_suggested_sentiment — AI's suggested sentiment (draft)", False),
        ("  5. ai_suggested_confidence — AI's self-reported confidence (0-1)", False),
        ("  6-11. gold_* fields — fill these in to CONFIRM or CORRECT the AI", False),
        ("     suggestion. Leave gold_annotation_status blank until you decide.", False),
        ("  12-14. labeler_notes / labeled_by / labeled_at — optional reviewer notes", False),
        ("", False),
        ("Hidden columns (present, not deleted): comment_id, split,", False),
        ("annotation_stage, is_few_shot_example, ai_suggested_secondary_intent,", False),
        ("ai_suggested_subcategory, ai_suggestion_model. Unhide via Excel/Numbers", False),
        ("if you need them; do not delete or reorder them.", False),
        ("", False),
        ("Row order: sorted by ai_suggested_confidence ascending (lowest", False),
        ("confidence first), so the cases most likely to need correction are", False),
        ("at the top. This sort is display-only and has no effect on the", False),
        ("underlying data file.", False),
        ("", False),
        ("Color coding:", True),
        ("  Amber row = ai_suggested_confidence < 0.85 (reporting-only flag,", False),
        ("  not a hard threshold — still review normally).", False),
        ("  Orange row = a known pre-existing anomaly: this row already had a", False),
        ("  gold_* label filled in before Phase B.2 ran (predates this AI run).", False),
        ("  Its AI suggestion disagrees with that old gold_* label — flagged", False),
        ("  for extra attention.", False),
        ("", False),
        ("IMPORTANT — this workbook does NOT write anything back automatically.", True),
        ("AI suggestions are never copied into gold_* by any script. A human", False),
        ("must explicitly decide and enter each gold_* value. This workbook is", False),
        ("for Phase B.3 review only; Phase B.3 itself has not started.", False),
    ]
    for i, (text, bold) in enumerate(lines, start=1):
        cell = ws_i.cell(row=i, column=1, value=text)
        if bold:
            cell.font = Font(bold=True, size=13 if i == 1 else 11)
    ws_i.column_dimensions["A"].width = 100

    # --- Review sheet ---------------------------------------------------------
    ws = wb.create_sheet("Review")
    for col_idx, col_name in enumerate(ALL_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    ws.freeze_panes = "A2"

    anomaly_ids = {"00bb4fc8146ff830b88709f0ddab7a164cd18312461b35c13e8f9d621c18089d"}

    for r_idx, (_, row) in enumerate(review.iterrows(), start=2):
        for c_idx, col_name in enumerate(ALL_COLUMNS, start=1):
            val = row[col_name]
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            if col_name == "comment_text":
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            if col_name in REVIEWER_EDITABLE_COLUMNS:
                cell.protection = Protection(locked=False)
        try:
            conf = float(row["ai_suggested_confidence"])
        except ValueError:
            conf = None
        row_fill = None
        if row["comment_id"] in anomaly_ids:
            row_fill = ANOMALY_FILL
        elif conf is not None and conf < 0.85:
            row_fill = LOW_CONF_FILL
        if row_fill is not None:
            for c_idx in range(1, len(ALL_COLUMNS) + 1):
                ws.cell(row=r_idx, column=c_idx).fill = row_fill

    n_rows = len(review)
    last_row = n_rows + 1

    # Column widths (visible columns only, reasonable defaults)
    width_map = {
        "comment_text": 60,
        "ai_suggested_primary_intent": 20,
        "ai_suggested_topic_subtype": 20,
        "ai_suggested_sentiment": 14,
        "ai_suggested_confidence": 12,
        "gold_primary_intent": 20,
        "gold_topic_subtype": 20,
        "gold_secondary_intent": 18,
        "gold_subcategory": 16,
        "gold_sentiment": 14,
        "gold_annotation_status": 18,
        "labeler_notes": 30,
        "labeled_by": 14,
        "labeled_at": 14,
    }
    for col_idx, col_name in enumerate(ALL_COLUMNS, start=1):
        letter = get_column_letter(col_idx)
        if col_name in width_map:
            ws.column_dimensions[letter].width = width_map[col_name]
        if col_name in HIDDEN_COLUMNS:
            ws.column_dimensions[letter].hidden = True

    # Dropdown validations on the gold_* columns a reviewer fills in.
    def _col_letter(name: str) -> str:
        return get_column_letter(ALL_COLUMNS.index(name) + 1)

    def _add_dv(col_name: str, options: list[str], allow_blank: bool = True) -> None:
        formula = '"' + ",".join(options) + '"'
        dv = DataValidation(type="list", formula1=formula, allow_blank=allow_blank)
        ws.add_data_validation(dv)
        letter = _col_letter(col_name)
        dv.add(f"{letter}2:{letter}{last_row}")

    _add_dv("gold_primary_intent", list(settings.PRIMARY_INTENTS))
    _add_dv("gold_topic_subtype", list(settings.TOPIC_SUBTYPES))
    _add_dv("gold_secondary_intent", list(settings.PRIMARY_INTENTS))
    _add_dv("gold_subcategory", list(settings.SUBCATEGORIES))
    _add_dv("gold_sentiment", list(settings.SENTIMENTS))
    _add_dv("gold_annotation_status", ["LABELED", "UNCERTAIN", "INSUFFICIENT_CONTEXT"])

    wb.save(out_path)
    return out_path


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        if value.time() == dt.time():
            return value.date().isoformat()
        return value.isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, bool):
        return "True" if value else "False"
    return str(value)


def _read_review_sheet(workbook_path: Path) -> pd.DataFrame:
    wb = load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        if "Review" not in wb.sheetnames:
            raise AssertionError("reviewed workbook is missing the Review sheet")
        ws = wb["Review"]
        rows = list(ws.iter_rows())
        if not rows:
            raise AssertionError("Review sheet is empty")
        if any(cell.data_type == "f" for row in rows for cell in row):
            raise AssertionError("formulas are not allowed in the Review sheet")

        headers = [_cell_text(cell.value) for cell in rows[0]]
        if headers != ALL_COLUMNS:
            raise AssertionError(
                "Review sheet columns were added, removed, or reordered"
            )
        values = [
            [_cell_text(cell.value) for cell in row]
            for row in rows[1:]
        ]
        return pd.DataFrame(values, columns=headers)
    finally:
        wb.close()


def prepare_reviewed_workbook_import(
    reviewed_workbook_path: Path,
    gold_csv_path: Path = GOLD_CSV,
    *,
    require_b3_complete: bool = True,
) -> pd.DataFrame:
    """
    Return a validated merged frame without writing it.

    Only reviewer-editable gold/trace fields can come from the workbook.
    Every other workbook value must exactly match the canonical CSV.
    """
    canonical = pd.read_csv(gold_csv_path, dtype=str, keep_default_na=False)
    verify_post_b2_b3_invariants(canonical, require_b3_complete=False)
    reviewed = _read_review_sheet(reviewed_workbook_path)

    if reviewed["comment_id"].duplicated().any():
        duplicate_ids = reviewed.loc[
            reviewed["comment_id"].duplicated(), "comment_id"
        ].tolist()
        raise AssertionError(
            f"Review sheet contains duplicate comment_id(s): {duplicate_ids[:5]}"
        )

    targets = canonical[canonical["annotation_stage"] == "ai_assisted_train"]
    expected_ids = set(targets["comment_id"])
    reviewed_ids = set(reviewed["comment_id"])
    if reviewed_ids != expected_ids:
        missing = sorted(expected_ids - reviewed_ids)
        unexpected = sorted(reviewed_ids - expected_ids)
        raise AssertionError(
            "Review sheet row set differs from canonical ai_assisted_train rows "
            f"(missing={missing[:5]}, unexpected={unexpected[:5]})"
        )

    canonical_targets = targets.set_index("comment_id")
    reviewed_targets = reviewed.set_index("comment_id")
    # comment_id is protected by the exact-set and duplicate checks above;
    # after indexing it is no longer a regular comparison column.
    for col in (c for c in PROTECTED_WORKBOOK_COLUMNS if c != "comment_id"):
        expected = canonical_targets.loc[reviewed_targets.index, col].astype(str)
        actual = reviewed_targets[col].astype(str)
        changed = actual != expected
        if changed.any():
            changed_ids = actual.index[changed].tolist()
            raise AssertionError(
                f"protected column {col!r} changed in reviewed workbook: "
                f"{changed_ids[:5]}"
            )

    # The one documented pre-B.2 gold anomaly may remain trace-less only while
    # its reviewer-editable values remain byte-identical to the canonical CSV.
    # Any B.3 change to it must acquire the same human trace as every other row.
    for cid in PREEXISTING_AI_ASSISTED_GOLD_IDS & reviewed_ids:
        old = canonical_targets.loc[cid, REVIEWER_EDITABLE_COLUMNS].astype(str)
        new = reviewed_targets.loc[cid, REVIEWER_EDITABLE_COLUMNS].astype(str)
        if not old.equals(new):
            if not new["labeled_by"].strip() or not new["labeled_at"].strip():
                raise AssertionError(
                    f"comment_id={cid}: changing the pre-existing anomaly "
                    "requires labeled_by and labeled_at"
                )

    merged = canonical.copy()
    canonical_index = merged.set_index("comment_id")
    for cid, row in reviewed_targets.iterrows():
        for col in REVIEWER_EDITABLE_COLUMNS:
            canonical_index.at[cid, col] = row[col]
    merged = canonical_index.loc[canonical["comment_id"]].reset_index()
    merged = merged[canonical.columns]

    # Independent post-merge checks: protected stages must remain untouched,
    # all B.2 suggestions must remain valid, and all imported human labels
    # must satisfy the gold-label business rules and trace requirements.
    verify_post_b2_b3_invariants(
        merged,
        require_b3_complete=require_b3_complete,
    )

    noneditable = [col for col in canonical.columns if col not in REVIEWER_EDITABLE_COLUMNS]
    if not canonical[noneditable].equals(merged[noneditable]):
        raise AssertionError("import changed a non-reviewer-editable canonical value")
    protected_rows = canonical["annotation_stage"] != "ai_assisted_train"
    if not canonical.loc[protected_rows].equals(merged.loc[protected_rows]):
        raise AssertionError("import changed a protected human_val/human_test/seed row")
    return merged


def import_reviewed_workbook(
    reviewed_workbook_path: Path,
    output_csv_path: Path,
    gold_csv_path: Path = GOLD_CSV,
    *,
    require_b3_complete: bool = True,
) -> Path:
    """Validate completely, then atomically write the requested CSV output."""
    merged = prepare_reviewed_workbook_import(
        reviewed_workbook_path,
        gold_csv_path,
        require_b3_complete=require_b3_complete,
    )
    output_csv_path = Path(output_csv_path)
    if not output_csv_path.parent.exists():
        raise FileNotFoundError(
            f"output directory does not exist: {output_csv_path.parent}"
        )

    temp_name = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="w",
            encoding="utf-8-sig",
            newline="",
            dir=output_csv_path.parent,
            prefix=f".{output_csv_path.name}.",
            suffix=".tmp",
            delete=False,
        ) as tmp:
            temp_name = tmp.name
            merged.to_csv(tmp, index=False)
        os.replace(temp_name, output_csv_path)
    finally:
        if temp_name and os.path.exists(temp_name):
            os.unlink(temp_name)
    return output_csv_path


def prepare_calibration_human_import(
    gold_csv_path: Path = GOLD_CSV,
    calibration_review_path: Path = CALIBRATION_REVIEW_JSON,
) -> pd.DataFrame:
    labels = calibration_human_labels_by_id(calibration_review_path)
    canonical = pd.read_csv(gold_csv_path, dtype=str, keep_default_na=False)
    return merge_partial_reviewer_labels(
        canonical,
        labels,
        require_empty_existing_gold=True,
    )


def import_calibration_human_labels(
    output_csv_path: Path,
    gold_csv_path: Path = GOLD_CSV,
    calibration_review_path: Path = CALIBRATION_REVIEW_JSON,
) -> Path:
    """Import completed B.2 calibration human labels into the canonical gold CSV."""
    merged = prepare_calibration_human_import(gold_csv_path, calibration_review_path)
    output_csv_path = Path(output_csv_path)
    if not output_csv_path.parent.exists():
        raise FileNotFoundError(
            f"output directory does not exist: {output_csv_path.parent}"
        )

    temp_name = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="w",
            encoding="utf-8-sig",
            newline="",
            dir=output_csv_path.parent,
            prefix=f".{output_csv_path.name}.",
            suffix=".tmp",
            delete=False,
        ) as tmp:
            temp_name = tmp.name
            merged.to_csv(tmp, index=False)
        os.replace(temp_name, output_csv_path)
    finally:
        if temp_name and os.path.exists(temp_name):
            os.unlink(temp_name)
    return output_csv_path


def import_b3_wizard_human_labels(
    output_csv_path: Path,
    gold_csv_path: Path = GOLD_CSV,
    review_path: Path = B3_REVIEW_JSON,
) -> Path:
    """Import completed B.3 wizard human labels into the canonical gold CSV."""
    return import_b3_wizard_labels(output_csv_path, gold_csv_path, review_path)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Build the B.3 review workbook or strictly import a reviewed copy."
    )
    parser.add_argument(
        "--import-reviewed",
        type=Path,
        help="Reviewed workbook to validate and import (never copies AI suggestions into gold fields).",
    )
    parser.add_argument(
        "--output-csv",
        type=Path,
        help="Explicit output path for --import-reviewed; required in import mode.",
    )
    parser.add_argument(
        "--allow-incomplete",
        action="store_true",
        help="Allow unreviewed target rows during an intentional partial B.3 import.",
    )
    parser.add_argument(
        "--import-calibration",
        action="store_true",
        help="Import completed B.2 calibration human labels into the canonical gold CSV.",
    )
    parser.add_argument(
        "--remaining-only",
        action="store_true",
        help="Build a B.3 workbook excluding rows already human-reviewed in calibration.",
    )
    parser.add_argument(
        "--import-b3-wizard",
        action="store_true",
        help="Import completed B.3 wizard human labels into the canonical gold CSV.",
    )
    parser.add_argument(
        "--b3-review",
        type=Path,
        default=B3_REVIEW_JSON,
        help="Path to b3_human_review.json for wizard import.",
    )
    parser.add_argument(
        "--calibration-review",
        type=Path,
        default=CALIBRATION_REVIEW_JSON,
        help="Path to b2_calibration_human_review.json for calibration import/exclusion.",
    )
    args = parser.parse_args()

    if args.import_calibration:
        if args.output_csv is None:
            parser.error("--output-csv is required with --import-calibration")
        path = import_calibration_human_labels(
            args.output_csv,
            calibration_review_path=args.calibration_review,
        )
    elif args.import_b3_wizard:
        if args.output_csv is None:
            parser.error("--output-csv is required with --import-b3-wizard")
        path = import_b3_wizard_human_labels(
            args.output_csv,
            review_path=args.b3_review,
        )
    elif args.import_reviewed:
        if args.output_csv is None:
            parser.error("--output-csv is required with --import-reviewed")
        path = import_reviewed_workbook(
            args.import_reviewed,
            args.output_csv,
            require_b3_complete=not args.allow_incomplete,
        )
    else:
        if args.output_csv is not None or args.allow_incomplete:
            parser.error("--output-csv/--allow-incomplete require an import mode")
        if args.remaining_only:
            exclude_ids = completed_calibration_comment_ids(args.calibration_review)
            path = build(
                exclude_comment_ids=exclude_ids,
                out_path=OUT_REMAINING_XLSX,
                instructions_title=(
                    "Phase B.3 Reviewer Workbook — Remaining ai_assisted_train Rows "
                    "(calibration-complete rows excluded)"
                ),
            )
        else:
            path = build()
    print(f"wrote {path}")
