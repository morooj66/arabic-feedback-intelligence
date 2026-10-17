"""Build the 75-record BLIND human-evaluation sample. LOCAL ONLY. No labels, no predictions, no external call.

Universe: the unlabeled "fresh candidate pool" from the sourcing audits (_scratch/fresh_candidate_pool_round1.json), the only
never-labelled pool in the project (no artifact literally named "Gold Test v2" exists). Everything already used is removed:
  * every development row (train/val) and the sealed human_test ids (ID manifest only; sealed texts are never retained or opened);
  * the round-1 / round-2 candidate rows that were already labelled;
  * rows the production pipeline would not classify (prefilter junk, exact duplicates);
  * text duplicates of development or candidate texts;
  * rows the privacy policy cannot sanitise (quarantined) and rows that are too short / too long to label.
Nothing about a record's labels or any model output is used to choose it. The targeted component uses only declared lexical cues.

Selection (seed fixed): NATURAL = simple random draw from the eligible pool; TARGETED = cue-enriched draws for rare classes
from the remaining eligible rows. The reviewer file shows only an anonymous id and the sanitised text; the component / stratum /
original id live in a separate coordinator-only mapping.

Run:  python scripts/offline_eval/blind_eval/build_sample.py
"""

from __future__ import annotations

import csv
import datetime as dt
import hashlib
import json
import random
import re
import sys
import unicodedata
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts" / "offline_eval" / "jev_dev_validation"))

SEED = 20261016
N_NATURAL = 45
TARGETED_PLAN = (("technical_cues", 8), ("complaint_cues", 8), ("entry_error_cues", 8), ("feature_request_cues", 6))
N_TOTAL = N_NATURAL + sum(n for _, n in TARGETED_PLAN)  # 75
MIN_CHARS, MAX_CHARS, MIN_WORDS = 15, 1200, 2

# Declared before sampling; lexical enrichment cues only (they surface candidates, they never label them).
CUES = {
    "technical_cues": ["لا يعمل", "لا تعمل", "لا يفتح", "لا تفتح", "لا يظهر", "لا تظهر", "ما يفتح", "ما يعمل", "الموقع لا", "الصفحة لا",
                       "التطبيق", "تسجيل الدخول", "الرابط", "تحميل", "يتعطل", "معطل", "error"],
    "complaint_cues": ["للأسف", "سيء", "سيئة", "إهمال", "اهمال", "تأخر", "تأخير", "لا أحد يرد", "ما حد", "غير مقبول", "مزعج", "استياء",
                       "شكوى", "أشتكي", "اشتكي", "مو معقول", "سخيف"],
    "entry_error_cues": ["خطأ", "غلط", "غير صحيح", "غير صحيحة", "تصويب", "الصواب", "مكرر", "تكرار", "ناقص", "لا يتناسب", "لا معنى"],
    "feature_request_cues": ["أتمنى أن", "اتمنى ان", "أتمنى إضافة", "خاصية", "ميزة", "تضيفون", "تطبيق للجوال", "نسخة", "pdf", "تحسين", "إتاحة", "قسم جديد"],
}


def norm_text(t: str) -> str:
    return " ".join(unicodedata.normalize("NFKC", str(t or "")).split()).lower()


def sid_of(source_file: str, source_row_index: int) -> str:
    return hashlib.sha256(f"{source_file}|{int(source_row_index)}".encode("utf-8")).hexdigest()


# ------------------------------------------------------------------------------ pure selection (tested)


def select_sample(eligible: list[dict], seed: int = SEED, n_natural: int = N_NATURAL, plan=TARGETED_PLAN) -> dict:
    """eligible: [{sid, text, ...}] with unique sid. Returns {"records": [...], "strata": {...}}; deterministic in `seed`.

    Natural first (uniform over ALL eligible), then targeted from what is left, in the declared stratum order; a record that matches
    several strata is assigned to the first. A stratum short of its target is reported, then topped up from strata with surplus in
    declared order (never from the natural pool), so counts and the shortfall are always explicit."""
    pool = sorted(eligible, key=lambda r: r["sid"])
    rng_nat = random.Random(f"{seed}-natural")
    natural = rng_nat.sample(pool, min(n_natural, len(pool)))
    taken = {r["sid"] for r in natural}
    rest = [r for r in pool if r["sid"] not in taken]
    members: dict[str, list[dict]] = {name: [] for name, _ in plan}
    for r in rest:
        t = r["text"]
        for name, _ in plan:
            if any(c.lower() in t.lower() for c in CUES[name]):
                members[name].append(r)
                break
    rng_t = random.Random(f"{seed}-targeted")
    chosen: dict[str, list[dict]] = {}
    for name, target in plan:
        m = members[name]
        chosen[name] = rng_t.sample(m, min(target, len(m)))
    short = {name: target - len(chosen[name]) for name, target in plan if len(chosen[name]) < target}
    redistributed = []
    for deficit_name, deficit in list(short.items()):
        for name, _ in plan:
            while deficit > 0 and name != deficit_name:
                left = [r for r in members[name] if r not in chosen[name]]
                if not left:
                    break
                pick = rng_t.choice(left)
                chosen[name].append(pick)
                redistributed.append((deficit_name, name, pick["sid"]))
                deficit -= 1
    records = [dict(r, component="natural", stratum="natural_random") for r in natural]
    for name, _ in plan:
        records += [dict(r, component="targeted", stratum=name) for r in chosen[name]]
    rng_order = random.Random(f"{seed}-order")
    rng_order.shuffle(records)
    for i, r in enumerate(records, 1):
        r["record_id"] = f"BE-{i:03d}"
    return {"records": records, "pool_sizes": {name: len(members[name]) for name, _ in plan}, "natural_pool": len(pool),
            "shortfall": short, "redistributed": redistributed}


# ----------------------------------------------------------------------------------- pool construction


def build_universe(log: dict) -> list[dict]:
    import pandas as pd
    from config import settings
    from rcs.ingest.cleaning import clean
    from rcs.ingest.loader import load_raw
    from rcs.prefilter.spam_rules import BurstConfig, apply_prefilter
    import jev_name_policy as NP

    raw_path = ROOT / "data" / "raw" / "historical-export.xlsx"
    raw = load_raw(raw_path, settings.SOURCE_COLUMNS)
    pool = json.loads((ROOT / "_scratch" / "fresh_candidate_pool_round1.json").read_text(encoding="utf-8"))
    pool_keys = {(r["source_file"], int(r["source_row_index"])) for r in pool}
    log["raw_rows"], log["fresh_pool_rows"] = len(raw), len(pool)

    # exclusions by id (the sealed human_test manifest is IDs only) + dev/candidate texts (never human_test texts)
    sealed = set()
    def walk(o):
        if isinstance(o, str) and len(o) == 64: sealed.add(o)
        elif isinstance(o, list): [walk(x) for x in o]
        elif isinstance(o, dict): [walk(v) for v in o.values()]
    walk(json.loads((ROOT / "data" / "gold" / "human_test_exclusion_ids.json").read_text()))
    gold_ids, dev_texts = set(), set()
    with open(ROOT / "data" / "gold" / "gold_dataset_round2_expanded.csv", newline="", encoding="utf-8-sig") as fh:
        for row in csv.DictReader(fh):          # streamed; rows of the sealed split are dropped immediately, never retained
            if row.get("split") == "test":
                gold_ids.add(row["comment_id"])
                continue
            gold_ids.add(row["comment_id"])
            dev_texts.add(norm_text(row["comment_text"]))
    cand_ids = set()
    for f in ("candidate_sourcing_audit_round1.csv", "candidate_sourcing_audit_round2.csv"):
        with open(ROOT / "data" / "gold" / f, newline="", encoding="utf-8-sig") as fh:
            cand_ids |= {r["comment_id"] for r in csv.DictReader(fh)}
    log["excluded_ids"] = {"human_test_manifest": len(sealed), "gold_dataset_ids": len(gold_ids), "round1_round2_candidates": len(cand_ids)}
    excl_ids = sealed | gold_ids | cand_ids
    cand_texts = {norm_text(t) for t, cid in zip(raw["original_comment"], raw["comment_id"]) if cid in cand_ids}

    df = clean(raw, dt.datetime(2026, 10, 16), settings.AGE_RECENT_MAX_DAYS, settings.AGE_DELAYED_MAX_DAYS)
    df = apply_prefilter(df, settings.KNOWN_JUNK_STRINGS, BurstConfig(min_count=settings.BURST_MIN_COUNT, window_hours=settings.BURST_WINDOW_HOURS))
    df = df[[(f, int(i)) in pool_keys for f, i in zip(df["source_file"], df["source_row_index"])]]
    log["after_pool_membership"] = len(df)
    steps = Counter()
    out, seen = [], set()
    for r in df.itertuples():
        if r.comment_id in excl_ids:
            steps["excluded_already_used_or_sealed_id"] += 1; continue
        if bool(r.is_prefiltered_junk):
            steps["excluded_prefilter_junk"] += 1; continue
        if r.is_exact_duplicate_of is not None and str(r.is_exact_duplicate_of) not in ("", "nan", "None"):
            steps["excluded_exact_duplicate"] += 1; continue
        n = norm_text(r.original_comment)
        if n in dev_texts or n in cand_texts:
            steps["excluded_text_duplicate_of_dev_or_candidate"] += 1; continue
        res = NP.sanitize(str(r.original_comment))
        if res.text is None:
            steps[f"excluded_privacy_policy:{res.quarantine_reason}"] += 1; continue
        txt = res.text
        if len(txt) < MIN_CHARS or len(txt) > MAX_CHARS or len(txt.split()) < MIN_WORDS:
            steps["excluded_length_or_too_few_words"] += 1; continue
        key = norm_text(txt)
        if key in seen:
            steps["excluded_duplicate_within_pool"] += 1; continue
        seen.add(key)
        out.append({"sid": r.comment_id, "source_file": r.source_file, "source_row_index": int(r.source_row_index), "text": txt,
                    "masks": sum(v for k, v in res.actions.items() if k in NP.IDENTITY_MASK_ACTIONS),
                    "lemma_present": str(getattr(r, "lemma", "")).strip() not in ("", "nan", "None")})
    log["exclusion_steps"] = dict(steps)
    log["eligible"] = len(out)
    return out


# ------------------------------------------------------------------------------------------ outputs


def write_outputs(sample: dict, log: dict, out_dir: Path) -> dict:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    import blind_guidelines as G

    rev = out_dir / "reviewer"; coord = out_dir / "coordinator_only"
    rev.mkdir(parents=True, exist_ok=True); coord.mkdir(parents=True, exist_ok=True)
    recs = sorted(sample["records"], key=lambda r: r["record_id"])

    wb = Workbook()
    ws0 = wb.active; ws0.title = "Read me first"
    for i, line in enumerate(G.README_LINES, 1):
        ws0.cell(row=i, column=1, value=line).alignment = Alignment(wrap_text=True, vertical="top")
    ws0.column_dimensions["A"].width = 120
    ws0["A1"].font = Font(bold=True, size=14)
    wsg = wb.create_sheet("Guidelines")
    for i, (kind, a, b) in enumerate(G.GUIDELINE_ROWS, 1):
        wsg.cell(row=i, column=1, value=a).alignment = Alignment(wrap_text=True, vertical="top")
        wsg.cell(row=i, column=2, value=b).alignment = Alignment(wrap_text=True, vertical="top")
        if kind == "h":
            for c in (1, 2):
                wsg.cell(row=i, column=c).font = Font(bold=True); wsg.cell(row=i, column=c).fill = PatternFill("solid", fgColor="DDEBDD")
    wsg.column_dimensions["A"].width = 34; wsg.column_dimensions["B"].width = 110
    wsr = wb.create_sheet("Review")
    wsr.sheet_view.rightToLeft = True
    headers = ["record_id", "comment_text", "intent", "topic", "sentiment", "insufficient_context", "reviewer_notes"]
    for c, h in enumerate(headers, 1):
        cell = wsr.cell(row=1, column=c, value=h); cell.font = Font(bold=True); cell.fill = PatternFill("solid", fgColor="DDEBDD")
    for r, rec in enumerate(recs, 2):
        wsr.cell(row=r, column=1, value=rec["record_id"])
        t = wsr.cell(row=r, column=2, value=rec["text"]); t.alignment = Alignment(wrap_text=True, vertical="top", readingOrder=2)
    last = len(recs) + 1
    for col, vals in (("C", G.INTENTS), ("D", G.TOPICS), ("E", G.SENTIMENTS), ("F", ["INSUFFICIENT_CONTEXT"])):
        dv = DataValidation(type="list", formula1='"' + ",".join(vals) + '"', allow_blank=True, showErrorMessage=True,
                            errorTitle="Not a valid code", error="Choose one of the listed codes (see the Guidelines sheet).")
        wsr.add_data_validation(dv); dv.add(f"{col}2:{col}{last}")
    for col, w in zip("ABCDEFG", (11, 90, 18, 26, 14, 22, 40)):
        wsr.column_dimensions[col].width = w
    wsr.freeze_panes = "C2"
    wsc = wb.create_sheet("Codes")
    for i, (code, ar, en) in enumerate(G.CODE_TABLE, 1):
        wsc.cell(row=i, column=1, value=code); wsc.cell(row=i, column=2, value=ar); wsc.cell(row=i, column=3, value=en)
    wsc.column_dimensions["A"].width = 28; wsc.column_dimensions["B"].width = 40; wsc.column_dimensions["C"].width = 70
    path = rev / "blind_eval_75_review.xlsx"
    wb.save(path)

    mapping = coord / "blind_eval_mapping.local.csv"
    with open(mapping, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["record_id", "comment_id", "source_file", "source_row_index", "component", "stratum", "lemma_present", "identity_masks_applied", "reviewer_text_sha256"])
        for rec in recs:
            w.writerow([rec["record_id"], rec["sid"], rec["source_file"], rec["source_row_index"], rec["component"], rec["stratum"],
                        int(rec["lemma_present"]), rec["masks"], hashlib.sha256(rec["text"].encode("utf-8")).hexdigest()])
    ids_digest = hashlib.sha256("\n".join(f"{r['record_id']}|{r['sid']}" for r in recs).encode()).hexdigest()
    texts_digest = hashlib.sha256("\n".join(f"{r['record_id']}|{r['text']}" for r in recs).encode("utf-8")).hexdigest()
    comp = Counter(r["component"] for r in recs); strata = Counter(r["stratum"] for r in recs)
    summary = {"seed": SEED, "n_total": len(recs), "components": dict(comp), "strata": dict(strata), "pool_sizes_for_targeted": sample["pool_sizes"],
               "natural_pool_size": sample["natural_pool"], "targeted_shortfall": sample["shortfall"], "redistributed": [list(x) for x in sample["redistributed"]],
               "sample_ids_sha256": ids_digest, "reviewer_texts_sha256": texts_digest, "pipeline_log": log,
               "records_with_identity_mask": sum(1 for r in recs if r["masks"]), "labels_prefilled": False, "model_predictions_used": False,
               "external_calls": 0}
    (coord / "sampling_summary.json").write_text(json.dumps(summary, indent=1))
    lines = ["# Blind evaluation sample v1 - COORDINATOR ONLY (never give this folder to reviewers)", "",
             f"Seed `{SEED}`; reproduce with `python scripts/offline_eval/blind_eval/build_sample.py`.", "",
             f"Total {len(recs)}: natural {comp['natural']} (uniform random from the eligible pool of {sample['natural_pool']}), targeted {comp['targeted']}.", "",
             "| stratum | target | actual | pool size (cue matches among remaining eligible) |", "|---|---|---|---|"]
    for name, target in TARGETED_PLAN:
        lines.append(f"| {name} | {target} | {strata.get(name, 0)} | {sample['pool_sizes'][name]} |")
    lines += ["", f"Shortfall before top-up: {sample['shortfall'] or 'none'}; top-ups: {len(sample['redistributed'])}.", "",
              "## Universe and exclusions", "", "```", json.dumps(log, indent=1), "```", "",
              "## Integrity", f"- sample id list sha256: `{ids_digest}`", f"- reviewer text sha256: `{texts_digest}`", "",
              "## What this sample is NOT", "- Not labelled. No label, prediction or model output was used to select it or is stored in it.",
              "- The natural component represents the ELIGIBLE pool (privacy-sanitisable, 15-1200 chars, non-junk, non-duplicate), not all traffic.",
              "- The targeted component is cue-enriched: use it only for rare-class diagnostics, never for an overall accuracy estimate.", "",
              "## After labelling", "Use scripts/offline_eval/blind_eval/evaluate_blind.py with the returned workbook and blind_eval_mapping.local.csv."]
    (coord / "SAMPLING_SUMMARY.md").write_text("\n".join(lines), encoding="utf-8")
    for p in (out_dir, rev, coord):
        p.chmod(0o700)
    for p in (path, mapping, coord / "sampling_summary.json", coord / "SAMPLING_SUMMARY.md"):
        p.chmod(0o600)
    return summary


def main() -> None:
    log: dict = {}
    eligible = build_universe(log)
    sample = select_sample(eligible)
    assert len(sample["records"]) == N_TOTAL, f"expected {N_TOTAL}, got {len(sample['records'])}"
    summary = write_outputs(sample, log, ROOT / "data" / "evaluation" / "blind_eval_v1")
    print(json.dumps({k: summary[k] for k in ("n_total", "components", "strata", "pool_sizes_for_targeted", "natural_pool_size",
                                              "targeted_shortfall", "records_with_identity_mask", "sample_ids_sha256")}, indent=1))
    print(json.dumps(log["exclusion_steps"], indent=1))


if __name__ == "__main__":
    main()
