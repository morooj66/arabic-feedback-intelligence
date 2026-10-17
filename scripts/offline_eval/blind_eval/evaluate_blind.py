"""Evaluation harness for the 75-record BLIND human evaluation. LOCAL ONLY (no network). Run it AFTER human labels exist.

    python scripts/offline_eval/blind_eval/evaluate_blind.py --labels <returned workbook.xlsx> [--second-labels <2nd reviewer.xlsx>]
                                                              [--jev-predictions jev_v1_predictions.jsonl]

Order of operations (enforced): 1) validate the returned labels against the frozen taxonomy and the sample; 2) only then compute
predictions. Compared systems: frozen RCS; RCS + opt-in WS/EC boundary layer (the candidate); optionally Jev v1 and Jev v1 + layer
when a predictions file is supplied (this script never calls Jev).

TWO RESULT FAMILIES THAT ARE NEVER COMBINED:
  A. NATURAL-DISTRIBUTION results: the 45 randomly drawn records only. These are the only numbers that may be read as accuracy on
     (the eligible part of) real traffic, with wide intervals because n is small.
  B. RARE-CLASS DIAGNOSTICS: the 30 cue-enriched records only. They show whether rare classes are detected (recall, false positives,
     boundary confusion). They are NOT an accuracy estimate and no pooled A+B accuracy is ever computed.
Records marked INSUFFICIENT_CONTEXT are excluded from scoring and counted separately. Gold labels of the development data are not read.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import random
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts" / "offline_eval" / "jev_dev_validation"))

import jev_metrics as M  # noqa: E402
from config import settings  # noqa: E402
from rcs.classification.topic_boundary_layer import boundary_topic  # noqa: E402

BASE = ROOT / "data" / "evaluation" / "blind_eval_v1"
TASKS = ("intent", "topic", "sentiment")
INSUFFICIENT = "INSUFFICIENT_CONTEXT"
SPEC_FILE = ROOT / "data" / "evaluation" / "jev_dev_validation" / "candidate_ws_ec_layer_spec.json"
WS, EC = "WORD_SUBMISSION", "ENTRY_CORRECTION"


class LabelError(ValueError):
    pass


# ----------------------------------------------------------------------------------------------- labels


def load_labels_xlsx(path: Path) -> dict[str, dict]:
    from openpyxl import load_workbook

    ws = load_workbook(path, data_only=True)["Review"]
    head = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    need = ["record_id", "comment_text", "intent", "topic", "sentiment", "insufficient_context", "reviewer_notes"]
    if head[: len(need)] != need:
        raise LabelError(f"unexpected Review header {head}")
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        clean = lambda v: (str(v).strip() if v is not None and str(v).strip() else None)
        out[str(row[0]).strip()] = {"text": row[1], "intent": clean(row[2]), "topic": clean(row[3]), "sentiment": clean(row[4]),
                                    "insufficient": clean(row[5]), "notes": clean(row[6])}
    return out


def validate_labels(labels: dict[str, dict], expected_ids: list[str]) -> dict:
    """Raises LabelError listing every problem (ids only). Returns {"scored": {id: {...}}, "insufficient": [ids]}."""
    problems = []
    missing = [i for i in expected_ids if i not in labels]
    extra = [i for i in labels if i not in set(expected_ids)]
    if missing:
        problems.append(f"missing records: {missing}")
    if extra:
        problems.append(f"unknown records: {extra}")
    allowed = {"intent": set(settings.PRIMARY_INTENTS), "topic": set(settings.TOPIC_SUBTYPES), "sentiment": set(settings.SENTIMENTS)}
    scored, insufficient = {}, []
    for i in expected_ids:
        r = labels.get(i)
        if r is None:
            continue
        if r["insufficient"] is not None:
            if r["insufficient"] != INSUFFICIENT:
                problems.append(f"{i}: insufficient_context must be {INSUFFICIENT} or empty")
            elif any(r[t] is not None for t in TASKS):
                problems.append(f"{i}: INSUFFICIENT_CONTEXT rows must leave intent/topic/sentiment empty")
            else:
                insufficient.append(i)
            continue
        for t in TASKS:
            if r[t] is None:
                problems.append(f"{i}: {t} is empty")
            elif r[t] not in allowed[t]:
                problems.append(f"{i}: {t} value {r[t]!r} is not in the frozen taxonomy")
        if all(r[t] in allowed[t] for t in TASKS):
            scored[i] = {t: r[t] for t in TASKS}
    if problems:
        raise LabelError("labels not ready:\n  - " + "\n  - ".join(problems))
    return {"scored": scored, "insufficient": insufficient}


def agreement(a: dict, b: dict) -> dict:
    """Two reviewers on the same records (scored by both): raw agreement and Cohen's kappa per task; disagreements listed by id."""
    ids = sorted(set(a) & set(b))
    out = {"n": len(ids), "tasks": {}, "disagreements": {}}
    for t in TASKS:
        la = [a[i][t] for i in ids]; lb = [b[i][t] for i in ids]
        po = sum(x == y for x, y in zip(la, lb)) / len(ids) if ids else 0.0
        cats = sorted(set(la) | set(lb))
        pe = sum((la.count(c) / len(ids)) * (lb.count(c) / len(ids)) for c in cats) if ids else 0.0
        out["tasks"][t] = {"raw_agreement": po, "cohen_kappa": (po - pe) / (1 - pe) if pe < 1 else 1.0}
        out["disagreements"][t] = [i for i, x, y in zip(ids, la, lb) if x != y]
    return out


# ----------------------------------------------------------------------------------------------- stats


def wilson(k, n, z=1.96):
    return M.wilson(k, n, z)


def mcnemar_exact_p(b: int, c: int) -> float:
    """Two-sided exact binomial p for discordant pairs (b = base right/cand wrong, c = base wrong/cand right)."""
    n = b + c
    if n == 0:
        return 1.0
    k = min(b, c)
    p = 2 * sum(math.comb(n, i) for i in range(0, k + 1)) / (2 ** n)
    return min(1.0, p)


def paired_boot(gold, a, b, labels, n=2000, seed=20261016):
    if len(gold) < 2:
        return None
    rng = random.Random(seed)
    da, df = [], []
    for _ in range(n):
        idx = [rng.randrange(len(gold)) for _ in range(len(gold))]
        g = [gold[i] for i in idx]; pa = [a[i] for i in idx]; pb = [b[i] for i in idx]
        da.append(M.accuracy(g, pb) - M.accuracy(g, pa))
        df.append(M.macro_f1(M.per_class(g, pb, labels), only_supported=True) - M.macro_f1(M.per_class(g, pa, labels), only_supported=True))
    da.sort(); df.sort()
    lo, hi = int(0.025 * n), int(0.975 * n) - 1
    return {"d_accuracy_ci95": (da[lo], da[hi]), "d_macro_f1_ci95": (df[lo], df[hi])}


# --------------------------------------------------------------------------------------------- evaluation


def with_layer(preds: dict, texts: dict, lemma: dict) -> dict:
    return {i: {**p, "topic": boundary_topic(p["topic"], texts[i], bool(lemma.get(i, False)))} for i, p in preds.items()}


def evaluate(gold: dict, systems: dict, component: dict, comparisons: list[tuple[str, str]]) -> dict:
    """gold: {id: {intent, topic, sentiment}} (scored only); systems: {name: {id: {intent, topic, sentiment}}};
    component: {id: 'natural'|'targeted'}; comparisons: [(base_name, candidate_name)]."""
    res: dict = {"A_natural": {}, "B_rare_class_diagnostics": {}, "never_pooled": True}
    ids_nat = [i for i in gold if component[i] == "natural"]
    ids_tar = [i for i in gold if component[i] == "targeted"]
    for name, P in systems.items():
        for i in gold:
            if i not in P:
                raise LabelError(f"system {name!r} has no prediction for {i}")
    # ---- A: natural distribution
    res["A_natural"]["n_scored"] = len(ids_nat)
    res["A_natural"]["gold_support"] = {t: dict(sorted(_count([gold[i][t] for i in ids_nat]).items())) for t in TASKS}
    res["A_natural"]["systems"] = {}
    for name, P in systems.items():
        res["A_natural"]["systems"][name] = {}
        for t in TASKS:
            g = [gold[i][t] for i in ids_nat]; p = [P[i][t] for i in ids_nat]
            s = M.score_task(t, g, p) if g else None
            res["A_natural"]["systems"][name][t] = None if s is None else {
                "n": s["n"], "accuracy": s["accuracy"], "accuracy_wilson95": wilson(sum(a == b for a, b in zip(g, p)), len(g)),
                "macro_f1_supported": s["macro_f1_supported_classes"], "macro_f1_ci95": s["ci95"]["macro_f1"] if s["ci95"] else None,
                "per_class": s["per_class"]}
    res["A_natural"]["comparisons"] = {}
    for base, cand in comparisons:
        if base not in systems or cand not in systems:
            continue
        entry = {}
        for t in TASKS:
            g = [gold[i][t] for i in ids_nat]; a = [systems[base][i][t] for i in ids_nat]; b = [systems[cand][i][t] for i in ids_nat]
            reg = [i for i, x, p, q in zip(ids_nat, g, a, b) if p == x and q != x]
            gain = [i for i, x, p, q in zip(ids_nat, g, a, b) if p != x and q == x]
            entry[t] = {"regressions": len(reg), "gains": len(gain), "regression_ids": reg, "changed": sum(p != q for p, q in zip(a, b)),
                        "mcnemar_exact_p": mcnemar_exact_p(len(reg), len(gain)), "paired_bootstrap": paired_boot(g, a, b, M.TASK_LABELS[t])}
        res["A_natural"]["comparisons"][f"{base} -> {cand}"] = entry
    # WS vs EC confusion inside A
    res["A_natural"]["ws_ec_confusion"] = {name: _ws_ec(gold, P, ids_nat) for name, P in systems.items()}

    # ---- B: rare-class diagnostics on the targeted component ONLY
    res["B_rare_class_diagnostics"]["n_scored"] = len(ids_tar)
    res["B_rare_class_diagnostics"]["gold_support"] = {t: dict(sorted(_count([gold[i][t] for i in ids_tar]).items())) for t in TASKS}
    diag = {}
    for name, P in systems.items():
        d = {}
        for task, cls, label in (("intent", "COMPLAINT", "COMPLAINT"), ("topic", "TECHNICAL_ISSUE", "TECHNICAL_ISSUE"),
                                 ("topic", "DICTIONARY_SUGGESTION", "DICTIONARY_SUGGESTION"), ("topic", EC, "ENTRY_CORRECTION"),
                                 ("topic", WS, "WORD_SUBMISSION")):
            sup = [i for i in ids_tar if gold[i][task] == cls]
            tp = [i for i in sup if P[i][task] == cls]
            fp = [i for i in ids_tar if gold[i][task] != cls and P[i][task] == cls]
            nat_fp = [i for i in ids_nat if gold[i][task] != cls and P[i][task] == cls]
            nat_neg = sum(1 for i in ids_nat if gold[i][task] != cls)
            d[label] = {"support": len(sup), "recall": (len(tp) / len(sup)) if sup else None, "recall_wilson95": wilson(len(tp), len(sup)) if sup else None,
                        "false_positives_in_targeted": len(fp),
                        "false_positives_in_natural": len(nat_fp), "natural_negatives": nat_neg,
                        "natural_false_positive_rate": (len(nat_fp) / nat_neg) if nat_neg else None}
        diag[name] = d
    res["B_rare_class_diagnostics"]["systems"] = diag
    res["B_rare_class_diagnostics"]["ws_ec_confusion"] = {name: _ws_ec(gold, P, ids_tar) for name, P in systems.items()}
    res["B_rare_class_diagnostics"]["comparisons"] = {}
    for base, cand in comparisons:
        if base in systems and cand in systems:
            ent = {}
            for t in TASKS:
                g = [gold[i][t] for i in ids_tar]; a = [systems[base][i][t] for i in ids_tar]; b = [systems[cand][i][t] for i in ids_tar]
                ent[t] = {"regressions": sum(1 for x, p, q in zip(g, a, b) if p == x and q != x),
                          "gains": sum(1 for x, p, q in zip(g, a, b) if p != x and q == x)}
            res["B_rare_class_diagnostics"]["comparisons"][f"{base} -> {cand}"] = ent
    return res


def disagreement_rows(gold: dict, systems: dict, component: dict) -> list[dict]:
    """One row per (system, task, record) where the system differs from the human label. Ids and codes only (no comment text)."""
    rows = []
    for name, P in systems.items():
        for i in sorted(gold):
            for t in TASKS:
                if P[i][t] != gold[i][t]:
                    rows.append({"system": name, "component": component[i], "record_id": i, "task": t,
                                 "human_label": gold[i][t], "system_prediction": P[i][t]})
    return rows


def top_confusions(rows: list[dict], k: int = 6) -> dict:
    """{system: {component: {task: [((human, predicted), count), ...]}}} - the most frequent disagreement pairs."""
    out: dict = {}
    for r in rows:
        d = out.setdefault(r["system"], {}).setdefault(r["component"], {}).setdefault(r["task"], {})
        key = (r["human_label"], r["system_prediction"])
        d[key] = d.get(key, 0) + 1
    return {s: {c: {t: sorted(v.items(), key=lambda kv: -kv[1])[:k] for t, v in tasks.items()} for c, tasks in comps.items()} for s, comps in out.items()}


def _count(xs):
    d = {}
    for x in xs:
        d[x] = d.get(x, 0) + 1
    return d


def _ws_ec(gold, P, ids):
    rows = [i for i in ids if gold[i]["topic"] in (WS, EC)]
    m = {g: {WS: 0, EC: 0, "OTHER": 0} for g in (WS, EC)}
    for i in rows:
        m[gold[i]["topic"]][P[i]["topic"] if P[i]["topic"] in (WS, EC) else "OTHER"] += 1
    return {"n": len(rows), "matrix_rows_gold_cols_pred": m}


# ------------------------------------------------------------------------------------------------ report


def render(res: dict, meta: dict) -> str:
    L = ["# Blind evaluation report", "",
         f"Sample: `{meta['sample_ids_sha256'][:16]}…`; labels file sha256 `{meta['labels_sha256'][:16]}…`; candidate spec `{meta['candidate_spec_sha256'][:16]}…`; "
         f"reviewers: {meta['reviewers']}; INSUFFICIENT_CONTEXT excluded: {meta['insufficient']}.", "",
         "> **Two result families, never combined.** A = natural-distribution records (read as accuracy, wide intervals). "
         "B = cue-enriched rare-class records (diagnostics only; NOT an accuracy estimate). No pooled A+B accuracy exists in this report.", ""]
    if meta.get("agreement"):
        L += ["## Reviewer agreement", ""]
        for t, v in meta["agreement"]["tasks"].items():
            L.append(f"- {t}: raw {v['raw_agreement']:.2f}, Cohen's kappa {v['cohen_kappa']:.2f}; disagreements on {len(meta['agreement']['disagreements'][t])} records")
        L.append("")
    A = res["A_natural"]
    L += [f"## A. Natural-distribution results (n = {A['n_scored']} scored records)", "", "Gold support: " + json.dumps(A["gold_support"]), ""]
    L += ["| task | system | accuracy (Wilson 95%) | macro-F1 (supported classes; bootstrap 95%) |", "|---|---|---|---|"]
    for t in TASKS:
        for name, v in A["systems"].items():
            r = v[t]
            if r is None:
                continue
            ci = r["macro_f1_ci95"]
            L.append(f"| {t} | {name} | {r['accuracy']:.3f} ({r['accuracy_wilson95'][0]:.2f}-{r['accuracy_wilson95'][1]:.2f}) | {r['macro_f1_supported']:.3f}" + (f" ({ci[0]:.2f}-{ci[1]:.2f})" if ci else "") + " |")
    for t in TASKS:
        for name, v in A["systems"].items():
            r = v[t]
            if r is None:
                continue
            L += ["", f"**{t} per-class - {name}**", "", "| class | P | R | F1 | support | predicted |", "|---|---|---|---|---|---|"]
            for c, x in r["per_class"].items():
                if x["support"] or x["predicted"]:
                    L.append(f"| {c} | {x['precision']:.2f} | {x['recall']:.2f} | {x['f1']:.2f} | {x['support']} | {x['predicted']} |")
    L += ["", "### Paired comparison (candidate vs base) on the natural records", "",
          "| comparison | task | gains | regressions | rows changed | exact McNemar p | paired 95% CI: accuracy change | macro-F1 change |", "|---|---|---|---|---|---|---|---|"]
    for k, ent in A["comparisons"].items():
        for t, v in ent.items():
            pb = v["paired_bootstrap"]
            L.append(f"| {k} | {t} | {v['gains']} | {v['regressions']} | {v['changed']} | {v['mcnemar_exact_p']:.3f} | "
                     + (f"{pb['d_accuracy_ci95'][0]:+.2f}..{pb['d_accuracy_ci95'][1]:+.2f} | {pb['d_macro_f1_ci95'][0]:+.2f}..{pb['d_macro_f1_ci95'][1]:+.2f}" if pb else "n/a | n/a") + " |")
    L += ["", "### WORD_SUBMISSION vs ENTRY_CORRECTION confusion (natural records; rows = gold)", ""]
    for name, v in A["ws_ec_confusion"].items():
        L.append(f"- {name} (n={v['n']}): {json.dumps(v['matrix_rows_gold_cols_pred'])}")
    B = res["B_rare_class_diagnostics"]
    L += ["", f"## B. Rare-class diagnostics (n = {B['n_scored']} cue-enriched records; NOT an accuracy estimate)", "", "Gold support: " + json.dumps(B["gold_support"]), "",
          "| system | class | support | recall (Wilson 95%) | false positives (targeted) | false positives (natural) / natural negatives |", "|---|---|---|---|---|---|"]
    for name, d in B["systems"].items():
        for c, v in d.items():
            if v["support"] == 0 and v["false_positives_in_targeted"] == 0 and v["false_positives_in_natural"] == 0:
                continue
            rec = "n/a" if v["recall"] is None else f"{v['recall']:.2f} ({v['recall_wilson95'][0]:.2f}-{v['recall_wilson95'][1]:.2f})"
            L.append(f"| {name} | {c} | {v['support']} | {rec} | {v['false_positives_in_targeted']} | {v['false_positives_in_natural']} / {v['natural_negatives']} |")
    L += ["", "### WORD_SUBMISSION vs ENTRY_CORRECTION confusion (targeted records; rows = gold)", ""]
    for name, v in B["ws_ec_confusion"].items():
        L.append(f"- {name} (n={v['n']}): {json.dumps(v['matrix_rows_gold_cols_pred'])}")
    L += ["", "### Regressions / gains on the targeted records (counts only)", ""]
    for k, ent in B["comparisons"].items():
        L.append(f"- {k}: " + ", ".join(f"{t} +{v['gains']}/-{v['regressions']}" for t, v in ent.items()))
    L += ["", "## How to read this", "",
          "* A is the only place an accuracy number is allowed to appear; with n of about 45 every interval is wide, so a difference is only meaningful when its paired interval excludes zero.",
          "* B answers 'is the rare class found, and at what false-positive cost?'. Enrichment inflates support and false-positive counts, so B figures must not be quoted as rates.",
          "* A candidate is considered to improve accuracy only if A shows no regression beyond noise AND B shows the targeted gain (boundary confusion falls, recall does not fall).",
          "* Development-set numbers from earlier work are not comparable and are not used here."]
    return "\n".join(L) + "\n"


# ----------------------------------------------------------------------------------------------------- CLI


def load_verified_intent_candidate(model_dir: Path):
    """Load the experimental Intent candidate ONLY if every file matches SHA256SUMS (so the exact frozen artifact is what gets scored)."""
    model_dir = Path(model_dir)
    sums = {}
    for line in (model_dir / "SHA256SUMS").read_text().splitlines():
        digest, name = line.split(maxsplit=1)
        sums[name.strip()] = digest
    for name, digest in sums.items():
        if hashlib.sha256((model_dir / name).read_bytes()).hexdigest() != digest:
            raise LabelError(f"intent candidate artifact {name} does not match SHA256SUMS; refusing to evaluate a modified model")
    if "model.joblib" not in sums:
        raise LabelError("SHA256SUMS does not cover model.joblib")
    import joblib

    return joblib.load(model_dir / "model.joblib"), sums["model.joblib"]


def with_intent_candidate(rcs_preds: dict, texts: dict, model) -> dict:
    """RCS predictions with ONLY the intent replaced by the candidate's; topic and sentiment stay the frozen RCS outputs."""
    ids = list(rcs_preds)
    intents = model.predict([texts[i] for i in ids])
    return {i: {**rcs_preds[i], "intent": intents[n]} for n, i in enumerate(ids)}


def compute_rcs_predictions(texts: dict[str, str]) -> dict:
    from rcs.classification import frozen_inference

    ids = list(texts)
    t = [texts[i] for i in ids]
    intent = frozen_inference.predict_intent(t)
    topic = frozen_inference.predict_topic(ids, t)
    sent = frozen_inference.predict_sentiment(ids, t)
    return {i: {"intent": intent[n][0], "topic": topic[i][0], "sentiment": sent[i][0]} for n, i in enumerate(ids)}


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--labels", required=True)
    ap.add_argument("--second-labels")
    ap.add_argument("--adjudicated", help="final adjudicated labels workbook (required when two reviewers disagree)")
    ap.add_argument("--jev-predictions", help="jsonl with eval_id (= record_id) + intent/topic/sentiment from a Jev run (not made by this script)")
    ap.add_argument("--intent-candidate-dir", help="experimental Intent candidate directory (verified against SHA256SUMS before use)")
    ap.add_argument("--out-dir", default=str(BASE / "results"))
    args = ap.parse_args()

    mapping = {r["record_id"]: r for r in csv.DictReader(open(BASE / "coordinator_only" / "blind_eval_mapping.local.csv", encoding="utf-8"))}
    ids = sorted(mapping)
    summary = json.loads((BASE / "coordinator_only" / "sampling_summary.json").read_text())
    la = validate_labels(load_labels_xlsx(Path(args.labels)), ids)
    final = la["scored"]; insufficient = la["insufficient"]; reviewers = 1; agree = None
    if args.second_labels:
        lb = validate_labels(load_labels_xlsx(Path(args.second_labels)), ids)
        agree = agreement(la["scored"], lb["scored"])
        reviewers = 2
        if any(agree["disagreements"][t] for t in TASKS) or set(la["insufficient"]) != set(lb["insufficient"]):
            if not args.adjudicated:
                raise LabelError("two reviewers disagree; supply --adjudicated with the final labels (disagreements: "
                                 + json.dumps({t: agree["disagreements"][t] for t in TASKS}) + ")")
            adj = validate_labels(load_labels_xlsx(Path(args.adjudicated)), ids)
            final, insufficient = adj["scored"], adj["insufficient"]
    # labels are locked: only now are predictions produced
    review_rows = load_labels_xlsx(Path(args.labels))
    texts = {i: str(review_rows[i]["text"]) for i in final}
    rcs = compute_rcs_predictions(texts)
    lemma = {i: bool(int(mapping[i]["lemma_present"])) for i in final}
    systems = {"Frozen RCS": rcs, "RCS + WS/EC layer (candidate)": with_layer(rcs, texts, lemma)}
    comparisons = [("Frozen RCS", "RCS + WS/EC layer (candidate)")]
    candidate_sha = None
    if args.intent_candidate_dir:
        model, candidate_sha = load_verified_intent_candidate(Path(args.intent_candidate_dir))
        systems["RCS + Intent candidate v1"] = with_intent_candidate(rcs, texts, model)
        comparisons.append(("Frozen RCS", "RCS + Intent candidate v1"))
    if args.jev_predictions:
        jev = {r["eval_id"]: r for r in (json.loads(l) for l in Path(args.jev_predictions).read_text(encoding="utf-8").splitlines() if l.strip())}
        jev = {i: {t: jev[i][t] for t in TASKS} for i in final}
        systems["Jev v1"] = jev
        systems["Jev v1 + WS/EC layer"] = with_layer(jev, texts, lemma)
        comparisons += [("Jev v1", "Jev v1 + WS/EC layer"), ("Frozen RCS", "Jev v1")]
    component = {i: mapping[i]["component"] for i in final}
    res = evaluate(final, systems, component, comparisons)
    dis = disagreement_rows(final, systems, component)
    res["top_disagreement_pairs"] = {s_: {c: {t: [[list(k), n] for k, n in v] for t, v in tasks.items()} for c, tasks in comps.items()}
                                     for s_, comps in top_confusions(dis).items()}
    spec = json.loads(SPEC_FILE.read_text())["sha256"] if SPEC_FILE.exists() else "unknown"
    meta = {"sample_ids_sha256": summary["sample_ids_sha256"], "labels_sha256": hashlib.sha256(Path(args.labels).read_bytes()).hexdigest(),
            "candidate_spec_sha256": spec, "intent_candidate_model_sha256": candidate_sha, "reviewers": reviewers, "insufficient": len(insufficient), "agreement": agree}
    out = Path(args.out_dir); out.mkdir(parents=True, exist_ok=True)
    (out / "blind_eval_report.md").write_text(render(res, meta), encoding="utf-8")
    (out / "blind_eval_results.json").write_text(json.dumps({"meta": meta, "results": res}, indent=1, default=str))
    (out / "rcs_predictions.json").write_text(json.dumps(rcs, indent=1))
    with open(out / "disagreements_with_human_labels.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=["system", "component", "record_id", "task", "human_label", "system_prediction"])
        w.writeheader(); w.writerows(dis)
    for f in out.iterdir():
        f.chmod(0o600)
    print(render(res, meta))


if __name__ == "__main__":
    main()
