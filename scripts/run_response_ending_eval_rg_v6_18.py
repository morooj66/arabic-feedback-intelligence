#!/usr/bin/env python3
"""Live rg-v6 ending/closing check on 18 GENERATE_DRAFT comments.

Does not overwrite response_quality_review_rg_v5_30.xlsx.
Does not send replies to visitors.
"""

from __future__ import annotations

import importlib.util
import json
import re
import sys
import time
from collections import Counter
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT / "src"))
sys.path.insert(0, str(PROJECT_ROOT))

spec = importlib.util.spec_from_file_location(
    "eq30", PROJECT_ROOT / "scripts" / "run_response_quality_eval_30.py"
)
eq = importlib.util.module_from_spec(spec)
assert spec.loader is not None
spec.loader.exec_module(eq)

import pandas as pd  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from config import settings  # noqa: E402
from rcs.response_generation.eligibility import EligibilityInput, determine_response_eligibility  # noqa: E402
from rcs.response_generation.service import GenerationRequest, build_live_gemini_client, generate_suggested_response  # noqa: E402
from rcs.review.taxonomy_labels import INTENT_AR, SENTIMENT_AR, TOPIC_AR  # noqa: E402
from rcs.routing.engine import determine_routing  # noqa: E402
from rcs.training.gold_splits import load_supervised_gold  # noqa: E402

OUT_XLSX = PROJECT_ROOT / "response_ending_review_rg_v6_18.xlsx"
PREV_FILES = [
    PROJECT_ROOT / "response_quality_review_rg_v5_30.xlsx",
    PROJECT_ROOT / "response_quality_review_rg_v4_15.xlsx",
    PROJECT_ROOT / "response_quality_review_rg_v3_15.xlsx",
    PROJECT_ROOT / "response_quality_review_30.xlsx",
]
SAMPLE_SEED = 48
SAMPLE_N = 18
SENT_SPLIT = re.compile(r"(?<=[.!?؟۔])\s+")
CLOSING_PATTERNS = (
    ("نشكرك على تواصلك", re.compile(r"نشكرك على تواصلكم?")),
    ("نسعد بتواصلك", re.compile(r"نسعد بتواصلكم?")),
    ("شاكرين ومقدرين", re.compile(r"شاكرين ومقدرين")),
    ("فريق المعجم", re.compile(r"فريق المعجم")),
    ("شكرًا لتواصلك", re.compile(r"شكرًا لتواصلكم?")),
)

SYSTEM_COLS = (
    "comment_id",
    "comment_text",
    "intent_ar",
    "topic_ar",
    "comment_tone_ar",
    "rg_v6_draft",
    "ending_sentence",
    "ending_added",
    "sentence_count",
    "matched_closing_phrases",
    "routing_decision",
    "generation_status",
    "prompt_version",
    "model_name",
)
HUMAN_COLS = ("awkward_ending", "human_notes")


def previous_eval_ids() -> set[str]:
    ids: set[str] = set()
    for path in PREV_FILES:
        if not path.exists():
            continue
        prev = pd.read_excel(path, sheet_name="Review", dtype=str)
        if "comment_id" in prev.columns:
            ids.update(prev["comment_id"].dropna().astype(str))
    return ids


def sentences(text: str) -> list[str]:
    parts = [p.strip() for p in SENT_SPLIT.split((text or "").strip()) if p.strip()]
    return parts or ([text.strip()] if text and text.strip() else [])


def ending_fields(draft: str) -> dict:
    sents = sentences(draft)
    last = sents[-1] if sents else ""
    added = "yes" if len(sents) >= 2 else "no"
    matched = [name for name, pattern in CLOSING_PATTERNS if pattern.search(last)]
    return {
        "ending_sentence": last,
        "ending_added": added,
        "sentence_count": len(sents),
        "matched_closing_phrases": " | ".join(matched),
        "_matched": matched,
    }


def select_18(scored: pd.DataFrame, rng, *, exclude_ids: set[str]) -> tuple[pd.DataFrame, dict]:
    used: set[str] = set()
    buckets: dict[str, int] = {}
    eligible = scored[~scored["comment_id"].isin(exclude_ids)]
    gen = eligible[eligible["pred_routing"] == "GENERATE_DRAFT"]

    def take_intent(label: str, n: int) -> None:
        buckets[f"gold_intent:{label}"] = len(eq._take(rng, gen[gen["gold_primary_intent"] == label], n, used))

    def take_topic(label: str, n: int) -> None:
        buckets[f"gold_topic:{label}"] = len(eq._take(rng, gen[gen["gold_topic_subtype"] == label], n, used))

    take_intent("THANKS_PRAISE", 3)
    take_intent("SUGGESTION", 3)
    take_intent("COMPLAINT", 3)
    take_intent("INQUIRY", 3)
    take_topic("WORD_SUBMISSION", 2)
    take_topic("ENTRY_CORRECTION", 1)
    take_topic("TECHNICAL_ISSUE", 1)
    leftover = SAMPLE_N - len(used)
    buckets["leftover_generate"] = len(eq._take(rng, gen, leftover, used))
    if len(used) != SAMPLE_N:
        raise SystemExit(f"sample size {len(used)} != {SAMPLE_N}")
    order = list(used)
    rng2 = eq.np.random.default_rng(SAMPLE_SEED + 1)
    rng2.shuffle(order)
    ordered = scored.set_index("comment_id").loc[order].reset_index()
    return ordered, {"selection_seed": SAMPLE_SEED, "bucket_fills": buckets, "generate_pool": int(len(gen))}


def write_workbook(rows: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Endings"
    ws.sheet_view.rightToLeft = True
    ws.freeze_panes = "A2"
    header_fill = PatternFill("solid", fgColor="1F4E78")
    human_fill = PatternFill("solid", fgColor="FFF2CC")
    header_font = Font(color="FFFFFF", bold=True)
    wrap = Alignment(wrap_text=True, vertical="top", horizontal="right")
    top = Alignment(vertical="top", wrap_text=True)
    thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    cols = list(SYSTEM_COLS) + list(HUMAN_COLS)
    for col_idx, name in enumerate(cols, 1):
        cell = ws.cell(1, col_idx, name)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin
    for r_i, row in enumerate(rows, 2):
        for c_i, name in enumerate(SYSTEM_COLS, 1):
            cell = ws.cell(r_i, c_i, row.get(name) if row.get(name) is not None else "")
            cell.border = thin
            cell.alignment = wrap if name in ("comment_text", "rg_v6_draft", "ending_sentence") else top
            if name == "comment_id":
                cell.number_format = "@"
        for c_i, _name in enumerate(HUMAN_COLS, 1 + len(SYSTEM_COLS)):
            cell = ws.cell(r_i, c_i, None)
            cell.fill = human_fill
            cell.border = thin
        ws.row_dimensions[r_i].height = 72
    for letter, width in {
        "A": 14, "B": 46, "C": 14, "D": 18, "E": 12, "F": 48, "G": 36, "H": 12,
        "I": 12, "J": 22, "K": 20, "L": 16, "M": 12, "N": 18, "O": 16, "P": 28,
    }.items():
        ws.column_dimensions[letter].width = width
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{1 + len(rows)}"
    ins = wb.create_sheet("Instructions")
    ins["A1"] = (
        "rg-v6 ending review. Automated ending_added = yes when the draft has two or more "
        "sentences (last sentence treated as the ending). Human columns stay blank. "
        "This is not a quality approval."
    )
    ins.column_dimensions["A"].width = 110
    wb.save(OUT_XLSX)


def main() -> None:
    if settings.RESPONSE_GENERATION_MODEL_NAME != "gemini-3.6-flash":
        raise SystemExit("Unexpected model; not overriding.")
    if settings.RESPONSE_GENERATION_PROMPT_VERSION != "rg-v6":
        raise SystemExit("Expected rg-v6.")
    exclusion = set(json.loads(eq.EXCLUSION_JSON.read_text())["test_ids"])
    gold = pd.read_csv(eq.GOLD_CSV, dtype=str, keep_default_na=False)
    human_test_ids = set(gold.loc[gold["annotation_stage"] == "human_test", "comment_id"])
    prev_ids = previous_eval_ids() | exclusion | human_test_ids
    pool = gold[
        (gold["annotation_stage"] != "human_test")
        & (~gold["comment_id"].isin(exclusion))
        & (gold["comment_text"].astype(str).str.strip() != "")
    ].copy()
    freeze = json.loads(eq.FREEZE_JSON.read_text())
    data = load_supervised_gold(eq.GOLD_CSV, include_test=False)
    intent_pipe = eq.fit_frozen_intent_pipeline(eq.remap_intent(data.intent.train))
    texts = pool["comment_text"].tolist()
    intent_labels = list(intent_pipe.classes_)
    intent_probs = intent_pipe.predict_proba(texts)
    import numpy as np

    pool["pred_intent"] = [intent_labels[int(np.argmax(p))] for p in intent_probs]
    pool["intent_conf"] = [float(np.max(p)) for p in intent_probs]
    topic_pred, topic_conf = eq.predict_topic(texts, eq.find_topic_checkpoint(), freeze)
    sent_pred, sent_conf = eq.predict_sentiment(texts, freeze)
    pool["pred_topic"] = topic_pred
    pool["topic_conf"] = topic_conf
    pool["pred_sentiment"] = sent_pred
    pool["sent_conf"] = sent_conf
    pool["pred_confidence"] = [
        min(i, t, s) for i, t, s in zip(pool["intent_conf"], pool["topic_conf"], pool["sent_conf"])
    ]
    routes, reasons, class_routes, needs_flags = [], [], [], []
    for rec in pool.itertuples(index=False):
        class_route = determine_routing(rec.pred_intent, rec.pred_topic, None, None, True)
        needs_review = rec.pred_confidence < settings.RESPONSE_LOW_CONFIDENCE_THRESHOLD
        el = determine_response_eligibility(
            EligibilityInput(
                comment_text=rec.comment_text,
                primary_intent=rec.pred_intent,
                topic_subtype=rec.pred_topic,
                sentiment=rec.pred_sentiment,
                confidence=rec.pred_confidence,
                needs_review=needs_review,
                is_content_meaningful=True,
                insufficient_context=rec.gold_annotation_status == "INSUFFICIENT_CONTEXT",
                trusted_grounding_available=False,
                classification_routing_decision=class_route.routing_decision,
            )
        )
        routes.append(el.routing_decision)
        reasons.append(el.reason_code)
        class_routes.append(class_route.routing_decision)
        needs_flags.append(needs_review)
    pool["pred_routing"] = routes
    pool["pred_reason"] = reasons
    pool["class_route"] = class_routes
    pool["pred_needs_review"] = needs_flags

    rng = eq.np.random.default_rng(SAMPLE_SEED)
    sample, selection_meta = select_18(pool, rng, exclude_ids=prev_ids)
    if set(sample["comment_id"]) & prev_ids:
        raise SystemExit("sampled a previous eval or sealed human_test id")

    inner = build_live_gemini_client()
    client = eq.CountingGeminiClient(inner)
    rows: list[dict] = []
    for rec in sample.itertuples(index=False):
        req = GenerationRequest(
            comment_id=str(rec.comment_id),
            comment_text=str(rec.comment_text),
            primary_intent=rec.pred_intent,
            topic_subtype=rec.pred_topic,
            sentiment=rec.pred_sentiment,
            confidence=float(rec.pred_confidence),
            needs_review=bool(rec.pred_needs_review),
            is_content_meaningful=True,
            insufficient_context=str(rec.gold_annotation_status) == "INSUFFICIENT_CONTEXT",
            trusted_grounding_available=False,
            classification_routing_decision=rec.class_route,
        )
        client.set_comment_id(str(rec.comment_id))
        result = generate_suggested_response(req, client=client)
        if result.routing_decision == "GENERATE_DRAFT":
            time.sleep(0.4)
        draft = result.suggested_response or ""
        extra = ending_fields(draft)
        rows.append(
            {
                "comment_id": str(rec.comment_id),
                "comment_text": str(rec.comment_text),
                "intent_ar": INTENT_AR.get(rec.pred_intent, rec.pred_intent),
                "topic_ar": TOPIC_AR.get(rec.pred_topic, rec.pred_topic),
                "comment_tone_ar": SENTIMENT_AR.get(rec.pred_sentiment, rec.pred_sentiment),
                "rg_v6_draft": draft,
                "ending_sentence": extra["ending_sentence"],
                "ending_added": extra["ending_added"],
                "sentence_count": extra["sentence_count"],
                "matched_closing_phrases": extra["matched_closing_phrases"],
                "routing_decision": result.routing_decision,
                "generation_status": result.generation_status,
                "prompt_version": result.prompt_version,
                "model_name": result.model_name,
                "_matched": extra["_matched"],
            }
        )

    write_workbook(rows)
    phrase_counts = Counter()
    for row in rows:
        for name in row["_matched"]:
            phrase_counts[name] += 1
    endings = [row["ending_sentence"] for row in rows if row["ending_added"] == "yes"]
    ending_dupes = {text: n for text, n in Counter(endings).items() if n > 1}
    report = {
        "selection": selection_meta,
        "n": len(rows),
        "gemini_calls": len(client.call_comment_ids),
        "generate_succeeded": sum(1 for r in rows if r["generation_status"] == "SUCCEEDED"),
        "ending_added_yes": sum(1 for r in rows if r["ending_added"] == "yes"),
        "ending_added_no": sum(1 for r in rows if r["ending_added"] == "no"),
        "closing_phrase_counts": dict(phrase_counts),
        "exact_duplicate_endings": ending_dupes,
        "model": settings.RESPONSE_GENERATION_MODEL_NAME,
        "prompt_version": settings.RESPONSE_GENERATION_PROMPT_VERSION,
        "xlsx": str(OUT_XLSX),
        "did_not_overwrite_rg_v5": True,
        "human_columns_blank": True,
        "precheck_is_not_human_score": True,
        "samples": [
            {
                "comment_id": r["comment_id"],
                "comment_text": r["comment_text"],
                "rg_v6_draft": r["rg_v6_draft"],
                "ending_sentence": r["ending_sentence"],
                "ending_added": r["ending_added"],
                "matched_closing_phrases": r["matched_closing_phrases"],
            }
            for r in rows
        ],
    }
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
