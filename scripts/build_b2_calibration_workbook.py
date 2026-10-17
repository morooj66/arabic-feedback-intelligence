#!/usr/bin/env python3
"""
Phase B.2 calibration experiment — build a 40-row reviewer workbook for
controlled human labeling before any further full Gemini rerun.

READ-ONLY with respect to data/gold/gold_dataset_round1.csv: reads the
canonical CSV (b2-v3 predictions) and the frozen pre-rerun B.3 export
(gold_dataset_round1_b3_review.xlsx, b2-v2 predictions) to side-by-side
compare both prompt versions on the same comments.

Writes a NEW workbook only (default: data/gold/b2_calibration_sample40.xlsx).
Does not import labels back into the canonical CSV.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT / "src"))
sys.path.insert(0, str(PROJECT_ROOT))

import pandas as pd  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill, Protection  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

from config import settings  # noqa: E402

GOLD_CSV = settings.DATA_DIR / "gold" / "gold_dataset_round1.csv"
B2_V2_XLSX = settings.DATA_DIR / "gold" / "gold_dataset_round1_b3_review.xlsx"
OUT_XLSX = settings.DATA_DIR / "gold" / "b2_calibration_sample40.xlsx"
OUT_MANIFEST = settings.DATA_DIR / "gold" / "b2_calibration_sample40_manifest.json"

CALIBRATION_SEED = 42
SAMPLE_SIZE = 40

# Quotas must sum to SAMPLE_SIZE.
STRATA: list[tuple[str, int]] = [
    ("short_bare_word", 6),
    ("suggestion_word_submission", 8),
    ("inquiry", 5),
    ("complaint_entry_correction", 5),
    ("spam_off_topic_personal", 5),
    ("dictionary_suggestion", 4),
    ("long_structured", 7),
]

VISIBLE_COLUMNS = [
    "comment_text",
    "b2_v2_primary_intent",
    "b2_v2_topic_subtype",
    "b2_v2_sentiment",
    "b2_v3_primary_intent",
    "b2_v3_topic_subtype",
    "b2_v3_sentiment",
    "gold_primary_intent",
    "gold_topic_subtype",
    "gold_sentiment",
    "gold_annotation_status",
]

HIDDEN_COLUMNS = [
    "comment_id",
    "calibration_stratum",
]

ALL_COLUMNS = VISIBLE_COLUMNS + HIDDEN_COLUMNS
REVIEWER_EDITABLE_COLUMNS = [
    "gold_primary_intent",
    "gold_topic_subtype",
    "gold_sentiment",
    "gold_annotation_status",
]

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
STRATUM_FILLS = {
    "short_bare_word": PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"),
    "suggestion_word_submission": PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid"),
    "inquiry": PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
    "complaint_entry_correction": PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"),
    "spam_off_topic_personal": PatternFill(start_color="EDEDED", end_color="EDEDED", fill_type="solid"),
    "dictionary_suggestion": PatternFill(start_color="E4DFEC", end_color="E4DFEC", fill_type="solid"),
    "long_structured": PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid"),
    "fill": PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid"),
}


def _rank_key(stratum: str, comment_id: str, seed: int) -> str:
    return hashlib.sha256(f"{seed}:{stratum}:{comment_id}".encode()).hexdigest()


def _is_short_bare_word(text: str) -> bool:
    t = text.strip()
    if not t or "\n" in t:
        return False
    words = t.split()
    return len(t) <= 40 and len(words) <= 5


def _is_long_structured(text: str) -> bool:
    t = text.strip()
    if len(t) >= 200:
        return True
    if t.count("\n") >= 2:
        return True
    markers = ("المعنى", "Phone Number", "محتوي الرسالة", "source:", "entryId")
    return any(marker in t for marker in markers)


def _stratum_predicate(stratum: str, row: pd.Series) -> bool:
    text = row["comment_text"]
    if stratum == "short_bare_word":
        return _is_short_bare_word(text)
    if stratum == "suggestion_word_submission":
        return (
            row["b2_v3_primary_intent"] == "SUGGESTION"
            and row["b2_v3_topic_subtype"] == "WORD_SUBMISSION"
        )
    if stratum == "inquiry":
        return row["b2_v3_primary_intent"] == "INQUIRY" or row["b2_v2_primary_intent"] == "INQUIRY"
    if stratum == "complaint_entry_correction":
        pairs = {
            (row["b2_v3_primary_intent"], row["b2_v3_topic_subtype"]),
            (row["b2_v2_primary_intent"], row["b2_v2_topic_subtype"]),
        }
        return ("COMPLAINT", "ENTRY_CORRECTION") in pairs
    if stratum == "spam_off_topic_personal":
        pairs = {
            (row["b2_v3_primary_intent"], row["b2_v3_topic_subtype"]),
            (row["b2_v2_primary_intent"], row["b2_v2_topic_subtype"]),
        }
        return ("SPAM", "OFF_TOPIC_PERSONAL") in pairs
    if stratum == "dictionary_suggestion":
        return (
            row["b2_v3_topic_subtype"] == "DICTIONARY_SUGGESTION"
            or row["b2_v2_topic_subtype"] == "DICTIONARY_SUGGESTION"
        )
    if stratum == "long_structured":
        return _is_long_structured(text)
    raise ValueError(f"unknown stratum: {stratum}")


def _pick_from_pool(pool: pd.DataFrame, stratum: str, n: int, seed: int) -> pd.DataFrame:
    if pool.empty or n <= 0:
        return pool.iloc[0:0].copy()
    ranked = pool.copy()
    ranked["_rank"] = ranked["comment_id"].map(lambda cid: _rank_key(stratum, cid, seed))
    return ranked.sort_values("_rank").head(n).drop(columns="_rank")


def load_merged_predictions(
    gold_csv_path: Path = GOLD_CSV,
    b2_v2_xlsx_path: Path = B2_V2_XLSX,
) -> pd.DataFrame:
    canonical = pd.read_csv(gold_csv_path, dtype=str, keep_default_na=False)
    targets = canonical[canonical["annotation_stage"] == "ai_assisted_train"].copy()
    if len(targets) != 300:
        raise AssertionError(f"expected 300 ai_assisted_train rows, found {len(targets)}")

    if not b2_v2_xlsx_path.exists():
        raise FileNotFoundError(
            f"{b2_v2_xlsx_path} not found — needed for frozen b2-v2 predictions"
        )
    v2 = pd.read_excel(b2_v2_xlsx_path, sheet_name="Review", dtype=str).fillna("")
    if len(v2) != 300:
        raise AssertionError(f"expected 300 rows in b2-v2 workbook, found {len(v2)}")
    if not (v2["ai_suggestion_model"].astype(str).str.contains("b2-v2").all()):
        raise AssertionError("b2-v2 workbook does not contain only b2-v2 predictions")

    merged = targets.merge(
        v2[
            [
                "comment_id",
                "ai_suggested_primary_intent",
                "ai_suggested_topic_subtype",
                "ai_suggested_sentiment",
            ]
        ],
        on="comment_id",
        how="inner",
        suffixes=("_v3", "_v2wb"),
    )
    if len(merged) != 300:
        raise AssertionError("comment_id mismatch between canonical CSV and b2-v2 workbook")

    merged = merged.rename(
        columns={
            "ai_suggested_primary_intent_v3": "b2_v3_primary_intent",
            "ai_suggested_topic_subtype_v3": "b2_v3_topic_subtype",
            "ai_suggested_sentiment_v3": "b2_v3_sentiment",
            "ai_suggested_primary_intent_v2wb": "b2_v2_primary_intent",
            "ai_suggested_topic_subtype_v2wb": "b2_v2_topic_subtype",
            "ai_suggested_sentiment_v2wb": "b2_v2_sentiment",
        }
    )
    if not (merged["ai_suggestion_model"].astype(str).str.contains("b2-v3").all()):
        raise AssertionError("canonical CSV is not entirely on b2-v3 predictions")
    return merged


def select_calibration_sample(
    merged: pd.DataFrame,
    *,
    seed: int = CALIBRATION_SEED,
    sample_size: int = SAMPLE_SIZE,
) -> tuple[pd.DataFrame, dict]:
    if sum(q for _, q in STRATA) != sample_size:
        raise AssertionError("strata quotas must sum to sample_size")

    selected_frames: list[pd.DataFrame] = []
    selected_ids: set[str] = set()
    composition: dict[str, list[str]] = {}

    for stratum, quota in STRATA:
        pool = merged[~merged["comment_id"].isin(selected_ids)]
        pool = pool[pool.apply(lambda row: _stratum_predicate(stratum, row), axis=1)]
        picked = _pick_from_pool(pool, stratum, quota, seed)
        picked = picked.copy()
        picked["calibration_stratum"] = stratum
        selected_frames.append(picked)
        selected_ids.update(picked["comment_id"].tolist())
        composition[stratum] = picked["comment_id"].tolist()

    selected = pd.concat(selected_frames, ignore_index=True)
    if len(selected) < sample_size:
        remaining = merged[~merged["comment_id"].isin(selected_ids)].copy()
        fill_n = sample_size - len(selected)
        fill = _pick_from_pool(remaining, "fill", fill_n, seed)
        fill["calibration_stratum"] = "fill"
        selected = pd.concat([selected, fill], ignore_index=True)
        composition["fill"] = fill["comment_id"].tolist()

    if len(selected) != sample_size:
        raise AssertionError(f"selected {len(selected)} rows, expected {sample_size}")

    selected = selected.sort_values("comment_id").reset_index(drop=True)
    manifest = {
        "seed": seed,
        "sample_size": sample_size,
        "strata_quotas": dict(STRATA),
        "composition_by_stratum": composition,
        "comment_ids_in_workbook_order": selected.sort_values("comment_id")["comment_id"].tolist(),
    }
    return selected, manifest


def build(
    gold_csv_path: Path = GOLD_CSV,
    b2_v2_xlsx_path: Path = B2_V2_XLSX,
    out_path: Path = OUT_XLSX,
    manifest_path: Path = OUT_MANIFEST,
    *,
    seed: int = CALIBRATION_SEED,
) -> tuple[Path, dict]:
    merged = load_merged_predictions(gold_csv_path, b2_v2_xlsx_path)
    sample, manifest = select_calibration_sample(merged, seed=seed)

    export = sample[
        [
            "comment_id",
            "comment_text",
            "b2_v2_primary_intent",
            "b2_v2_topic_subtype",
            "b2_v2_sentiment",
            "b2_v3_primary_intent",
            "b2_v3_topic_subtype",
            "b2_v3_sentiment",
            "gold_primary_intent",
            "gold_topic_subtype",
            "gold_sentiment",
            "gold_annotation_status",
            "calibration_stratum",
        ]
    ].copy()

    wb = Workbook()
    ws_i = wb.active
    ws_i.title = "Instructions"
    instructions = [
        ("Phase B.2 Calibration Sample — 40-row human ground truth", True),
        ("", False),
        ("Purpose: label these 40 comments so we can compare b2-v2 vs b2-v3", False),
        ("and evaluate a proposed topic-aware few-shot calibration.", False),
        ("This is NOT Phase B.3 and does NOT write back to the canonical CSV.", False),
        ("", False),
        ("Fill gold_primary_intent, gold_topic_subtype, gold_sentiment, and", False),
        ("gold_annotation_status from the comment text only.", False),
        ("Use LABELED when confident; UNCERTAIN or INSUFFICIENT_CONTEXT otherwise.", False),
        ("", False),
        ("b2_v2_* columns are frozen predictions from prompt version b2-v2.", False),
        ("b2_v3_* columns are frozen predictions from prompt version b2-v3.", False),
        ("Row colors map to the hidden calibration_stratum (sampling bucket).", False),
    ]
    for i, (text, bold) in enumerate(instructions, start=1):
        cell = ws_i.cell(row=i, column=1, value=text)
        if bold:
            cell.font = Font(bold=True, size=13 if i == 1 else 11)
    ws_i.column_dimensions["A"].width = 100

    ws = wb.create_sheet("Review")
    for col_idx, col_name in enumerate(ALL_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    ws.freeze_panes = "A2"

    for r_idx, (_, row) in enumerate(export.iterrows(), start=2):
        stratum = row["calibration_stratum"]
        row_fill = STRATUM_FILLS.get(stratum)
        for c_idx, col_name in enumerate(ALL_COLUMNS, start=1):
            val = row[col_name]
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            if col_name == "comment_text":
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            if col_name in REVIEWER_EDITABLE_COLUMNS:
                cell.protection = Protection(locked=False)
            else:
                cell.protection = Protection(locked=True)
            if row_fill is not None:
                cell.fill = row_fill

    last_row = len(export) + 1
    width_map = {
        "comment_text": 60,
        "b2_v2_primary_intent": 18,
        "b2_v2_topic_subtype": 20,
        "b2_v2_sentiment": 14,
        "b2_v3_primary_intent": 18,
        "b2_v3_topic_subtype": 20,
        "b2_v3_sentiment": 14,
        "gold_primary_intent": 18,
        "gold_topic_subtype": 20,
        "gold_sentiment": 14,
        "gold_annotation_status": 20,
    }
    for col_idx, col_name in enumerate(ALL_COLUMNS, start=1):
        letter = get_column_letter(col_idx)
        if col_name in width_map:
            ws.column_dimensions[letter].width = width_map[col_name]
        if col_name in HIDDEN_COLUMNS:
            ws.column_dimensions[letter].hidden = True

    def _col_letter(name: str) -> str:
        return get_column_letter(ALL_COLUMNS.index(name) + 1)

    def _add_dv(col_name: str, options: list[str], allow_blank: bool = True) -> None:
        formula = '"' + ",".join(options) + '"'
        dv = DataValidation(type="list", formula1=formula, allow_blank=allow_blank)
        ws.add_data_validation(dv)
        letter = _col_letter(col_name)
        dv.add(f"{letter}2:{letter}{last_row}")

    _add_dv("gold_primary_intent", list(settings.PRIMARY_INTENTS))
    _add_dv("gold_topic_subtype", list(settings.TOPIC_SUBTYPES))
    _add_dv("gold_sentiment", list(settings.SENTIMENTS))
    _add_dv("gold_annotation_status", ["LABELED", "UNCERTAIN", "INSUFFICIENT_CONTEXT"])

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return out_path, manifest


def report_composition(sample: pd.DataFrame, manifest: dict) -> str:
    lines = [
        f"Calibration sample: {len(sample)} rows (seed={manifest['seed']})",
        "",
        "Rows per sampling stratum:",
    ]
    for stratum, _quota in STRATA:
        ids = manifest["composition_by_stratum"].get(stratum, [])
        lines.append(f"  {stratum}: {len(ids)}")
    if manifest["composition_by_stratum"].get("fill"):
        lines.append(f"  fill: {len(manifest['composition_by_stratum']['fill'])}")

    lines.extend(["", "b2-v2 vs b2-v3 agreement on sample:"])
    agree_intent = (sample["b2_v2_primary_intent"] == sample["b2_v3_primary_intent"]).sum()
    agree_topic = (sample["b2_v2_topic_subtype"] == sample["b2_v3_topic_subtype"]).sum()
    agree_pair = (
        (sample["b2_v2_primary_intent"] == sample["b2_v3_primary_intent"])
        & (sample["b2_v2_topic_subtype"] == sample["b2_v3_topic_subtype"])
    ).sum()
    lines.append(f"  primary_intent match: {agree_intent}/{len(sample)}")
    lines.append(f"  topic_subtype match: {agree_topic}/{len(sample)}")
    lines.append(f"  intent+topic match: {agree_pair}/{len(sample)}")

    lines.extend(["", "b2-v3 pair distribution in sample:"])
    pairs = sample.groupby(["b2_v3_primary_intent", "b2_v3_topic_subtype"]).size().sort_values(ascending=False)
    for (intent, topic), count in pairs.items():
        lines.append(f"  {intent}/{topic}: {count}")

    lines.extend(["", "Sample rows (comment_id | stratum | b2-v2 -> b2-v3):"])
    for _, row in sample.sort_values("comment_id").iterrows():
        text = row["comment_text"].replace("\n", " ")[:70]
        lines.append(
            f"  {row['comment_id'][:12]}… | {row['calibration_stratum']} | "
            f"{row['b2_v2_primary_intent']}/{row['b2_v2_topic_subtype']} -> "
            f"{row['b2_v3_primary_intent']}/{row['b2_v3_topic_subtype']} | {text!r}"
        )
    return "\n".join(lines)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Build the 40-row B.2 calibration workbook.")
    parser.add_argument("--gold-csv", type=Path, default=GOLD_CSV)
    parser.add_argument("--b2-v2-xlsx", type=Path, default=B2_V2_XLSX)
    parser.add_argument("--out", type=Path, default=OUT_XLSX)
    parser.add_argument("--manifest", type=Path, default=OUT_MANIFEST)
    parser.add_argument("--seed", type=int, default=CALIBRATION_SEED)
    args = parser.parse_args()

    out_path, manifest = build(
        gold_csv_path=args.gold_csv,
        b2_v2_xlsx_path=args.b2_v2_xlsx,
        out_path=args.out,
        manifest_path=args.manifest,
        seed=args.seed,
    )
    sample = load_merged_predictions(args.gold_csv, args.b2_v2_xlsx)
    selected_ids = set(manifest["comment_ids_in_workbook_order"])
    sample = sample[sample["comment_id"].isin(selected_ids)].copy()
    sample["calibration_stratum"] = sample["comment_id"].map(
        {
            cid: stratum
            for stratum, ids in manifest["composition_by_stratum"].items()
            for cid in ids
        }
    )
    print(f"Wrote {out_path}")
    print(f"Wrote {args.manifest}")
    print()
    print(report_composition(sample, manifest))
