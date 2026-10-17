"""Completed-export end-to-end (real Chrome -> real Vite frontend -> real backend). SYNTHETIC data only. Reuses the seed of verify_reviewer_flow.py.

    python scripts/e2e/verify_reviewer_flow.py seed <workdir>
    python scripts/e2e/serve_e2e_mock.py --workdir <workdir> --port 8794
    python scripts/e2e/verify_export_flow.py run <workdir> http://localhost:5193 http://127.0.0.1:8794
Reviewers decide through the real API, then the ADMIN exports through the real "المراجعات المكتملة" page; the downloaded .xlsx is compared with the database.
"""
from __future__ import annotations

import hashlib
import http.cookiejar
import io
import json
import re
import sqlite3
import sys
import urllib.request
from pathlib import Path


def main(work: Path, base: str, api: str) -> int:
    import openpyxl
    from playwright.sync_api import sync_playwright

    U = json.loads((work / "e2e_users.json").read_text())
    DB = work / "e2e.sqlite"
    R: list[dict] = []

    def check(name, ok, detail=""):
        R.append({"check": name, "result": "PASS" if ok else "FAIL", "detail": str(detail)[:300]})
        print(("PASS" if ok else "FAIL"), "-", name, ("| " + str(detail)[:200]) if not ok and detail else "", flush=True)
        return ok

    def q(sql, *a):
        c = sqlite3.connect(f"file:{DB}?mode=ro", uri=True); r = c.execute(sql, a).fetchall(); c.close(); return r

    def fingerprint():
        rows = q("select comment_id, review_status, reviewer_action, reviewer_user_id, row_version, final_reply, reviewer_edit, no_response_reason from review_items order by comment_id")
        ev = q("select event_id, review_item_id, event_type from review_item_events order by event_id")
        return hashlib.sha256(repr((rows, ev)).encode()).hexdigest(), len(ev)

    class Api:
        def __init__(self, who):
            self.jar = http.cookiejar.CookieJar(); self.op = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(self.jar))
            self.post("/api/login", {"email": U[who]["email"], "password": U[who]["password"]})

        def post(self, path, body):
            r = urllib.request.Request(api + path, data=json.dumps(body).encode(), headers={"Content-Type": "application/json"}, method="POST")
            try:
                with self.op.open(r, timeout=60) as resp:
                    return resp.status, json.loads(resp.read().decode() or "{}")
            except urllib.error.HTTPError as e:
                return e.code, {}

        def get(self, path):
            with self.op.open(api + path, timeout=60) as resp:
                return json.loads(resp.read().decode())

    # -------- reviewers decide through the real API: A -> 25 (mixed kinds), B -> 10
    for who, prefix, n in (("reviewer_a", "a", 25), ("reviewer_b", "b", 10)):
        a = Api(who)
        for i in range(n):
            cid = f"{prefix}{i:03d}"
            a.post(f"/api/items/{cid}/claim", {})
            kind = i % 3
            if kind == 0 and q("select suggested_final_reply from review_items where comment_id=?", cid)[0][0]:
                s, _ = a.post(f"/api/items/{cid}/send-as-is", {})
            elif kind == 1:
                s, _ = a.post(f"/api/items/{cid}/edit-and-send", {"edited_response": f"نشكركم على ملاحظتكم رقم {cid}، وستُدرس بعناية."})
            else:
                s, _ = a.post(f"/api/items/{cid}/no-response", {"reason": "لا يتطلب ردًا"})
            assert s == 200, (cid, s)
    done = q("select comment_id from review_items where reviewer_action is not null")
    total_done = len(done)
    pending_ids = {r[0] for r in q("select comment_id from review_items where reviewer_action is null")}
    check("setup: reviewers completed a mix of decisions (more than one page)", total_done > 20, total_done)
    before = fingerprint()

    with sync_playwright() as p:
        br = p.chromium.launch(channel="chrome", headless=True)
        pg = br.new_context(viewport={"width": 1400, "height": 900}, accept_downloads=True).new_page(); pg.on("dialog", lambda d: d.accept())
        pg.goto(base + "/"); pg.wait_for_timeout(600)
        pg.fill('input[placeholder="أدخل البريد الإلكتروني"]', U["admin"]["email"]); pg.fill('input[placeholder="أدخل كلمة المرور"]', U["admin"]["password"])
        pg.click("button[type=submit]"); pg.wait_for_timeout(1800)
        pg.click("text=المراجعات المكتملة >> nth=0"); pg.wait_for_timeout(2000)
        body = re.sub(r"\s+", " ", pg.inner_text("body"))
        check("page: lists every completed review (not only one page)", f"{total_done} مراجعة مكتملة" in body, re.findall(r"\d+ مراجعة مكتملة", body))
        btn = pg.query_selector("button:has-text('تصدير المراجعات المكتملة')")
        check("page: the export button announces the number that will be exported", btn is not None and f"({total_done})" in btn.inner_text(), btn and btn.inner_text())

        def download(button_text):
            with pg.expect_download() as dl:
                pg.click(f"button:has-text('{button_text}')")
            path = work / "ui_export.xlsx"; dl.value.save_as(str(path))
            wb = openpyxl.load_workbook(path)
            h1, *r1 = [list(r) for r in wb["الردود المراجعة"].iter_rows(values_only=True)]
            h2, *r2 = [list(r) for r in wb["بيانات التدقيق"].iter_rows(values_only=True)]
            return h1, r1, h2, r2

        h1, r1, h2, r2 = download("تصدير المراجعات المكتملة")
        ids = [r[h2.index("comment_id")] for r in r2]
        check("export all: the file contains every completed review and nothing else", len(ids) == total_done == len(set(ids)) and not set(ids) & pending_ids, (len(ids), total_done))
        pg.wait_for_timeout(800)
        check("export all: the page reports the actual number exported", f"تم تصدير {total_done} من {total_done}" in re.sub(r"\s+", " ", pg.inner_text("body")), "")
        db = {r[0]: r for r in q("select comment_id, reviewer_action, final_reply, reviewer_edit, no_response_reason, review_status from review_items where reviewer_action is not null")}
        ok = True
        for a_, b_ in zip(r1, r2):
            cid = b_[h2.index("comment_id")]
            _, act, final, edit, reason, status = db[cid]
            ok = ok and a_[h1.index("إجراء المراجع")] == act and (a_[h1.index("الرد النهائي")] or None) == (final or None) and (a_[h1.index("الرد بعد التعديل")] or None) == (edit or None)
            ok = ok and ((b_[h2.index("no_response_reason")] or None) == (reason or None)) and a_[h1.index("حالة المراجعة")] == status
        check("export all: decisions, exact final/edited replies, no-response reasons and statuses match the database row by row", ok, "")
        check("export all: no-response rows carry no reply", all(not r[h1.index("الرد النهائي")] for r in r1 if r[h1.index("إجراء المراجع")] == "NO_RESPONSE_NEEDED"), "")

        pg.fill('input[aria-label="عدد المراجعات المطلوب تصديرها"]', "7"); pg.wait_for_timeout(300)
        h1, r1, h2, r2 = download("تصدير المراجعات المكتملة (7)")
        check("export N: exactly the requested number of completed reviews", len(r1) == 7 and all(r[h1.index("إجراء المراجع")] for r in r1), len(r1))
        pg.fill('input[aria-label="عدد المراجعات المطلوب تصديرها"]', ""); pg.wait_for_timeout(300)

        ra = U["reviewer_a"]["user_id"]
        pg.select_option('select[aria-label="اختر المراجع"]', value=ra); pg.wait_for_timeout(1500)
        a_done = len(q("select comment_id from review_items where reviewer_action is not null and reviewer_user_id=?", ra))
        h1, r1, h2, r2 = download("تصدير المراجعات المكتملة")
        check("filtered export: only the selected reviewer's completed reviews, across all pages", len(r1) == a_done > 20 and all(db[r[h2.index("comment_id")]][0] for r in r2) and
              {x[0] for x in q("select comment_id from review_items where reviewer_user_id=? and reviewer_action is not null", ra)} == {r[h2.index("comment_id")] for r in r2}, (len(r1), a_done))
        check("export is read-only: statuses, assignments, versions, texts and audit history are byte-identical", fingerprint() == before, "")
        br.close()
    out = work / "export_flow_report.json"
    out.write_text(json.dumps({"results": R, "pass": sum(r["result"] == "PASS" for r in R), "fail": sum(r["result"] == "FAIL" for r in R)}, ensure_ascii=False, indent=1))
    print("\nPASS", sum(r["result"] == "PASS" for r in R), "FAIL", sum(r["result"] == "FAIL" for r in R))
    return 0 if all(r["result"] == "PASS" for r in R) else 1


if __name__ == "__main__":
    sys.exit(main(Path(sys.argv[2]), sys.argv[3], sys.argv[4]))
