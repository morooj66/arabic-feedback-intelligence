#!/usr/bin/env python3
"""
Human-evaluation FORMATTING ONLY over the already-existing
data/feedback_research/rg_v8_context_ab_eval_blind.csv.

Does not call Gemini. Does not touch review_items or any operational
database. Does not rerun the A/B experiment. Never reads
rg_v8_context_ab_eval_key.json — the hidden A/B key is not available to
this script at all, so it cannot leak into the output even by accident.

Produces data/feedback_research/rg_v8_context_ab_eval_blind_review.xlsx:
the same 10 blind cases, same response_A/response_B text, reformatted
for a non-technical human reviewer (Arabic headers, RTL layout, wrapped
text, frozen header, autofilter, a dropdown-constrained judgment column).

  python scripts/build_rg_v8_context_ab_eval_review_xlsx.py
"""

from __future__ import annotations

import csv
import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT / "src"))
sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

from rcs.review.taxonomy_labels import INTENT_AR, SENTIMENT_AR, TOPIC_AR  # noqa: E402

SRC_CSV = PROJECT_ROOT / "data" / "feedback_research" / "rg_v8_context_ab_eval_blind.csv"
OUT_XLSX = PROJECT_ROOT / "data" / "feedback_research" / "rg_v8_context_ab_eval_blind_review.xlsx"

HEADERS = [
    "رقم الحالة",
    "تعليق الزائر",
    "المدخل المرتبط / lemma",
    "التصنيف الرئيسي",
    "الموضوع",
    "المشاعر",
    "الرد A",
    "الرد B",
    "الحكم",
    "ملاحظة المراجع",
]
JUDGMENT_CHOICES = ["A أفضل", "B أفضل", "متساويان", "كلاهما غير مناسب"]

# column letter -> width, tuned for readable wrapped Arabic text
COLUMN_WIDTHS = {
    "A": 12,   # رقم الحالة
    "B": 48,   # تعليق الزائر
    "C": 22,   # lemma
    "D": 16,   # التصنيف الرئيسي
    "E": 18,   # الموضوع
    "F": 12,   # المشاعر
    "G": 48,   # الرد A
    "H": 48,   # الرد B
    "I": 20,   # الحكم
    "J": 32,   # ملاحظة المراجع
}

HEADER_FILL = PatternFill("solid", fgColor="1F5C3E")  # brand-green, header only
HEADER_FONT = Font(bold=True, size=11, color="FFFFFF")
BODY_FONT = Font(size=11)


def _estimate_row_height(*texts: str, chars_per_line: int = 58, line_height: float = 15.5) -> float:
    """Rough wrapped-text line count -> row height, so long responses
    aren't visually cut off. Not pixel-perfect (Excel doesn't expose true
    text-measurement to openpyxl), just a reasonable heuristic."""
    max_lines = 1
    for text in texts:
        if not text:
            continue
        lines = 0
        for physical_line in str(text).split("\n"):
            lines += max(1, -(-len(physical_line) // chars_per_line))  # ceil div
        max_lines = max(max_lines, lines)
    return min(max(max_lines * line_height, 32), 320)


def main() -> int:
    if not SRC_CSV.exists():
        print(f"ERROR: source blind CSV not found: {SRC_CSV}", file=sys.stderr)
        return 2

    with SRC_CSV.open(encoding="utf-8-sig") as f:
        rows = list(csv.DictReader(f))
    if len(rows) != 10:
        print(f"WARNING: expected 10 cases in the source CSV, found {len(rows)}.", file=sys.stderr)

    wb = Workbook()
    ws = wb.active
    ws.title = "التقييم"
    ws.sheet_view.rightToLeft = True

    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True, readingOrder=2)
    ws.row_dimensions[1].height = 34

    for row in rows:
        ws.append(
            [
                row["case_id"],
                row["comment_text"],
                row["lemma"] or "",
                INTENT_AR.get(row["intent"], row["intent"]),
                TOPIC_AR.get(row["topic"], row["topic"]),
                SENTIMENT_AR.get(row["sentiment"], row["sentiment"]),
                row["response_A"],
                row["response_B"],
                "",  # judgment: left blank for the reviewer
                "",  # note: left blank for the reviewer
            ]
        )

    for col_letter, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[col_letter].width = width

    for row_idx, csv_row in enumerate(rows, start=2):
        for cell in ws[row_idx]:
            cell.font = BODY_FONT
            cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="right", readingOrder=2)
        ws.row_dimensions[row_idx].height = _estimate_row_height(
            csv_row["comment_text"], csv_row["response_A"], csv_row["response_B"]
        )

    last_row = ws.max_row
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{last_row}"

    judgment_col = get_column_letter(HEADERS.index("الحكم") + 1)
    dv = DataValidation(
        type="list",
        formula1='"' + ",".join(JUDGMENT_CHOICES) + '"',
        allow_blank=True,
    )
    dv.error = "الرجاء اختيار أحد الخيارات المتاحة من القائمة"
    dv.errorTitle = "قيمة غير صالحة"
    dv.prompt = "اختر: A أفضل / B أفضل / متساويان / كلاهما غير مناسب"
    dv.promptTitle = "الحكم"
    ws.add_data_validation(dv)
    dv.add(f"{judgment_col}2:{judgment_col}{last_row}")

    OUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_XLSX)
    print(f"wrote {OUT_XLSX} ({len(rows)} cases)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
