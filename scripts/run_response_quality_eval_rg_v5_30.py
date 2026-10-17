#!/usr/bin/env python3
"""Live rg-v5 quality sample: 30 comments (generate + no-response).

Does not overwrite previous workbooks. Does not change eligibility, frozen ML,
or reviewer workflow. Does not send replies to visitors.
"""

from __future__ import annotations

import importlib.util
import json
import re
import sys
import time
from collections import Counter
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT / "src"))
sys.path.insert(0, str(PROJECT_ROOT))

spec = importlib.util.spec_from_file_location(
    "eq30", PROJECT_ROOT / "scripts" / "run_response_quality_eval_30.py"
)
eq = importlib.util.module_from_spec(spec)
assert spec.loader is not None
spec.loader.exec_module(eq)

import pandas as pd  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

from config import settings  # noqa: E402
from rcs.response_generation.eligibility import EligibilityInput, determine_response_eligibility  # noqa: E402
from rcs.response_generation.service import GenerationRequest, build_live_gemini_client, generate_suggested_response  # noqa: E402
from rcs.review.taxonomy_labels import INTENT_AR, SENTIMENT_AR, TOPIC_AR, assert_canonical_coverage  # noqa: E402
from rcs.routing.engine import determine_routing  # noqa: E402
from rcs.training.gold_splits import load_supervised_gold  # noqa: E402

OUT_XLSX = PROJECT_ROOT / "response_quality_review_rg_v5_30.xlsx"
PREV_FILES = [
    PROJECT_ROOT / "response_quality_review_rg_v4_15.xlsx",
    PROJECT_ROOT / "response_quality_review_rg_v3_15.xlsx",
    PROJECT_ROOT / "response_quality_review_30.xlsx",
]
SAMPLE_SEED = 47
SAMPLE_N = 30

OPENING_PATTERNS = (
    ("شكرًا لتواصلك", re.compile(r"^شكرًا لتواصلكم?")),
    ("أهلًا بك", re.compile(r"^أهل[اًا] بك")),
    ("نشكر لك", re.compile(r"^نشكر لك")),
    ("وصلت رسالتكم", re.compile(r"^وصلت رسالتكم")),
)
CLOSING_PATTERNS = (
    ("فريق المعجم", re.compile(r"فريق المعجم\s*[.۔]?$")),
    ("نسعد بتواصلك", re.compile(r"نسعد بتواصلكم?\s*[.۔]?$")),
)

SYSTEM_COLS = (
    "comment_id",
    "comment_text",
    "intent_ar",
    "topic_ar",
    "comment_tone_ar",
    "intent_code",
    "topic_code",
    "sentiment_code",
    "routing_decision",
    "reason_code",
    "generation_status",
    "gemini_suggested_response",
    "prompt_version",
    "model_name",
)
HUMAN_COLS = (
    "human_relevance_score",
    "human_tone_score",
    "human_specificity_score",
    "human_safety_factuality_score",
    "human_edit_level",
    "generic_or_repetitive",
    "unsupported_claim",
    "human_notes",
)


def previous_eval_ids() -> set[str]:
    ids: set[str] = set()
    for path in PREV_FILES:
        if not path.exists():
            continue
        prev = pd.read_excel(path, sheet_name="Review", dtype=str)
        if "comment_id" in prev.columns:
            ids.update(prev["comment_id"].dropna().astype(str))
    return ids


def select_30(scored: pd.DataFrame, rng, *, exclude_ids: set[str]) -> tuple[pd.DataFrame, dict]:
    used: set[str] = set()
    buckets: dict[str, int] = {}
    eligible = scored[~scored["comment_id"].isin(exclude_ids)]
    skip = eligible[eligible["pred_routing"] == "NO_RESPONSE_RECOMMENDED"]
    gen = eligible[eligible["pred_routing"] == "GENERATE_DRAFT"]

    buckets["pred_NO_RESPONSE_RECOMMENDED"] = len(eq._take(rng, skip, 6, used))

    def take_intent(label: str, n: int) -> None:
        buckets[f"gold_intent:{label}"] = len(eq._take(rng, gen[gen["gold_primary_intent"] == label], n, used))

    def take_topic(label: str, n: int) -> None:
        buckets[f"gold_topic:{label}"] = len(eq._take(rng, gen[gen["gold_topic_subtype"] == label], n, used))

    def take_sent(label: str, n: int) -> None:
        buckets[f"gold_sentiment:{label}"] = len(eq._take(rng, gen[gen["gold_sentiment"] == label], n, used))

    take_intent("THANKS_PRAISE", 3)
    take_intent("SUGGESTION", 3)
    take_intent("COMPLAINT", 3)
    take_intent("INQUIRY", 3)
    take_topic("WORD_SUBMISSION", 2)
    take_topic("ENTRY_CORRECTION", 2)
    take_topic("DICTIONARY_SUGGESTION", 1)
    take_topic("TECHNICAL_ISSUE", 2)
    take_topic("WORD_OF_YEAR_NOMINATION", 2)
    take_topic("OFF_TOPIC_PERSONAL", 1)
    take_topic("GENERAL", 1)
    take_sent("POSITIVE", 1)
    take_sent("NEGATIVE", 1)
    leftover = SAMPLE_N - len(used)
    buckets["leftover_generate"] = len(eq._take(rng, gen, leftover, used))
    if len(used) < SAMPLE_N:
        buckets["leftover_any"] = len(eq._take(rng, eligible, SAMPLE_N - len(used), used))
    if len(used) != SAMPLE_N:
        raise SystemExit(f"sample size {len(used)} != {SAMPLE_N}")
    order = list(used)
    rng2 = eq.np.random.default_rng(SAMPLE_SEED + 1)
    rng2.shuffle(order)
    ordered = scored.set_index("comment_id").loc[order].reset_index()
    return ordered, {
        "selection_seed": SAMPLE_SEED,
        "bucket_fills": buckets,
        "generate_pool": int(len(gen)),
        "skip_pool": int(len(skip)),
        "excluded_previous_eval_ids": len(exclude_ids),
    }


def repetition_stats(rows: list[dict]) -> dict:
    drafts = [str(r.get("gemini_suggested_response") or "").strip() for r in rows if r.get("gemini_suggested_response")]
    openings = {name: 0 for name, _ in OPENING_PATTERNS}
    closings = {name: 0 for name, _ in CLOSING_PATTERNS}
    for draft in drafts:
        for name, pattern in OPENING_PATTERNS:
            if pattern.search(draft):
                openings[name] += 1
        for name, pattern in CLOSING_PATTERNS:
            if pattern.search(draft):
                closings[name] += 1
    norms = [" ".join(d.split()) for d in drafts]
    dupes = {text: n for text, n in Counter(norms).items() if n > 1}
    return {
        "draft_count": len(drafts),
        "exact_duplicate_drafts": dupes,
        "opening_phrase_counts": openings,
        "closing_phrase_counts": closings,
    }


def write_workbook(rows: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Review"
    ws.sheet_view.rightToLeft = True
    ws.freeze_panes = "A2"
    header_fill = PatternFill("solid", fgColor="1F4E78")
    human_fill = PatternFill("solid", fgColor="FFF2CC")
    header_font = Font(color="FFFFFF", bold=True)
    wrap = Alignment(wrap_text=True, vertical="top", horizontal="right")
    top = Alignment(vertical="top", wrap_text=True)
    thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    cols = list(SYSTEM_COLS) + list(HUMAN_COLS)
    for col_idx, name in enumerate(cols, 1):
        cell = ws.cell(1, col_idx, name)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin
    for r_i, row in enumerate(rows, 2):
        for c_i, name in enumerate(SYSTEM_COLS, 1):
            cell = ws.cell(r_i, c_i, row.get(name) if row.get(name) is not None else "")
            cell.border = thin
            cell.alignment = wrap if name in ("comment_text", "gemini_suggested_response") else top
            if name == "comment_id":
                cell.number_format = "@"
        for c_i, name in enumerate(HUMAN_COLS, 1 + len(SYSTEM_COLS)):
            cell = ws.cell(r_i, c_i, None)
            cell.fill = human_fill
            cell.border = thin
            cell.alignment = wrap if name == "human_notes" else top
        ws.row_dimensions[r_i].height = 78
    for letter, width in {
        "A": 14, "B": 48, "C": 16, "D": 20, "E": 14, "F": 16, "G": 24, "H": 14,
        "I": 24, "J": 26, "K": 16, "L": 52, "M": 12, "N": 18, "O": 14, "P": 14,
        "Q": 16, "R": 16, "S": 16, "T": 16, "U": 18, "V": 32,
    }.items():
        ws.column_dimensions[letter].width = width
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{1 + len(rows)}"
    last = 1 + len(rows)
    score_dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True)
    ws.add_data_validation(score_dv)
    score_dv.add(f"O2:R{last}")
    edit_dv = DataValidation(type="list", formula1='"SEND_AS_IS,MINOR_EDIT,MAJOR_EDIT,UNUSABLE"', allow_blank=True)
    ws.add_data_validation(edit_dv)
    edit_dv.add(f"S2:S{last}")
    yn = DataValidation(type="list", formula1='"YES,NO"', allow_blank=True)
    ws.add_data_validation(yn)
    yn.add(f"T2:U{last}")
    ins = wb.create_sheet("Instructions")
    ins["A1"] = (
        "rg-v5 response-quality review (30 comments). Automated checks are not a quality approval. "
        "Human-review columns remain blank. This is a project quality-review artifact."
    )
    ins.column_dimensions["A"].width = 110
    wb.save(OUT_XLSX)


def main() -> None:
    assert_canonical_coverage()
    if settings.RESPONSE_GENERATION_MODEL_NAME != "gemini-3.6-flash":
        raise SystemExit("Unexpected model; not overriding.")
    if settings.RESPONSE_GENERATION_PROMPT_VERSION != "rg-v5":
        raise SystemExit("Expected rg-v5.")
    exclusion = set(json.loads(eq.EXCLUSION_JSON.read_text())["test_ids"])
    gold = pd.read_csv(eq.GOLD_CSV, dtype=str, keep_default_na=False)
    human_test_ids = set(gold.loc[gold["annotation_stage"] == "human_test", "comment_id"])
    prev_ids = previous_eval_ids()
    pool = gold[
        (gold["annotation_stage"] != "human_test")
        & (~gold["comment_id"].isin(exclusion))
        & (gold["comment_text"].astype(str).str.strip() != "")
    ].copy()
    freeze = json.loads(eq.FREEZE_JSON.read_text())
    data = load_supervised_gold(eq.GOLD_CSV, include_test=False)
    intent_pipe = eq.fit_frozen_intent_pipeline(eq.remap_intent(data.intent.train))
    texts = pool["comment_text"].tolist()
    intent_labels = list(intent_pipe.classes_)
    intent_probs = intent_pipe.predict_proba(texts)
    import numpy as np

    pool["pred_intent"] = [intent_labels[int(np.argmax(p))] for p in intent_probs]
    pool["intent_conf"] = [float(np.max(p)) for p in intent_probs]
    topic_pred, topic_conf = eq.predict_topic(texts, eq.find_topic_checkpoint(), freeze)
    sent_pred, sent_conf = eq.predict_sentiment(texts, freeze)
    pool["pred_topic"] = topic_pred
    pool["topic_conf"] = topic_conf
    pool["pred_sentiment"] = sent_pred
    pool["sent_conf"] = sent_conf
    pool["pred_confidence"] = [
        min(i, t, s) for i, t, s in zip(pool["intent_conf"], pool["topic_conf"], pool["sent_conf"])
    ]
    routes, reasons, class_routes, needs_flags = [], [], [], []
    for rec in pool.itertuples(index=False):
        class_route = determine_routing(rec.pred_intent, rec.pred_topic, None, None, True)
        needs_review = rec.pred_confidence < settings.RESPONSE_LOW_CONFIDENCE_THRESHOLD
        el = determine_response_eligibility(
            EligibilityInput(
                comment_text=rec.comment_text,
                primary_intent=rec.pred_intent,
                topic_subtype=rec.pred_topic,
                sentiment=rec.pred_sentiment,
                confidence=rec.pred_confidence,
                needs_review=needs_review,
                is_content_meaningful=True,
                insufficient_context=rec.gold_annotation_status == "INSUFFICIENT_CONTEXT",
                trusted_grounding_available=False,
                classification_routing_decision=class_route.routing_decision,
            )
        )
        routes.append(el.routing_decision)
        reasons.append(el.reason_code)
        class_routes.append(class_route.routing_decision)
        needs_flags.append(needs_review)
    pool["pred_routing"] = routes
    pool["pred_reason"] = reasons
    pool["class_route"] = class_routes
    pool["pred_needs_review"] = needs_flags

    rng = eq.np.random.default_rng(SAMPLE_SEED)
    sample, selection_meta = select_30(pool, rng, exclude_ids=prev_ids)
    if set(sample["comment_id"]) & prev_ids:
        raise SystemExit("sample reused a previous evaluation comment_id")
    if set(sample["comment_id"]) & exclusion or set(sample["comment_id"]) & human_test_ids:
        raise SystemExit("sampled sealed human_test")

    inner = build_live_gemini_client()
    client = eq.CountingGeminiClient(inner)
    rows: list[dict] = []
    for rec in sample.itertuples(index=False):
        req = GenerationRequest(
            comment_id=str(rec.comment_id),
            comment_text=str(rec.comment_text),
            primary_intent=rec.pred_intent,
            topic_subtype=rec.pred_topic,
            sentiment=rec.pred_sentiment,
            confidence=float(rec.pred_confidence),
            needs_review=bool(rec.pred_needs_review),
            is_content_meaningful=True,
            insufficient_context=str(rec.gold_annotation_status) == "INSUFFICIENT_CONTEXT",
            trusted_grounding_available=False,
            classification_routing_decision=rec.class_route,
        )
        before = len(client.call_comment_ids)
        client.set_comment_id(str(rec.comment_id))
        result = generate_suggested_response(req, client=client)
        called = len(client.call_comment_ids) > before
        if called:
            time.sleep(0.4)
        rows.append(
            {
                "comment_id": str(rec.comment_id),
                "comment_text": str(rec.comment_text),
                "intent_ar": INTENT_AR.get(rec.pred_intent, rec.pred_intent),
                "topic_ar": TOPIC_AR.get(rec.pred_topic, rec.pred_topic),
                "comment_tone_ar": SENTIMENT_AR.get(rec.pred_sentiment, rec.pred_sentiment),
                "intent_code": rec.pred_intent,
                "topic_code": rec.pred_topic,
                "sentiment_code": rec.pred_sentiment,
                "routing_decision": result.routing_decision,
                "reason_code": result.reason_code,
                "generation_status": result.generation_status,
                "gemini_suggested_response": result.suggested_response,
                "prompt_version": result.prompt_version,
                "model_name": result.model_name,
                "_gemini_called": called,
                "_gold_intent": rec.gold_primary_intent,
                "_gold_topic": rec.gold_topic_subtype,
                "_gold_sentiment": getattr(rec, "gold_sentiment", ""),
            }
        )

    write_workbook(rows)
    findings: dict[str, list[str]] = {}
    for row in rows:
        issues = eq.precheck_row(
            {
                "routing_decision": row["routing_decision"],
                "generation_status": row["generation_status"],
                "gemini_suggested_response": row["gemini_suggested_response"],
            },
            bool(row["_gemini_called"]),
        )
        for issue in issues:
            findings.setdefault(issue, []).append(row["comment_id"])
    report = {
        "selection": selection_meta,
        "n": len(rows),
        "routing_distribution": dict(Counter(r["routing_decision"] for r in rows)),
        "generation_status_distribution": dict(Counter(r["generation_status"] for r in rows)),
        "intent_ar": dict(Counter(r["intent_ar"] for r in rows)),
        "topic_ar": dict(Counter(r["topic_ar"] for r in rows)),
        "comment_tone_ar": dict(Counter(r["comment_tone_ar"] for r in rows)),
        "gold_intent": dict(Counter(r["_gold_intent"] for r in rows)),
        "gold_topic": dict(Counter(r["_gold_topic"] for r in rows)),
        "gemini_calls": len(client.call_comment_ids),
        "no_response_count": sum(1 for r in rows if r["routing_decision"] == "NO_RESPONSE_RECOMMENDED"),
        "empty_drafts": sum(
            1
            for r in rows
            if r["routing_decision"] == "GENERATE_DRAFT" and not (r["gemini_suggested_response"] or "").strip()
        ),
        "model": settings.RESPONSE_GENERATION_MODEL_NAME,
        "prompt_version": settings.RESPONSE_GENERATION_PROMPT_VERSION,
        "xlsx": str(OUT_XLSX),
        "human_columns_blank": True,
        "visitor_send": False,
        "automated_precheck": findings,
        "repetition": repetition_stats(rows),
        "canonical_taxonomy": {
            "intent": list(INTENT_AR.values()),
            "topic": list(TOPIC_AR.values()),
            "sentiment": list(SENTIMENT_AR.values()),
        },
        "precheck_is_not_human_score": True,
    }
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
