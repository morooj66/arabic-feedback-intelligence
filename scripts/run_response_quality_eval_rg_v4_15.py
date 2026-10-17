#!/usr/bin/env python3
"""Live rg-v4 quality sample: 15 GENERATE_DRAFT comments only.

Does not reuse the rg-v3 15-sample. Does not change eligibility, frozen ML,
or reviewer workflow. Does not send replies to visitors.
"""

from __future__ import annotations

import importlib.util
import json
import sys
import time
from collections import Counter
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT / "src"))
sys.path.insert(0, str(PROJECT_ROOT))

spec = importlib.util.spec_from_file_location(
    "eq30", PROJECT_ROOT / "scripts" / "run_response_quality_eval_30.py"
)
eq = importlib.util.module_from_spec(spec)
assert spec.loader is not None
spec.loader.exec_module(eq)

import pandas as pd  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

from config import settings  # noqa: E402
from rcs.response_generation.eligibility import EligibilityInput, determine_response_eligibility  # noqa: E402
from rcs.response_generation.service import GenerationRequest, build_live_gemini_client, generate_suggested_response  # noqa: E402
from rcs.routing.engine import determine_routing  # noqa: E402
from rcs.training.gold_splits import load_supervised_gold  # noqa: E402

OUT_XLSX = PROJECT_ROOT / "response_quality_review_rg_v4_15.xlsx"
PREV_XLSX = PROJECT_ROOT / "response_quality_review_rg_v3_15.xlsx"
SAMPLE_SEED = 45
SAMPLE_N = 15

SYSTEM_COLS = (
    "sample_no",
    "comment_id",
    "comment_text",
    "intent",
    "topic",
    "sentiment",
    "reason_code",
    "gemini_suggested_response",
    "generation_status",
)
HUMAN_COLS = eq.HUMAN_COLS


def select_15(scored: pd.DataFrame, rng, *, exclude_ids: set[str]) -> tuple[pd.DataFrame, dict]:
    gen = scored[
        (scored["pred_routing"] == "GENERATE_DRAFT") & (~scored["comment_id"].isin(exclude_ids))
    ]
    used: set[str] = set()
    buckets: dict[str, int] = {}

    def take_intent(label: str, n: int) -> None:
        buckets[f"gold_intent:{label}"] = len(eq._take(rng, gen[gen["gold_primary_intent"] == label], n, used))

    def take_topic(label: str, n: int) -> None:
        buckets[f"gold_topic:{label}"] = len(eq._take(rng, gen[gen["gold_topic_subtype"] == label], n, used))

    take_intent("THANKS_PRAISE", 3)
    take_intent("SUGGESTION", 3)
    take_intent("COMPLAINT", 2)
    take_intent("INQUIRY", 2)
    take_topic("TECHNICAL_ISSUE", 2)
    take_topic("ENTRY_CORRECTION", 2)
    leftover = SAMPLE_N - len(used)
    buckets["leftover_generate"] = len(eq._take(rng, gen, leftover, used))
    if len(used) != SAMPLE_N:
        raise SystemExit(f"sample size {len(used)} != {SAMPLE_N}")
    if set(scored.set_index("comment_id").loc[list(used)]["pred_routing"].unique()) != {"GENERATE_DRAFT"}:
        raise SystemExit("sample includes a non-GENERATE_DRAFT row")
    order = list(used)
    rng2 = eq.np.random.default_rng(SAMPLE_SEED + 1)
    rng2.shuffle(order)
    ordered = scored.set_index("comment_id").loc[order].reset_index()
    return ordered, {
        "selection_seed": SAMPLE_SEED,
        "bucket_fills": buckets,
        "generate_pool": int(len(gen)),
        "excluded_previous_sample_ids": len(exclude_ids),
    }


def write_workbook(rows: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Review"
    ws.sheet_view.rightToLeft = True
    ws.freeze_panes = "A2"
    header_fill = PatternFill("solid", fgColor="1F4E78")
    human_fill = PatternFill("solid", fgColor="FFF2CC")
    header_font = Font(color="FFFFFF", bold=True)
    wrap = Alignment(wrap_text=True, vertical="top", horizontal="right")
    top = Alignment(vertical="top", wrap_text=True)
    thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    cols = list(SYSTEM_COLS) + list(HUMAN_COLS)
    for col_idx, name in enumerate(cols, 1):
        cell = ws.cell(1, col_idx, name)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin
    for r_i, row in enumerate(rows, 2):
        for c_i, name in enumerate(SYSTEM_COLS, 1):
            cell = ws.cell(r_i, c_i, row.get(name) if row.get(name) is not None else "")
            cell.border = thin
            cell.alignment = wrap if name in ("comment_text", "gemini_suggested_response") else top
            if name == "comment_id":
                cell.number_format = "@"
        for c_i, name in enumerate(HUMAN_COLS, 1 + len(SYSTEM_COLS)):
            cell = ws.cell(r_i, c_i, None)
            cell.fill = human_fill
            cell.border = thin
            cell.alignment = wrap if name == "human_notes" else top
        ws.row_dimensions[r_i].height = 80
    widths = {"A": 11, "B": 44, "C": 48, "D": 16, "E": 24, "F": 12, "G": 28, "H": 52, "I": 16, "J": 18, "K": 16, "L": 24, "M": 16, "N": 28, "O": 32}
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{1 + len(rows)}"
    last = 1 + len(rows)
    score_dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True)
    ws.add_data_validation(score_dv)
    score_dv.add(f"J2:L{last}")
    edit_dv = DataValidation(type="list", formula1='"SEND_AS_IS,MINOR_EDIT,MAJOR_EDIT,UNUSABLE"', allow_blank=True)
    ws.add_data_validation(edit_dv)
    edit_dv.add(f"M2:M{last}")
    hall_dv = DataValidation(type="list", formula1='"YES,NO"', allow_blank=True)
    ws.add_data_validation(hall_dv)
    hall_dv.add(f"N2:N{last}")
    ins = wb.create_sheet("Instructions")
    ins["A1"] = "rg-v4 human inspection (15 GENERATE_DRAFT comments). Automated checks are not a quality approval."
    ins.column_dimensions["A"].width = 100
    wb.save(OUT_XLSX)


def main() -> None:
    if settings.RESPONSE_GENERATION_MODEL_NAME != "gemini-3.6-flash":
        raise SystemExit("Unexpected model; not overriding.")
    if settings.RESPONSE_GENERATION_PROMPT_VERSION != "rg-v4":
        raise SystemExit("Expected rg-v4.")
    exclusion = set(json.loads(eq.EXCLUSION_JSON.read_text())["test_ids"])
    gold = pd.read_csv(eq.GOLD_CSV, dtype=str, keep_default_na=False)
    human_test_ids = set(gold.loc[gold["annotation_stage"] == "human_test", "comment_id"])
    pool = gold[
        (gold["annotation_stage"] != "human_test")
        & (~gold["comment_id"].isin(exclusion))
        & (gold["comment_text"].astype(str).str.strip() != "")
    ].copy()
    freeze = json.loads(eq.FREEZE_JSON.read_text())
    data = load_supervised_gold(eq.GOLD_CSV, include_test=False)
    intent_pipe = eq.fit_frozen_intent_pipeline(eq.remap_intent(data.intent.train))
    texts = pool["comment_text"].tolist()
    intent_labels = list(intent_pipe.classes_)
    intent_probs = intent_pipe.predict_proba(texts)
    import numpy as np

    pool["pred_intent"] = [intent_labels[int(np.argmax(p))] for p in intent_probs]
    pool["intent_conf"] = [float(np.max(p)) for p in intent_probs]
    topic_pred, topic_conf = eq.predict_topic(texts, eq.find_topic_checkpoint(), freeze)
    sent_pred, sent_conf = eq.predict_sentiment(texts, freeze)
    pool["pred_topic"] = topic_pred
    pool["topic_conf"] = topic_conf
    pool["pred_sentiment"] = sent_pred
    pool["sent_conf"] = sent_conf
    pool["pred_confidence"] = [
        min(i, t, s) for i, t, s in zip(pool["intent_conf"], pool["topic_conf"], pool["sent_conf"])
    ]
    routes = []
    reasons = []
    class_routes = []
    needs_flags = []
    for rec in pool.itertuples(index=False):
        class_route = determine_routing(rec.pred_intent, rec.pred_topic, None, None, True)
        needs_review = rec.pred_confidence < settings.RESPONSE_LOW_CONFIDENCE_THRESHOLD
        el = determine_response_eligibility(
            EligibilityInput(
                comment_text=rec.comment_text,
                primary_intent=rec.pred_intent,
                topic_subtype=rec.pred_topic,
                sentiment=rec.pred_sentiment,
                confidence=rec.pred_confidence,
                needs_review=needs_review,
                is_content_meaningful=True,
                insufficient_context=rec.gold_annotation_status == "INSUFFICIENT_CONTEXT",
                trusted_grounding_available=False,
                classification_routing_decision=class_route.routing_decision,
            )
        )
        routes.append(el.routing_decision)
        reasons.append(el.reason_code)
        class_routes.append(class_route.routing_decision)
        needs_flags.append(needs_review)
    pool["pred_routing"] = routes
    pool["pred_reason"] = reasons
    pool["class_route"] = class_routes
    pool["pred_needs_review"] = needs_flags

    prev_ids: set[str] = set()
    if PREV_XLSX.exists():
        prev = pd.read_excel(PREV_XLSX, sheet_name="Review", dtype=str)
        if "comment_id" in prev.columns:
            prev_ids = set(prev["comment_id"].dropna().astype(str))
    rng = eq.np.random.default_rng(SAMPLE_SEED)
    sample, selection_meta = select_15(pool, rng, exclude_ids=prev_ids)
    if set(sample["comment_id"]) & prev_ids:
        raise SystemExit("sample reused a previous rg-v3 comment_id")
    if set(sample["comment_id"]) & exclusion or set(sample["comment_id"]) & human_test_ids:
        raise SystemExit("sampled sealed human_test")
    if (sample["pred_routing"] != "GENERATE_DRAFT").any():
        raise SystemExit("NO_RESPONSE_RECOMMENDED leaked into the 15-sample")

    inner = build_live_gemini_client()
    client = eq.CountingGeminiClient(inner)
    rows: list[dict] = []
    for i, rec in enumerate(sample.itertuples(index=False), 1):
        req = GenerationRequest(
            comment_id=str(rec.comment_id),
            comment_text=str(rec.comment_text),
            primary_intent=rec.pred_intent,
            topic_subtype=rec.pred_topic,
            sentiment=rec.pred_sentiment,
            confidence=float(rec.pred_confidence),
            needs_review=bool(rec.pred_needs_review),
            is_content_meaningful=True,
            insufficient_context=str(rec.gold_annotation_status) == "INSUFFICIENT_CONTEXT",
            trusted_grounding_available=False,
            classification_routing_decision=rec.class_route,
        )
        before = len(client.call_comment_ids)
        client.set_comment_id(str(rec.comment_id))
        result = generate_suggested_response(req, client=client)
        called = len(client.call_comment_ids) > before
        if called:
            time.sleep(0.4)
        row = {
            "sample_no": i,
            "comment_id": str(rec.comment_id),
            "comment_text": str(rec.comment_text),
            "intent": rec.pred_intent,
            "topic": rec.pred_topic,
            "sentiment": rec.pred_sentiment,
            "reason_code": result.reason_code,
            "gemini_suggested_response": result.suggested_response,
            "generation_status": result.generation_status,
            "routing_decision": result.routing_decision,
            "_gemini_called": called,
            "_gold_intent": rec.gold_primary_intent,
            "_gold_topic": rec.gold_topic_subtype,
        }
        rows.append(row)

    write_workbook(rows)
    findings: dict[str, list[int]] = {}
    for row in rows:
        issues = eq.precheck_row(row, bool(row["_gemini_called"]))
        for issue in issues:
            findings.setdefault(issue, []).append(int(row["sample_no"]))
    report = {
        "selection": selection_meta,
        "n": len(rows),
        "routing_distribution": dict(Counter(r["routing_decision"] for r in rows)),
        "gold_intent": dict(Counter(r["_gold_intent"] for r in rows)),
        "gold_topic": dict(Counter(r["_gold_topic"] for r in rows)),
        "intent_distribution": dict(Counter(r["intent"] for r in rows)),
        "topic_distribution": dict(Counter(r["topic"] for r in rows)),
        "generation_status_distribution": dict(Counter(r["generation_status"] for r in rows)),
        "gemini_calls": len(client.call_comment_ids),
        "model": settings.RESPONSE_GENERATION_MODEL_NAME,
        "prompt_version": settings.RESPONSE_GENERATION_PROMPT_VERSION,
        "pairs": [
            {"comment_id": r["comment_id"], "comment_text": r["comment_text"], "gemini_suggested_response": r["gemini_suggested_response"]}
            for r in rows
        ],
        "human_columns_blank": True,
        "visitor_send": False,
        "eligibility_unchanged": True,
        "automated_precheck": findings,
        "precheck_is_not_human_score": True,
    }
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
